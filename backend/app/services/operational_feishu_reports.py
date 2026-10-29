"""Audited Feishu reports for the Collection Rate V3 and PDD V2 dashboards.

This module deliberately keeps source selection and metric rules separate. It
shares only the report lifecycle and the already approved Feishu transport
primitives with the Sem Movimentacao pilot.
"""
from __future__ import annotations

import hashlib
import hmac
import io
import json
import logging
import os
import re
import stat
import tempfile
import time as monotonic_clock
import uuid
import zipfile
from datetime import date, datetime, time, timedelta, timezone
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any

import httpx
from sqlalchemy import Connection, Engine, text

from app.core.config import PROJECT_ROOT, Settings, get_settings
from app.services.data_foundation.contracts import get_contract
from app.services.data_foundation.mapping_revision import read_revision
from app.services.data_foundation.queries import _assert_map_available, _query_plan
from app.services.feishu_publication_notifications import RECIPIENT_EMAIL
from app.services.feishu_report_recipients import recipient_from_audit, recipient_options, resolve_recipient
from app.services.no_movement_feishu_reports import (
    FILES_URL, HTTP_TIMEOUT, IMAGES_URL, MAX_FEISHU_MESSAGE_BYTES,
    MAX_FILE_BYTES, MESSAGE_URL, RECIPIENT_KEY, RECIPIENT_LABEL, TOKEN_URL,
    _FeishuRejected, _feishu_interactive_content, _feishu_payload,
)
from app.services.pdd_collection_failure_operational_analysis import (
    build_bilingual_analysis as build_pdd_analysis,
)

logger = logging.getLogger(__name__)
TABLE = "operational_feishu_reports"
RULES_VERSION = "feishu-operational-v1"
REPORT_FETCH_BATCH_SIZE = 2_000
MAX_GENERATION_SECONDS = 240.0
SOURCES: dict[str, dict[str, Any]] = {
    "collection_rate": {
        "contract": 3, "table": "fact_collection_rate_v3", "flag": "feishu_collection_rate_report_enabled",
        "date_field": "deadline_at", "sheet_mode": "single", "sheet_name": "Dados",
        "analysis": "collection_rate",
    },
    "pdd_collection_failure": {
        "contract": 2, "table": "fact_pdd_collection_failure_v2", "flag": "feishu_pdd_collection_failure_report_enabled",
        "date_field": "data_date", "sheet_mode": "daily", "analysis": "pdd_collection_failure",
    },
}


class OperationalFeishuReportError(ValueError):
    def __init__(self, code: str, status_code: int = 422) -> None:
        super().__init__(code)
        self.code = code
        self.status_code = status_code


def _check_generation_deadline(deadline: float) -> None:
    if monotonic_clock.monotonic() > deadline:
        raise OperationalFeishuReportError("operational_feishu_generation_timeout", 504)


def _now() -> datetime:
    return datetime.now(timezone.utc).replace(tzinfo=None)


def _uuid(value: str) -> bytes:
    try:
        return uuid.UUID(value).bytes
    except (ValueError, TypeError, AttributeError):
        raise OperationalFeishuReportError("operational_feishu_report_id_invalid", 422) from None


def _source(source_id: str) -> dict[str, Any]:
    source = SOURCES.get(source_id)
    if source is None:
        raise OperationalFeishuReportError("operational_feishu_source_invalid", 404)
    return source


def _feature_enabled(source_id: str, settings: Settings) -> bool:
    return bool(getattr(settings, _source(source_id)["flag"], False))


def _require_feature_enabled(source_id: str, settings: Settings) -> None:
    if not _feature_enabled(source_id, settings):
        raise OperationalFeishuReportError("operational_feishu_send_disabled", 503)


def _artifact_root(staging_directory: Path) -> Path:
    root = staging_directory.expanduser().resolve()
    if not root.is_absolute() or root.is_relative_to(PROJECT_ROOT.resolve()):
        raise OperationalFeishuReportError("operational_feishu_storage_unavailable", 503)
    directory = root / ".operational-feishu-reports"
    if directory.is_symlink():
        raise OperationalFeishuReportError("operational_feishu_storage_unavailable", 503)
    try:
        directory.mkdir(mode=0o700, parents=True, exist_ok=True)
        directory = directory.resolve()
        directory.relative_to(root)
        directory.chmod(0o700)
    except (OSError, ValueError):
        raise OperationalFeishuReportError("operational_feishu_storage_unavailable", 503) from None
    return directory


def _artifact_path(staging_directory: Path, report_id: uuid.UUID, ext: str) -> Path:
    if ext not in {"xlsx", "png"}:
        raise OperationalFeishuReportError("operational_feishu_artifact_invalid", 409)
    return _artifact_root(staging_directory) / f"{report_id.hex}.{ext}"


def _store(staging_directory: Path, report_id: uuid.UUID, artifacts: dict[str, bytes]) -> None:
    try:
        for ext, value in artifacts.items():
            if ext not in {"xlsx", "png"} or not value or len(value) > MAX_FILE_BYTES:
                raise OperationalFeishuReportError("operational_feishu_artifact_size_invalid", 413)
            path = _artifact_path(staging_directory, report_id, ext)
            fd = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
            with os.fdopen(fd, "wb") as handle:
                handle.write(value)
                handle.flush()
                os.fsync(handle.fileno())
            path.chmod(0o600)
    except Exception as error:
        for ext in ("xlsx", "png"):
            try:
                _artifact_path(staging_directory, report_id, ext).unlink(missing_ok=True)
            except OSError:
                pass
        if isinstance(error, OperationalFeishuReportError):
            raise
        raise OperationalFeishuReportError("operational_feishu_storage_unavailable", 503) from None


def _read_artifact(staging_directory: Path, report_id: uuid.UUID, ext: str, expected: object) -> bytes:
    path = _artifact_path(staging_directory, report_id, ext)
    try:
        info = path.stat()
        if path.is_symlink() or not stat.S_ISREG(info.st_mode) or info.st_size <= 0 or info.st_size > MAX_FILE_BYTES:
            raise OSError
        content = path.read_bytes()
    except OSError:
        raise OperationalFeishuReportError("operational_feishu_artifacts_changed", 409) from None
    digest = hashlib.sha256(content).hexdigest()
    if not isinstance(expected, str) or not hmac.compare_digest(digest, expected):
        raise OperationalFeishuReportError("operational_feishu_artifacts_changed", 409)
    return content


def _open_verified_artifact(staging_directory: Path, report_id: uuid.UUID, ext: str, expected: object):
    path = _artifact_path(staging_directory, report_id, ext)
    try:
        handle = path.open("rb")
        info = os.fstat(handle.fileno())
        if path.is_symlink() or not stat.S_ISREG(info.st_mode) or info.st_size <= 0 or info.st_size > MAX_FILE_BYTES:
            handle.close()
            raise OSError
        digest = hashlib.sha256()
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
        if not isinstance(expected, str) or not hmac.compare_digest(digest.hexdigest(), expected):
            handle.close()
            raise OperationalFeishuReportError("operational_feishu_artifacts_changed", 409)
        handle.seek(0)
        return handle
    except OperationalFeishuReportError:
        raise
    except OSError:
        raise OperationalFeishuReportError("operational_feishu_artifacts_changed", 409) from None


def _publication(engine: Engine, source_id: str, publication_id: str, connection: Connection | None = None) -> dict[str, Any]:
    source = _source(source_id)
    binary_id = _uuid(publication_id)
    if connection is None:
        with engine.connect() as active:
            return _publication(engine, source_id, publication_id, active)
    # These are daily sources. The dashboards do not use source_snapshot_head:
    # collection_rate resolves the current publication per date by greatest
    # published version, while PDD reads the source_daily_head pointer per date.
    # Accept a publication only when every non-empty date it owns is still the
    # active partition under that same source-specific rule. This also rejects
    # a superseded publication without requiring a snapshot-level head row.
    if source_id == "collection_rate":
        current_partition = """
            NOT EXISTS (
                SELECT 1
                FROM source_publication_dates AS candidate_date
                WHERE candidate_date.source_id = p.source_id
                  AND candidate_date.publication_id = p.publication_id
                  AND candidate_date.row_count > 0
                  AND EXISTS (
                      SELECT 1
                      FROM source_publications AS newer
                      JOIN source_publication_dates AS newer_date
                        ON newer_date.source_id = newer.source_id
                       AND newer_date.publication_id = newer.publication_id
                       AND newer_date.data_date = candidate_date.data_date
                       AND newer_date.row_count > 0
                      WHERE newer.source_id = p.source_id
                        AND newer.contract_version = p.contract_version
                        AND newer.publication_state = 'PUBLISHED'
                        AND newer.version_no > p.version_no
                  )
            )
            AND EXISTS (
                SELECT 1 FROM source_publication_dates AS active_date
                WHERE active_date.source_id = p.source_id
                  AND active_date.publication_id = p.publication_id
                  AND active_date.row_count > 0
            )
        """
    else:
        current_partition = """
            NOT EXISTS (
                SELECT 1
                FROM source_publication_dates AS candidate_date
                WHERE candidate_date.source_id = p.source_id
                  AND candidate_date.publication_id = p.publication_id
                  AND candidate_date.row_count > 0
                  AND NOT EXISTS (
                      SELECT 1 FROM source_daily_head AS active_head
                      WHERE active_head.source_id = candidate_date.source_id
                        AND active_head.data_date = candidate_date.data_date
                        AND active_head.publication_id = candidate_date.publication_id
                  )
            )
            AND EXISTS (
                SELECT 1 FROM source_publication_dates AS active_date
                JOIN source_daily_head AS active_head
                  ON active_head.source_id = active_date.source_id
                 AND active_head.data_date = active_date.data_date
                 AND active_head.publication_id = active_date.publication_id
                WHERE active_date.source_id = p.source_id
                  AND active_date.publication_id = p.publication_id
                  AND active_date.row_count > 0
            )
        """
    row = connection.execute(text(f"""
        SELECT p.publication_id, p.version_no, p.contract_version, p.map_version_id,
               p.period_start, p.snapshot_as_of, p.published_at, p.publication_state
        FROM source_publications AS p
        WHERE p.source_id=:source_id AND p.publication_id=:publication_id
          AND p.publication_state='PUBLISHED' AND p.contract_version=:contract_version
          AND p.row_count = (
              SELECT COALESCE(SUM(metadata.row_count), 0)
              FROM source_publication_dates AS metadata
              WHERE metadata.source_id=p.source_id AND metadata.publication_id=p.publication_id
          )
          AND p.period_start = (
              SELECT MIN(metadata.data_date)
              FROM source_publication_dates AS metadata
              WHERE metadata.source_id=p.source_id AND metadata.publication_id=p.publication_id
          )
          AND p.period_end = (
              SELECT MAX(metadata.data_date)
              FROM source_publication_dates AS metadata
              WHERE metadata.source_id=p.source_id AND metadata.publication_id=p.publication_id
          )
          AND {current_partition}
    """), {"source_id": source_id, "publication_id": binary_id, "contract_version": source["contract"]}).mappings().first()
    if row is None:
        raise OperationalFeishuReportError("operational_feishu_publication_not_current", 409)
    result = dict(row)
    map_id = bytes(result.get("map_version_id") or b"")
    if len(map_id) != 16:
        raise OperationalFeishuReportError("operational_feishu_publication_mapping_invalid", 409)
    result["map_version_id"] = map_id
    return result


def _assert_report_publication_matches(report: Any, publication: dict[str, Any]) -> None:
    if (
        int(report["publication_version"]) != int(publication["version_no"])
        or int(report["contract_version"]) != int(publication["contract_version"])
    ):
        raise OperationalFeishuReportError("operational_feishu_publication_not_current", 409)
    if bytes(report["map_version_id"]) != bytes(publication["map_version_id"]):
        raise OperationalFeishuReportError("operational_feishu_mapping_changed", 409)


def _regions(engine: Engine, map_id: bytes, connection: Connection | None = None) -> list[str]:
    if connection is None:
        with engine.connect() as active:
            return _regions(engine, map_id, active)
    rows = connection.execute(text("""
        SELECT DISTINCT regional_name FROM base_mapping_entries
        WHERE map_version_id=:map_id AND regional_name IS NOT NULL AND regional_name<>''
        ORDER BY regional_name
    """), {"map_id": map_id}).scalars().all()
    return [str(value) for value in rows]


def _rows_query(engine: Engine, source_id: str, publication: dict[str, Any], region_code: str | None):
    source = _source(source_id)
    contract = get_contract(source_id, source["contract"])
    columns = [column.key for column in contract.columns]
    selection = ", ".join(f"f.`{key}` AS `{key}`" for key in columns)
    if region_code:
        region_predicate = " AND BINARY m.regional_name=BINARY :region_code" if engine.dialect.name == "mysql" else " AND m.regional_name COLLATE BINARY=:region_code"
    else:
        region_predicate = ""
    sheet_field = "f.source_sheet_name" if source["sheet_mode"] == "daily" else "NULL AS source_sheet_name"
    query = text(f"""
        SELECT {selection}, f.source_row_no, {sheet_field},
               f.resolved_base_code, f.mapping_resolution,
               m.regional_name AS mapped_region, m.base_name AS mapped_base_name,
               m.rm_responsible AS mapped_rm
        FROM {source['table']} f
        LEFT JOIN base_mapping_entries m
          ON m.map_version_id=f.map_version_id AND m.base_code=f.resolved_base_code
        WHERE f.source_id=:source_id AND f.publication_id=:publication_id
          AND f.map_version_id=:map_id {region_predicate}
        ORDER BY {('f.data_date, f.source_row_no' if source['sheet_mode']=='daily' else 'f.deadline_at, f.source_row_no')}
    """)
    params: dict[str, object] = {"source_id": source_id, "publication_id": bytes(publication["publication_id"]), "map_id": bytes(publication["map_version_id"])}
    if region_code:
        params["region_code"] = region_code
    return query, params


def _iter_rows(connection: Connection, query, params: dict[str, object]):
    """Yield ordered fact rows without asking SQLAlchemy to buffer the result."""
    result = connection.execution_options(yield_per=REPORT_FETCH_BATCH_SIZE).execute(query, params).mappings()
    try:
        for row in result:
            yield row
    finally:
        result.close()


def _fetch_rows(engine: Engine, source_id: str, publication: dict[str, Any], region_code: str | None, connection: Connection | None = None) -> list[dict[str, Any]]:
    query, params = _rows_query(engine, source_id, publication, region_code)
    if connection is not None:
        return [dict(row) for row in _iter_rows(connection, query, params)]
    with engine.connect() as active:
        return [dict(row) for row in _iter_rows(active, query, params)]


def _digest_add(digest: Any, row: dict[str, Any], columns: list[str]) -> None:
    payload = json.dumps([_stable(row.get(key)) for key in columns], ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    digest.update(len(payload).to_bytes(8, "big"))
    digest.update(payload)


class _ScopedContentDigest:
    """Bounded-memory digest of the exact source-contract rows in UI order."""

    CHUNK_BYTES = 1024 * 1024

    def __init__(self) -> None:
        self._buffer = bytearray()
        self._chunk_digests: list[bytes] = []
        self._byte_count = 0

    def add(self, row: dict[str, Any], columns: list[str]) -> None:
        payload = json.dumps([_stable(row.get(key)) for key in columns], ensure_ascii=False, separators=(",", ":")).encode("utf-8") + b"\n"
        self._byte_count += len(payload)
        offset = 0
        while offset < len(payload):
            remaining = self.CHUNK_BYTES - len(self._buffer)
            take = min(remaining, len(payload) - offset)
            self._buffer.extend(payload[offset:offset + take])
            offset += take
            if len(self._buffer) == self.CHUNK_BYTES:
                self._chunk_digests.append(hashlib.sha256(self._buffer).digest())
                self._buffer.clear()

    def hexdigest(self) -> str:
        chunk_digests = list(self._chunk_digests)
        if self._buffer:
            chunk_digests.append(hashlib.sha256(self._buffer).digest())
        manifest = b"".join(chunk_digests) + self._byte_count.to_bytes(8, "big")
        return hashlib.sha256(manifest).hexdigest()


def _digest_rows(connection: Connection, engine: Engine, source_id: str, publication: dict[str, Any], region_code: str | None) -> tuple[bytes, int]:
    contract = get_contract(source_id, _source(source_id)["contract"])
    columns = [column.key for column in contract.columns]
    digest = hashlib.sha256()
    map_id = bytes(publication["map_version_id"])
    digest.update(map_id)
    query, params = _rows_query(engine, source_id, publication, region_code)
    count = 0
    for row in _iter_rows(connection, query, params):
        _digest_add(digest, row, columns)
        count += 1
    return digest.digest(), count


def get_options(engine: Engine, *, source_id: str, publication_id: str, region_code: str | None = None, settings: Settings | None = None) -> dict[str, object]:
    config = settings or get_settings()
    _require_feature_enabled(source_id, config)
    publication = _publication(engine, source_id, publication_id)
    source = _source(source_id)
    published_at = publication.get("published_at")
    if isinstance(published_at, (datetime, date)):
        published_text = published_at.isoformat(sep=" ") if isinstance(published_at, datetime) else published_at.isoformat()
    else:
        published_text = str(published_at) if published_at else None
    regions = _regions(engine, bytes(publication["map_version_id"]))
    if region_code is not None and region_code not in regions:
        raise OperationalFeishuReportError("operational_feishu_region_invalid", 422)
    return {
        "sourceId": source_id, "publicationId": str(uuid.UUID(bytes=bytes(publication["publication_id"]))),
        "publicationVersion": int(publication["version_no"]), "contractVersion": int(publication["contract_version"]),
        "mapVersionId": str(uuid.UUID(bytes=bytes(publication["map_version_id"]))),
        "publishedAt": published_text, "regions": regions,
        "scopeRegionCode": region_code,
        "recipients": recipient_options(region_code, config),
        "sendEnabled": True,
    }


def get_availability(
    engine: Engine | None,
    *,
    source_id: str,
    publication_id: str,
    settings: Settings | None = None,
) -> dict[str, bool]:
    """Expose only per-source send availability to authorized dashboard UI."""
    config = settings or get_settings()
    enabled = _feature_enabled(source_id, config)
    if not enabled:
        # Fail closed without querying publication metadata while disabled.
        return {"enabled": False, "publicationValid": False}
    try:
        if engine is None:
            raise OperationalFeishuReportError("operational_feishu_publication_not_current", 409)
        _publication(engine, source_id, publication_id)
    except OperationalFeishuReportError:
        return {"enabled": True, "publicationValid": False}
    return {"enabled": True, "publicationValid": True}


def _stable(value: object) -> object:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, bytes):
        return value.hex()
    return value


def _digest(rows: list[dict[str, Any]], columns: list[str], map_id: bytes) -> bytes:
    digest = hashlib.sha256()
    digest.update(map_id)
    for row in rows:
        _digest_add(digest, row, columns)
    return digest.digest()


def _as_date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value)[:10])


def _percent(value: object) -> str:
    if value is None:
        return "—"
    try:
        amount = Decimal(str(value)) * Decimal(100)
    except Exception:
        return "—"
    return f"{amount.quantize(Decimal('0.1'), rounding=ROUND_HALF_UP):.1f}%"


def _new_analysis_totals(source_id: str) -> dict[str, Any]:
    if source_id == "collection_rate":
        return {"rowCount": 0, "attempted": 0, "planned": 0, "totalOrders": 0, "unmappedPlanned": 0, "groups": {}, "regions": {}}
    return {"rowCount": 0, "totalOrders": 0, "unscanned": 0, "unmappedTotal": 0, "unmappedUnscanned": 0, "groups": {}}


def _accumulate_analysis_row(source_id: str, totals: dict[str, Any], row: dict[str, Any], operational_date: date | None) -> None:
    source = _source(source_id)
    value = row.get(source["date_field"])
    if value is None or (operational_date is not None and _as_date(value) != operational_date):
        return
    totals["rowCount"] += 1
    mapped = bool(row.get("resolved_base_code") and row.get("mapped_region") and row.get("mapping_resolution") in {"NAME_APPROVED", "CODE_MATCH"})
    if source_id == "collection_rate":
        attempted = int(row.get("collected_plus_attempts_quantity") or 0)
        planned = int(row.get("quantity_to_collect") or 0)
        totals["attempted"] += attempted
        totals["planned"] += planned
        totals["totalOrders"] += int(row.get("order_quantity") or 0)
        region_name = str(row.get("mapped_region") or "Sem regional") if mapped else "Sem regional"
        regional = totals["regions"].setdefault(region_name, {"region": region_name, "attempted": 0, "planned": 0})
        regional["attempted"] += attempted
        regional["planned"] += planned
        if mapped:
            code = str(row["resolved_base_code"])
            item = totals["groups"].setdefault(code, {"code": code, "name": row.get("mapped_base_name") or code, "region": row.get("mapped_region"), "responsible": row.get("mapped_rm"), "attempted": 0, "planned": 0})
            item["attempted"] += attempted
            item["planned"] += planned
        else:
            totals["unmappedPlanned"] += planned
        return
    total = int(row.get("total_orders") or 0)
    unscanned = int(row.get("orders_unscanned_at_collection") or 0)
    totals["totalOrders"] += total
    totals["unscanned"] += unscanned
    if mapped:
        code = str(row["resolved_base_code"])
        item = totals["groups"].setdefault(code, {"code": code, "name": row.get("mapped_base_name") or code, "region": row.get("mapped_region"), "owner": row.get("mapped_rm"), "total": 0, "unscanned": 0})
        item["total"] += total
        item["unscanned"] += unscanned
    else:
        totals["unmappedTotal"] += total
        totals["unmappedUnscanned"] += unscanned


def _build_collection_rate_attempt_analysis(
    *, operational_date: date, attempted: int, planned: int, total_orders: int,
    bases: list[dict[str, Any]], unmapped_planned: int,
    regional_totals: list[dict[str, Any]] | None = None,
) -> tuple[str, dict[str, Any]]:
    overall_rate = Decimal(attempted) / Decimal(planned) if planned > 0 else None

    def rate(numerator: int, denominator: int) -> Decimal | None:
        return Decimal(numerator) / Decimal(denominator) if denominator > 0 else None

    mapped_regions: dict[str, dict[str, Any]] = {
        str(item["region"]): {**item, "bases": []}
        for item in (regional_totals or [])
    }
    for base in bases:
        region_name = str(base.get("region") or "").strip()
        if not region_name:
            continue
        group = mapped_regions.setdefault(region_name, {"region": region_name, "attempted": 0, "planned": 0, "bases": []})
        if int(base.get("planned") or 0) > 0:
            group["bases"].append(base)

    def impact(numerator: int, denominator: int) -> Decimal:
        if overall_rate is None or denominator <= 0:
            return Decimal(0)
        return max(Decimal(0), overall_rate * denominator - numerator)

    regions = [
        {**group, "rate": rate(group["attempted"], group["planned"]),
         "deviationImpact": impact(group["attempted"], group["planned"])}
        for group in mapped_regions.values() if group["planned"] > 0
    ]
    best = sorted(regions, key=lambda item: (-item["rate"], -item["planned"], item["region"]))[:3]
    worst = sorted(regions, key=lambda item: (item["rate"], -item["planned"], item["region"]))[:3]
    regional_priorities = sorted(
        (item for item in regions if overall_rate is not None and item["rate"] < overall_rate),
        key=lambda item: (-item["deviationImpact"], -item["planned"], item["region"]),
    )[:5]
    base_priorities = []
    for base in bases:
        base_rate = rate(int(base.get("attempted") or 0), int(base.get("planned") or 0))
        if base_rate is not None and overall_rate is not None and base_rate < overall_rate:
            base_priorities.append({**base, "rate": base_rate, "deviationImpact": impact(int(base.get("attempted") or 0), int(base.get("planned") or 0))})
    base_priorities.sort(key=lambda item: (-item["deviationImpact"], -item["planned"], str(item.get("region") or ""), str(item.get("name") or "")))
    base_priorities = base_priorities[:8]

    def pct(value: Decimal | None) -> str:
        return _percent(value).replace(".", ",") if value is not None else "—"

    def volume(value: object) -> str:
        return f"{int(value or 0):,}".replace(",", ".")

    date_label = operational_date.strftime("%d/%m/%Y")
    lines = [
        f"📊 **TAXA DE TENTATIVA {date_label} | 尝试揽收率**",
        f"Geral: **{pct(overall_rate)}** | Coleta + tentativas: **{volume(attempted)}** / quantidade a coletar: **{volume(planned)}**",
        f"整体：**{pct(overall_rate)}** | 揽收 + 尝试：**{volume(attempted)}** / 待揽收数量：**{volume(planned)}**",
        f"Volume operacional (pedidos): **{volume(total_orders)}**",
        f"运营订单量：**{volume(total_orders)}**",
        "Meta oficial única da Taxa de Tentativa não identificada; a meta da Taxa de Coleta não foi reutilizada.",
        "未确认尝试揽收率统一官方目标；未套用普通揽收率目标。",
        "",
        "**Melhores regionais / 区域较高比率**",
        *(f"• **{item['region']}** — {pct(item['rate'])} | volume {volume(item['planned'])}" for item in best),
        "",
        "**Menores taxas observadas / 观察到的较低比率**",
        *(f"• **{item['region']}** — {pct(item['rate'])} | volume {volume(item['planned'])}" for item in worst),
        "",
        "**Desvios em relação ao consolidado / 与汇总值的偏差**",
        *(f"• **{item['region']}** — {pct(item['rate'])} | volume {volume(item['planned'])} | desvio relativo estimado {volume(round(item['deviationImpact']))}" for item in regional_priorities),
        "",
        "**Bases para acompanhamento / 网点跟进**",
        *(f"• **{item.get('name') or item.get('code')}** ({item.get('code') or 'código não informado'}) — {item.get('region')} | {pct(item['rate'])} | volume {volume(item.get('planned'))} | RM: {item.get('responsible') or 'Sem RM'}" for item in base_priorities),
        "",
        "Os desvios são comparações estatísticas com o consolidado, não comprovam causas, meta atingida ou backlog observado. Investigue os registros e monitore novamente o mesmo indicador.",
        "偏差仅为与汇总值的统计比较，不能证明运营原因、目标达成或实际积压。请核查记录并持续监控同一指标。",
    ]
    if not best:
        lines.append("• Sem volume regional mapeado.")
    if not worst:
        lines.append("• Sem volume regional mapeado.")
    if not regional_priorities:
        lines.append("• Nenhuma regional abaixo do consolidado neste recorte.")
    if not base_priorities:
        lines.append("• Sem bases abaixo do consolidado com volume positivo.")
    if unmapped_planned > 0:
        lines.extend([
            f"Volume sem mapeamento mantido no consolidado e não atribuído a bases: {volume(unmapped_planned)}.",
            f"未映射数量仍计入汇总且未分配到网点：{volume(unmapped_planned)}。",
        ])
    return "\n".join(lines), {
        "metric": {"numeratorField": "collected_plus_attempts_quantity", "denominatorField": "quantity_to_collect", "attempted": attempted, "planned": planned, "rate": str(overall_rate) if overall_rate is not None else None},
        "regions": [
            {"region": item["region"], "rate": item["rate"], "planned": item["planned"], "bases": [
                {"baseName": base.get("name"), "baseCode": base.get("code"), "responsible": base.get("responsible"), "rate": rate(int(base.get("attempted") or 0), int(base.get("planned") or 0))}
                for base in item["bases"] if overall_rate is not None and rate(int(base.get("attempted") or 0), int(base.get("planned") or 0)) is not None and rate(int(base.get("attempted") or 0), int(base.get("planned") or 0)) < overall_rate
            ]} for item in regional_priorities
        ],
        "unmappedPlanned": unmapped_planned,
        "targetRate": None,
    }


def _finish_analysis(source_id: str, publication: dict[str, Any], region_code: str | None, operational_date: date, totals: dict[str, Any]) -> tuple[str, dict[str, Any], dict[str, int]]:
    if not totals["rowCount"]:
        raise OperationalFeishuReportError("operational_feishu_no_records", 404)
    if source_id == "collection_rate":
        message, result = _build_collection_rate_attempt_analysis(
            operational_date=operational_date, attempted=totals["attempted"], planned=totals["planned"],
            total_orders=totals["totalOrders"], bases=list(totals["groups"].values()),
            unmapped_planned=totals["unmappedPlanned"], regional_totals=list(totals.get("regions", {}).values()),
        )
        summary = {"attempted": totals["attempted"], "planned": totals["planned"], "totalOrders": totals["totalOrders"], "rateNumerator": totals["attempted"], "rateDenominator": totals["planned"]}
    else:
        message, result = build_pdd_analysis(
            operational_date=operational_date, publication_version=int(publication["version_no"]),
            total_orders=totals["totalOrders"], orders_unscanned=totals["unscanned"],
            bases=list(totals["groups"].values()), unmapped_total=totals["unmappedTotal"],
            unmapped_unscanned=totals["unmappedUnscanned"],
        )
        summary = {"collected": totals["totalOrders"] - totals["unscanned"], "planned": totals["totalOrders"], "totalOrders": totals["totalOrders"], "rateNumerator": totals["unscanned"], "rateDenominator": totals["totalOrders"]}
    if region_code:
        message = f"📍 Regional: **{region_code}** | 区域: **{region_code}**\n\n" + message
    return message, {"operationalDate": operational_date.isoformat(), "result": result}, summary


def _analysis(source_id: str, publication: dict[str, Any], rows: list[dict[str, Any]], region_code: str | None, operational_date: date) -> tuple[str, dict[str, Any], dict[str, int]]:
    totals = _new_analysis_totals(source_id)
    for row in rows:
        _accumulate_analysis_row(source_id, totals, row, operational_date)
    return _finish_analysis(source_id, publication, region_code, operational_date, totals)


def _xlsx(source_id: str, rows: list[dict[str, Any]], operational_date: date, region_code: str | None) -> tuple[str, bytes]:
    from openpyxl import Workbook

    source = _source(source_id)
    contract = get_contract(source_id, source["contract"])
    workbook = Workbook(write_only=True)
    columns = [column.key for column in contract.columns]
    if source["sheet_mode"] == "single":
        grouped = [(source["sheet_name"], rows)]
    else:
        by_day: dict[date, list[dict[str, Any]]] = {}
        for row in rows:
            if row.get("data_date") is not None:
                by_day.setdefault(_as_date(row["data_date"]), []).append(row)
        grouped = [(day.isoformat(), by_day[day]) for day in sorted(by_day)]
    for title, data_rows in grouped:
        sheet = workbook.create_sheet(title=title[:31])
        sheet.append([column.header for column in contract.columns])
        for row in data_rows:
            # Preserve the complete official contract width even when its
            # trailing columns are null in every row of the scoped extract.
            sheet.append([row.get(key) if row.get(key) is not None else "" for key in columns])
    workbook.properties.title = f"{source_id} — relatório oficial"
    workbook.properties.subject = f"Publication scoped report; region={region_code or 'ALL'}"
    stable = datetime.combine(operational_date, time.min)
    workbook.properties.created = stable
    workbook.properties.modified = stable
    buffer = io.BytesIO()
    workbook.save(buffer)
    raw = buffer.getvalue()
    canonical = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as source_zip:
        with zipfile.ZipFile(canonical, "w", zipfile.ZIP_DEFLATED) as target_zip:
            for item in source_zip.infolist():
                stable_item = zipfile.ZipInfo(item.filename, date_time=(1980, 1, 1, 0, 0, 0))
                stable_item.compress_type = item.compress_type
                stable_item.external_attr = item.external_attr
                target_zip.writestr(stable_item, source_zip.read(item.filename), compress_type=item.compress_type)
    content = canonical.getvalue()
    if not content or len(content) > MAX_FILE_BYTES or not zipfile.is_zipfile(io.BytesIO(content)):
        raise OperationalFeishuReportError("operational_feishu_xlsx_invalid", 413)
    safe_region = re.sub(r"[^A-Za-z0-9_-]+", "_", region_code or "nacional").strip("_") or "regional"
    prefix = "taxa_coleta" if source_id == "collection_rate" else "falha_coleta_pdd"
    return f"{prefix}_{safe_region}_{operational_date.isoformat()}.xlsx", content


class _OutputTooLarge(Exception):
    pass


class _LimitedWriter:
    """Seekable file wrapper that enforces the attachment ceiling while writing."""
    def __init__(self, handle, limit: int, deadline: float | None = None):
        self._handle = handle
        self._limit = limit
        self._deadline = deadline
        self._max_extent = 0

    def write(self, value: bytes) -> int:
        if self._deadline is not None:
            _check_generation_deadline(self._deadline)
        end = self._handle.tell() + len(value)
        if max(self._max_extent, end) > self._limit:
            raise _OutputTooLarge
        written = self._handle.write(value)
        self._max_extent = max(self._max_extent, end)
        return written

    def __getattr__(self, name: str):
        return getattr(self._handle, name)


def _temp_artifact(root: Path, suffix: str) -> Path:
    handle = tempfile.NamedTemporaryFile(prefix=".operational-feishu-", suffix=suffix, dir=root, delete=False)
    path = Path(handle.name)
    handle.close()
    try:
        path.chmod(0o600)
    except OSError:
        path.unlink(missing_ok=True)
        raise
    return path


def _write_xlsx_stream(source_id: str, rows, operational_date: date, region_code: str | None, raw_path: Path, final_path: Path, deadline: float, *, summary_rows=None, period_label: str | None = None) -> str:
    from openpyxl import Workbook
    from openpyxl.writer.excel import ExcelWriter
    from openpyxl.worksheet import _writer as worksheet_writer

    source = _source(source_id)
    contract = get_contract(source_id, source["contract"])
    columns = [column.key for column in contract.columns]
    headers = [column.header for column in contract.columns]
    workbook = Workbook(write_only=True)
    workbook.properties.title = f"{source_id} — relatório oficial"
    workbook.properties.subject = f"Publication scoped report; region={region_code or 'ALL'}"
    stable = datetime.combine(operational_date, time.min)
    workbook.properties.created = stable
    workbook.properties.modified = stable
    active_sheet = None
    active_day = None
    if source["sheet_mode"] == "single":
        active_sheet = workbook.create_sheet(title=source["sheet_name"])
        active_sheet.append(headers)
    before_temps = set(worksheet_writer.ALL_TEMP_FILES)
    try:
        for row in rows:
            _check_generation_deadline(deadline)
            if source["sheet_mode"] == "daily":
                if row.get("data_date") is None:
                    continue
                row_day = _as_date(row["data_date"])
                if active_day != row_day:
                    active_day = row_day
                    active_sheet = workbook.create_sheet(title=row_day.isoformat()[:31])
                    active_sheet.append(headers)
            active_sheet.append([row.get(key) if row.get(key) is not None else "" for key in columns])
        if summary_rows is not None:
            summary_sheet = workbook.create_sheet(title="Resumo da tentativa")
            for summary_row in summary_rows():
                summary_sheet.append(list(summary_row))
        if source["sheet_mode"] == "daily" and active_sheet is None:
            raise OperationalFeishuReportError("operational_feishu_no_records", 404)
        with raw_path.open("w+b") as handle:
            archive = zipfile.ZipFile(_LimitedWriter(handle, MAX_FILE_BYTES, deadline), "w", zipfile.ZIP_DEFLATED, allowZip64=True)
            try:
                ExcelWriter(workbook, archive).save()
            except _OutputTooLarge:
                # openpyxl's save_workbook helper does not close its archive if
                # a write fails. Detach it so ZipFile.__del__ cannot seek on a
                # file handle that this context is about to close.
                archive.fp = None
                archive._fileRefCnt = 0
                raise
            except BaseException:
                try:
                    archive.close()
                except Exception:
                    archive.fp = None
                    archive._fileRefCnt = 0
                raise
            handle.flush()
            os.fsync(handle.fileno())
        with zipfile.ZipFile(raw_path, "r") as original:
            with final_path.open("w+b") as target_handle:
                normalized = zipfile.ZipFile(_LimitedWriter(target_handle, MAX_FILE_BYTES, deadline), "w", zipfile.ZIP_DEFLATED)
                try:
                    for item in original.infolist():
                        stable_item = zipfile.ZipInfo(item.filename, date_time=(1980, 1, 1, 0, 0, 0))
                        stable_item.compress_type = item.compress_type
                        stable_item.external_attr = item.external_attr
                        stable_item.internal_attr = item.internal_attr
                        stable_item.create_system = item.create_system
                        with original.open(item, "r") as source_handle, normalized.open(stable_item, "w", force_zip64=True) as target_entry:
                            while chunk := source_handle.read(1024 * 1024):
                                _check_generation_deadline(deadline)
                                target_entry.write(chunk)
                    normalized.close()
                except _OutputTooLarge:
                    normalized.fp = None
                    normalized._fileRefCnt = 0
                    raise
                except BaseException:
                    try:
                        normalized.close()
                    except Exception:
                        normalized.fp = None
                        normalized._fileRefCnt = 0
                    raise
                target_handle.flush()
                os.fsync(target_handle.fileno())
        final_path.chmod(0o600)
        if final_path.stat().st_size <= 0 or not zipfile.is_zipfile(final_path):
            raise OperationalFeishuReportError("operational_feishu_xlsx_invalid", 413)
    except _OutputTooLarge:
        for sheet in workbook.worksheets:
            try:
                sheet.close()
            except Exception:
                pass
        try:
            workbook.close()
        except Exception:
            pass
        raise OperationalFeishuReportError("operational_feishu_xlsx_invalid", 413) from None
    except BaseException:
        if "archive" in locals() and archive.fp is not None:
            try:
                archive.close()
            except Exception:
                archive.fp = None
                archive._fileRefCnt = 0
        for sheet in workbook.worksheets:
            try:
                sheet.close()
            except Exception:
                pass
        try:
            workbook.close()
        except Exception:
            pass
        raise
    finally:
        for temp_name in set(worksheet_writer.ALL_TEMP_FILES) - before_temps:
            try:
                Path(temp_name).unlink(missing_ok=True)
                worksheet_writer.ALL_TEMP_FILES.remove(temp_name)
            except (OSError, ValueError):
                pass
    safe_region = re.sub(r"[^A-Za-z0-9_-]+", "_", region_code or "nacional").strip("_") or "regional"
    prefix = "taxa_coleta" if source_id == "collection_rate" else "falha_coleta_pdd"
    suffix = (period_label or operational_date.isoformat()).replace("/", "-").replace(" ", "_")
    return f"{prefix}_{safe_region}_{suffix}.xlsx"


def _normalized_scope_text(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _selected_scope_value(filters: dict[str, list[str]], key: str, value: str) -> bool:
    selected = filters.get(key) or []
    return not selected or value in selected


def _collection_scope_row(row: dict[str, Any], scope: dict[str, Any], report_region: str | None) -> dict[str, Any] | None:
    filters = scope["filters"]
    resolution = _normalized_scope_text(row.get("effective_mapping_resolution") or row.get("mapping_resolution")).upper()
    base_code = _normalized_scope_text(row.get("effective_base_code") or row.get("resolved_base_code"))
    mapped_region = _normalized_scope_text(row.get("mapped_region"))
    approved = bool(base_code and mapped_region and resolution in {"NAME_APPROVED", "CODE_MATCH"})
    region = mapped_region if approved else "Sem regional"
    rm_area = _normalized_scope_text(row.get("mapped_rm_area")) if approved else "Sem região do RM"
    rm_area = rm_area or "Sem região do RM"
    rm = _normalized_scope_text(row.get("mapped_rm")) if approved else "Sem RM"
    rm = rm or "Sem RM"
    source_base_name = str(row.get("source_base_name") or "")
    base = (_normalized_scope_text(row.get("mapped_base_name")) or _normalized_scope_text(source_base_name)) if approved else source_base_name
    base = base or "Sem base"
    origin = _normalized_scope_text(row.get("order_origin")) or "Sem origem"
    product_type = _normalized_scope_text(row.get("product_type")) or "Sem tipo"
    if report_region and region != report_region:
        return None
    values = {"regions": region, "rms": rm, "rmAreas": rm_area, "bases": base, "origins": origin, "productTypes": product_type}
    if not all(_selected_scope_value(filters, key, value) for key, value in values.items()):
        return None
    if filters.get("rgms") and region not in set(filters.get("rgmRegions") or []):
        return None
    result = dict(row)
    result.update({"resolved_base_code": base_code if approved else None,
                   "mapping_resolution": resolution if approved else "UNRESOLVED",
                   "mapped_region": mapped_region if approved else None,
                   "mapped_base_name": row.get("mapped_base_name") if approved else None,
                   "mapped_rm": row.get("mapped_rm") if approved else None})
    return result


def _stream_collection_rate_scope_payload(
    engine: Engine, *, publication_id: str, report_region: str | None,
    scope: dict[str, Any], identity: dict[str, Any], expected_mapping_revision: str | None,
    staging_directory: Path, report_id: uuid.UUID,
) -> dict[str, Any]:
    root = _artifact_root(staging_directory)
    deadline = monotonic_clock.monotonic() + MAX_GENERATION_SECONDS
    date_from = date.fromisoformat(scope["dateFrom"])
    date_to = date.fromisoformat(scope["dateTo"])
    start_revision = read_revision(engine)
    if not expected_mapping_revision or start_revision["revision"] != expected_mapping_revision:
        raise OperationalFeishuReportError("operational_feishu_scope_changed", 409)
    if scope.get("mappingRevision") != expected_mapping_revision:
        raise OperationalFeishuReportError("operational_feishu_scope_changed", 409)
    start_manifest = _collection_rate_period_manifest(engine, scope["dateFrom"], scope["dateTo"])
    anchor_publication_id = str(uuid.UUID(publication_id))
    if not any(len(item.split(":")) == 3 and item.split(":")[1] == anchor_publication_id for item in start_manifest):
        raise OperationalFeishuReportError("operational_feishu_scope_changed", 409)
    map_enabled = _assert_map_available(engine)
    map_id = uuid.UUID(str(start_revision["mappingVersionId"])).bytes
    plan = _query_plan("collection_rate", identity, date_from=date_from, date_to=date_to,
                       contract_version=3, sql_dialect=engine.dialect.name,
                       include_unit_metadata=map_enabled)
    raw_path = _temp_artifact(root, ".xlsx.raw")
    normalized_path = _temp_artifact(root, ".xlsx.normalized")
    xlsx_path = _artifact_path(staging_directory, report_id, "xlsx")
    totals = _new_analysis_totals("collection_rate")
    contract_columns = [column.key for column in get_contract("collection_rate", 3).columns]
    digest = hashlib.sha256()
    digest.update(map_id)
    identity_digest = hashlib.sha256()
    content_digest = _ScopedContentDigest()
    count = 0
    active_publications: set[str] = set()
    filtered_regions: set[str] = set()

    def selected_rows(connection):
        nonlocal count
        result = connection.execution_options(yield_per=REPORT_FETCH_BATCH_SIZE).execute(text(f"""
            SELECT {plan.select_sql}
            {plan.from_sql}
            WHERE {plan.where_sql}
            ORDER BY fact.data_date DESC, fact.source_row_no ASC
        """), plan.parameters).mappings()
        try:
            for mapped in result:
                _check_generation_deadline(deadline)
                row = dict(mapped)
                if not _normalized_scope_text(row.get("publication_id")):
                    continue
                row["publication_id"] = str(uuid.UUID(bytes=bytes(row["publication_id"])))
                selected = _collection_scope_row(row, scope, report_region)
                if selected is None:
                    continue
                key = f"{selected['publication_id']}:{int(selected['source_row_no'])}"
                identity_digest.update(key.encode("utf-8"))
                identity_digest.update(b"\n")
                content_digest.add(selected, contract_columns)
                _digest_add(digest, selected, contract_columns)
                active_publications.add(selected["publication_id"])
                filtered_regions.add(_normalized_scope_text(selected.get("mapped_region")) or "Sem regional")
                _accumulate_analysis_row("collection_rate", totals, selected, None)
                count += 1
                yield selected
        finally:
            result.close()

    try:
        # Verify the selected anchor publication before reading the active daily
        # partitions. The period query itself follows the Data Foundation's
        # per-date current-publication and official-map rules.
        anchor = _publication(engine, "collection_rate", publication_id)
        regions = set(_regions(engine, uuid.UUID(start_revision["mappingVersionId"]).bytes))
        if report_region is not None and report_region not in regions:
            raise OperationalFeishuReportError("operational_feishu_region_invalid", 422)
        summary_result: dict[str, Any] = {}
        report_publication = dict(anchor)
        report_publication["map_version_id"] = map_id
        report_publication["scope_publication_ids"] = []
        report_publication["scope_map_revision"] = str(start_revision["revision"])

        def summary_sheet_rows():
            message, result, summary = _finish_analysis("collection_rate", report_publication, report_region, date_to, totals)
            summary_result.update({"message": message, "analysisResult": result, "totals": summary})
            rate = Decimal(totals["attempted"]) / Decimal(totals["planned"]) if totals["planned"] else None
            rows = [
                ["Indicador", "Valor"],
                ["Período analisado", f"{date_from.isoformat()} a {date_to.isoformat()}"],
                ["Escopo", scope.get("scopeLabel") or "Filtros do painel"],
                ["Taxa de Tentativa consolidada", _percent(rate) if rate is not None else "—"],
                ["Numerador: coletados + tentativas", totals["attempted"]],
                ["Denominador: quantidade a coletar", totals["planned"]],
                ["Pedidos", totals["totalOrders"]],
                ["Volume sem mapeamento (incluído no consolidado)", totals["unmappedPlanned"]],
                [], ["Regional (inclui sem regional)", "Numerador", "Denominador", "Taxa de Tentativa"],
            ]
            for region, aggregate in sorted(totals["regions"].items()):
                numerator = int(aggregate["attempted"])
                denominator = int(aggregate["planned"])
                rows.append([region, numerator, denominator, _percent(Decimal(numerator) / Decimal(denominator)) if denominator else "—"])
            rows.extend([[], ["Base", "Código", "Regional", "RM", "Numerador", "Denominador", "Taxa de Tentativa"]])
            for base in sorted(totals["groups"].values(), key=lambda item: (str(item.get("region") or ""), str(item.get("name") or ""), str(item.get("code") or ""))):
                denominator = int(base.get("planned") or 0)
                numerator = int(base.get("attempted") or 0)
                rows.append([base.get("name") or base.get("code"), base.get("code"), base.get("region"), base.get("responsible") or "Sem RM", numerator, denominator, _percent(Decimal(numerator) / Decimal(denominator)) if denominator else "—"])
            return rows

        with engine.connect() as connection:
            with connection.begin():
                file_name = _write_xlsx_stream(
                    "collection_rate", selected_rows(connection), date_to, report_region,
                    raw_path, normalized_path, deadline, summary_rows=summary_sheet_rows,
                    period_label=f"{date_from.isoformat()}_{date_to.isoformat()}",
                )
                if count == 0:
                    raise OperationalFeishuReportError("operational_feishu_no_records", 404)
                actual_identity_digest = identity_digest.hexdigest()
                if not hmac.compare_digest(actual_identity_digest, str(scope.get("expectedRowIdentitySha256") or "")):
                    raise OperationalFeishuReportError("operational_feishu_scope_changed", 409)
                if not hmac.compare_digest(content_digest.hexdigest(), str(scope.get("expectedRowContentSha256") or "")):
                    raise OperationalFeishuReportError("operational_feishu_scope_changed", 409)
                if sorted(active_publications) != sorted(scope.get("publicationIds") or []):
                    raise OperationalFeishuReportError("operational_feishu_scope_changed", 409)
                if scope.get("expectedRowCount") != count:
                    raise OperationalFeishuReportError("operational_feishu_scope_changed", 409)
                dataset_digest = digest.digest()
        end_revision = read_revision(engine)
        end_manifest = _collection_rate_period_manifest(engine, scope["dateFrom"], scope["dateTo"])
        if end_revision["revision"] != start_revision["revision"] or end_manifest != start_manifest:
            raise OperationalFeishuReportError("operational_feishu_scope_changed", 409)
        os.replace(normalized_path, xlsx_path)
        report_publication["scope_publication_ids"] = sorted(active_publications)
        message = summary_result["message"]
        analysis_result = summary_result["analysisResult"]
        analysis_result["periodLabel"] = f"{date_from.isoformat()} a {date_to.isoformat()}"
        analysis_result["scopeLabel"] = scope.get("scopeLabel") or "Filtros do painel"
        summary = summary_result["totals"]
        png = _png("collection_rate", report_publication, report_region, summary, analysis_result)
        _check_generation_deadline(deadline)
        _store(staging_directory, report_id, {"png": png})
        return {"publication": report_publication, "analysis": message,
                "analysisResult": analysis_result, "totals": summary, "digest": dataset_digest,
                "recordCount": count, "fileName": file_name, "xlsxBytes": xlsx_path.stat().st_size,
                "pngBytes": len(png), "operationalDate": date_to,
                "activePublications": sorted(active_publications),
                "filteredRegions": sorted(filtered_regions),
                "periodPublicationManifest": start_manifest,
                "mapRevision": str(start_revision["revision"]),
                "artifacts": {"xlsxSha256": _file_digest(xlsx_path), "pngSha256": hashlib.sha256(png).hexdigest()}}
    except Exception:
        for path in (raw_path, normalized_path, xlsx_path):
            try:
                path.unlink(missing_ok=True)
            except OSError:
                pass
        raise
    finally:
        for path in (raw_path, normalized_path):
            try:
                path.unlink(missing_ok=True)
            except OSError:
                pass


def _collection_rate_period_manifest(engine: Engine, date_from: str, date_to: str) -> list[str]:
    with engine.connect() as connection:
        rows = connection.execute(text("""
        SELECT ranked.data_date, ranked.publication_id, ranked.version_no
        FROM (
            SELECT metadata.data_date, publication.publication_id, publication.version_no,
                   metadata.row_count,
                   ROW_NUMBER() OVER (PARTITION BY metadata.data_date ORDER BY publication.version_no DESC) AS publication_rank
            FROM source_publication_dates AS metadata
            JOIN source_publications AS publication
              ON publication.source_id=metadata.source_id AND publication.publication_id=metadata.publication_id
            WHERE metadata.source_id='collection_rate' AND publication.contract_version=3
              AND publication.publication_state='PUBLISHED'
              AND metadata.data_date>=:date_from AND metadata.data_date<=:date_to
        ) AS ranked
        WHERE ranked.publication_rank=1 AND ranked.row_count>0
        ORDER BY ranked.data_date
        """), {"date_from": date_from, "date_to": date_to}).mappings().all()
    return [f"{_as_date(row['data_date']).isoformat()}:{uuid.UUID(bytes=bytes(row['publication_id']))}:{int(row['version_no'])}" for row in rows]


def _report_scope(engine: Engine, report: Any) -> dict[str, Any] | None:
    context = report.get("delivery_context") if hasattr(report, "get") else None
    if isinstance(context, str):
        try:
            context = json.loads(context)
        except (TypeError, ValueError):
            return None
    if not isinstance(context, dict):
        return None
    scope = context.get("dashboardScope")
    if not isinstance(scope, dict):
        return None
    current_revision = read_revision(engine)
    period_manifest = _collection_rate_period_manifest(engine, str(scope.get("dateFrom")), str(scope.get("dateTo")))
    try:
        anchor_publication_id = str(uuid.UUID(bytes=bytes(report.get("publication_id") or b"")))
    except (TypeError, ValueError):
        raise OperationalFeishuReportError("operational_feishu_scope_changed", 409) from None
    anchor_versions = [item.split(":") for item in period_manifest if len(item.split(":")) == 3 and item.split(":")[1] == anchor_publication_id]
    if (str(report.get("source_id")) != "collection_rate"
            or current_revision.get("revision") != scope.get("mapRevision")
            or bytes(report.get("map_version_id") or b"") != uuid.UUID(str(current_revision["mappingVersionId"])).bytes
            or period_manifest != scope.get("periodPublicationManifest")
            or not anchor_versions
            or int(report.get("contract_version") or 0) != 3
            or any(int(report.get("publication_version") or 0) != int(item[2]) for item in anchor_versions)):
        raise OperationalFeishuReportError("operational_feishu_scope_changed", 409)
    return scope


def _file_digest(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _stream_payload(engine: Engine, *, source_id: str, publication_id: str, region_code: str | None,
                    staging_directory: Path, report_id: uuid.UUID) -> dict[str, Any]:
    root = _artifact_root(staging_directory)
    deadline = monotonic_clock.monotonic() + MAX_GENERATION_SECONDS
    raw_path: Path | None = None
    normalized_path: Path | None = None
    xlsx_path = _artifact_path(staging_directory, report_id, "xlsx")
    png_path = _artifact_path(staging_directory, report_id, "png")
    try:
        raw_path = _temp_artifact(root, ".xlsx.raw")
        normalized_path = _temp_artifact(root, ".xlsx.normalized")
        with engine.connect() as connection:
            with connection.begin():
                publication = _publication(engine, source_id, publication_id, connection)
                regions = _regions(engine, bytes(publication["map_version_id"]), connection)
                if region_code is not None and region_code not in regions:
                    raise OperationalFeishuReportError("operational_feishu_region_invalid", 422)
                day_value = connection.execute(text("""
                    SELECT MAX(data_date) FROM source_publication_dates
                    WHERE source_id=:source_id AND publication_id=:publication_id AND row_count>0
                """), {"source_id": source_id, "publication_id": bytes(publication["publication_id"])}).scalar_one_or_none()
                if day_value is None:
                    raise OperationalFeishuReportError("operational_feishu_operational_date_unavailable", 409)
                operational_date = _as_date(day_value)
                contract = get_contract(source_id, _source(source_id)["contract"])
                columns = [column.key for column in contract.columns]
                digest = hashlib.sha256()
                digest.update(bytes(publication["map_version_id"]))
                count = 0
                totals = _new_analysis_totals(source_id)
                query, params = _rows_query(engine, source_id, publication, region_code)

                def rows_and_account():
                    nonlocal count
                    for mapped_row in _iter_rows(connection, query, params):
                        _check_generation_deadline(deadline)
                        row = dict(mapped_row)
                        _digest_add(digest, row, columns)
                        _accumulate_analysis_row(source_id, totals, row, operational_date)
                        count += 1
                        yield row

                file_name = _write_xlsx_stream(source_id, rows_and_account(), operational_date, region_code, raw_path, normalized_path, deadline)
                if count == 0:
                    raise OperationalFeishuReportError("operational_feishu_no_records", 404)
                message, analysis_result, summary = _finish_analysis(source_id, publication, region_code, operational_date, totals)
                dataset_digest = digest.digest()
        os.replace(normalized_path, xlsx_path)
        png = _png(source_id, publication, region_code, summary, analysis_result)
        _check_generation_deadline(deadline)
        _store(staging_directory, report_id, {"png": png})
        return {
            "publication": publication, "analysis": message, "analysisResult": analysis_result,
            "totals": summary, "digest": dataset_digest, "recordCount": count,
            "fileName": file_name, "xlsxBytes": xlsx_path.stat().st_size, "pngBytes": len(png),
            "operationalDate": operational_date, "artifacts": {"xlsxSha256": _file_digest(xlsx_path), "pngSha256": hashlib.sha256(png).hexdigest()},
        }
    except Exception:
        for path in (raw_path, normalized_path, xlsx_path, png_path):
            try:
                path.unlink(missing_ok=True)
            except OSError:
                pass
        raise
    finally:
        for path in (raw_path, normalized_path):
            if path is not None:
                try:
                    path.unlink(missing_ok=True)
                except OSError:
                    pass


def _png(source_id: str, publication: dict[str, Any], region_code: str | None, totals: dict[str, int], result: dict[str, Any]) -> bytes:
    from PIL import Image, ImageDraw, ImageFont

    width, height = 1400, 760
    image = Image.new("RGB", (width, height), "#f4f6f9")
    draw = ImageDraw.Draw(image)
    def font(size: int, bold: bool = False):
        candidates = (
            ["C:/Windows/Fonts/msyhbd.ttc", "C:/Windows/Fonts/arialbd.ttf"] if bold
            else ["C:/Windows/Fonts/msyh.ttc", "C:/Windows/Fonts/arial.ttf"]
        ) + [
            "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        ]
        for candidate in candidates:
            try:
                return ImageFont.truetype(candidate, size=size)
            except OSError:
                continue
        return ImageFont.load_default()

    title, body, small = font(34, True), font(23), font(18)
    draw.rounded_rectangle((24, 20, width - 24, height - 20), radius=20, fill="white", outline="#e5e8ef", width=2)
    title_text = "TAXA DE COLETA V3" if source_id == "collection_rate" else "FALHA DE COLETA PDD"
    if source_id == "collection_rate" and result.get("periodLabel"):
        title_text = "TAXA DE TENTATIVA V3 | 尝试揽收率"
    draw.text((58, 48), "J&T EXPRESS  |  RELATÓRIO OPERACIONAL", fill="#d71920", font=small)
    draw.text((58, 84), title_text, fill="#17191f", font=title)
    if source_id == "collection_rate" and result.get("periodLabel"):
        draw.text((58, 140), f"Período: {result['periodLabel']}    Regional: {region_code or 'Todas as regionais'}", fill="#4d5360", font=body)
        scope_text = str(result.get("scopeLabel") or "Todos")
        scope_lines = [scope_text[index:index + 110] for index in range(0, min(len(scope_text), 220), 110)] or ["Todos"]
        if len(scope_text) > 220:
            scope_lines[-1] = scope_lines[-1][:107] + "…"
        for index, scope_line in enumerate(scope_lines):
            draw.text((58, 178 + index * 22), f"Filtros: {scope_line}" if index == 0 else f"         {scope_line}", fill="#4d5360", font=small)
        draw.text((58, 222), f"Publicações diárias: {len(publication.get('scope_publication_ids') or [])}    Contrato V{publication['contract_version']}", fill="#4d5360", font=small)
        cards_top, cards_bottom, cards_label, cards_value, rows_heading, row_top = 260, 375, 278, 316, 416, 458
    else:
        draw.text((58, 140), f"Data operacional: {result['operationalDate']}    Regional: {region_code or 'Todas as regionais'}", fill="#4d5360", font=body)
        draw.text((58, 180), f"Publicação v{publication['version_no']}    Contrato V{publication['contract_version']}", fill="#4d5360", font=small)
        cards_top, cards_bottom, cards_label, cards_value, rows_heading, row_top = 245, 360, 264, 300, 405, 450
    if source_id == "collection_rate":
        denominator, numerator = totals["rateDenominator"], totals["rateNumerator"]
        rate = Decimal(numerator) / Decimal(denominator) if denominator else None
        rate_text = _percent(rate) if rate is not None else "—"
        cards = [("VOLUME A COLETAR", f"{denominator:,}".replace(",", ".")), ("COLETA + TENTATIVAS", f"{numerator:,}".replace(",", ".")), ("TAXA DE TENTATIVA", rate_text), ("PEDIDOS", f"{totals['totalOrders']:,}".replace(",", "."))]
    else:
        den, num = totals["rateDenominator"], totals["rateNumerator"]
        rate = Decimal(num) / Decimal(den) if den else None
        cards = [("TOTAL DE PEDIDOS", f"{totals['totalOrders']:,}".replace(",", ".")), ("NÃO BIPADOS NA COLETA", f"{num:,}".replace(",", ".")), ("TAXA DE FALHA", _percent(rate) if rate is not None else "—")]
    x = 58
    card_width = 300
    for label, value in cards:
        draw.rounded_rectangle((x, cards_top, x + card_width, cards_bottom), radius=14, fill="#fff2f3", outline="#f1d5d7")
        draw.text((x + 18, cards_label), label, fill="#6b7079", font=small)
        draw.text((x + 18, cards_value), value, fill="#d71920", font=title)
        x += card_width + 24
    draw.text((58, rows_heading), "Principais unidades mapeadas", fill="#17191f", font=body)
    base_rows = []
    results = result.get("result", {})
    if source_id == "collection_rate":
        for region in results.get("regions", []):
            for base in region.get("bases", []):
                base_rows.append((base.get("baseName") or base.get("baseCode"), _percent(base.get("rate")), base.get("responsible")))
    else:
        for base in results.get("priorityBases", []):
            base_rows.append((base.get("baseName") or base.get("baseCode"), _percent(base.get("failureRate")), base.get("responsible")))
    for index, (name, ratio, owner) in enumerate(base_rows[:10]):
        draw.text((70, row_top + index * 25), f"• {name} — {ratio or '—'} | {owner or 'Sem responsável'}", fill="#333842", font=small)
    if not base_rows:
        draw.text((70, row_top + 2), "Sem unidades prioritárias nesta publicação.", fill="#555b64", font=small)
    draw.text((58, height - 62), "Prévia; o Excel contém todos os registros autorizados do recorte.", fill="#777d87", font=small)
    output = io.BytesIO()
    image.save(output, format="PNG", optimize=True)
    data = output.getvalue()
    if not data or len(data) > MAX_FILE_BYTES:
        raise OperationalFeishuReportError("operational_feishu_png_invalid", 413)
    return data


def _payload(engine: Engine, source_id: str, publication_id: str, region_code: str | None, connection: Connection | None = None, include_artifacts: bool = True) -> dict[str, Any]:
    publication = _publication(engine, source_id, publication_id, connection)
    regions = _regions(engine, bytes(publication["map_version_id"]))
    if region_code is not None and region_code not in regions:
        raise OperationalFeishuReportError("operational_feishu_region_invalid", 422)
    rows = _fetch_rows(engine, source_id, publication, region_code, connection)
    if not rows:
        raise OperationalFeishuReportError("operational_feishu_no_records", 404)
    contract = get_contract(source_id, _source(source_id)["contract"])
    columns = [column.key for column in contract.columns]
    digest = _digest(rows, columns, bytes(publication["map_version_id"]))
    analysis = None
    result = None
    totals = None
    file_name = ""
    xlsx = png = None
    operational_date: date | None = None
    if include_artifacts:
        if connection is None:
            with engine.connect() as active:
                day_value = active.execute(text("""
                    SELECT MAX(data_date) FROM source_publication_dates
                    WHERE source_id=:source_id AND publication_id=:publication_id AND row_count>0
                """), {"source_id": source_id, "publication_id": bytes(publication["publication_id"])}).scalar_one_or_none()
        else:
            day_value = connection.execute(text("""
                SELECT MAX(data_date) FROM source_publication_dates
                WHERE source_id=:source_id AND publication_id=:publication_id AND row_count>0
            """), {"source_id": source_id, "publication_id": bytes(publication["publication_id"])}).scalar_one_or_none()
        if day_value is None:
            raise OperationalFeishuReportError("operational_feishu_operational_date_unavailable", 409)
        operational_date = _as_date(day_value)
        analysis, result, totals = _analysis(source_id, publication, rows, region_code, operational_date)
        file_name, xlsx = _xlsx(source_id, rows, operational_date, region_code)
        png = _png(source_id, publication, region_code, totals, result)
    return {"publication": publication, "rows": rows, "regionCode": region_code, "analysis": analysis,
            "analysisResult": result, "totals": totals, "digest": digest, "fileName": file_name,
            "xlsx": xlsx, "png": png, "operationalDate": operational_date, "columns": columns}


def prepare_report(engine: Engine, *, source_id: str, publication_id: str, region_code: str | None,
                   recipient_key: str, actor_user_id: int, staging_directory: Path,
                   settings: Settings | None = None, dashboard_scope: dict[str, Any] | None = None,
                   identity: dict[str, Any] | None = None, expected_mapping_revision: str | None = None) -> dict[str, object]:
    config = settings or get_settings()
    _source(source_id)
    _require_feature_enabled(source_id, config)
    try:
        recipient = resolve_recipient(recipient_key, region_code, config)
    except ValueError as error:
        code = "operational_feishu_recipient_invalid" if str(error) == "feishu_report_recipient_invalid" else str(error)
        raise OperationalFeishuReportError(code, 403) from None
    report_id = uuid.uuid4()
    if source_id == "collection_rate" and dashboard_scope is not None:
        if identity is None:
            raise OperationalFeishuReportError("matrix_scope_required", 403)
        payload = _stream_collection_rate_scope_payload(
            engine, publication_id=publication_id, region_code=region_code,
            scope=dashboard_scope, identity=identity, expected_mapping_revision=expected_mapping_revision,
            staging_directory=staging_directory, report_id=report_id,
        )
        dashboard_scope = {**dashboard_scope,
                           "publicationIds": payload["activePublications"],
                           "periodPublicationManifest": payload["periodPublicationManifest"],
                           "mapRevision": payload["mapRevision"],
                           "expectedRowIdentitySha256": str(dashboard_scope["expectedRowIdentitySha256"]),
                           "expectedRowCount": int(payload["recordCount"]),
                           "scopeLabel": str(dashboard_scope.get("scopeLabel") or "Filtros do painel")}
    else:
        payload = _stream_payload(engine, source_id=source_id, publication_id=publication_id,
                                  region_code=region_code, staging_directory=staging_directory, report_id=report_id)
    now = _now()
    hashes = payload["artifacts"]
    previously_sent = False
    try:
        with engine.begin() as connection:
            prior_delivery = connection.execute(text(f"""
                SELECT report_id, report_state FROM {TABLE}
                WHERE source_id=:source_id AND publication_id=:publication_id AND region_key=:region_key
                  AND recipient_key=:recipient_key AND rules_version=:rules_version
                  AND report_state IN ('SENDING','PARTIAL','PENDING','SENT')
                ORDER BY CASE WHEN report_state IN ('SENDING','PARTIAL','PENDING') THEN 0 ELSE 1 END
                LIMIT 1
            """), {"source_id": source_id, "publication_id": bytes(payload["publication"]["publication_id"]),
                  "region_key": region_code or "ALL", "recipient_key": recipient_key, "rules_version": RULES_VERSION}).mappings().first()
            if prior_delivery and prior_delivery["report_state"] in {"SENDING", "PARTIAL", "PENDING"}:
                raise OperationalFeishuReportError("operational_feishu_delivery_unresolved", 409)
            previously_sent = bool(prior_delivery and prior_delivery["report_state"] == "SENT")
            connection.execute(text(f"""
                INSERT INTO {TABLE}
                  (report_id,source_id,publication_id,publication_version,contract_version,map_version_id,
                   region_key,recipient_key,recipient_email,recipient_type,recipient_identifier,recipient_name,
                   actor_user_id,analysis_text,record_count,
                   dataset_sha256,rules_version,report_state,delivery_context,created_at,updated_at)
                VALUES (:report_id,:source_id,:publication_id,:publication_version,:contract_version,:map_version_id,
                   :region_key,:recipient_key,:recipient_email,:recipient_type,:recipient_identifier,:recipient_name,
                   :actor_user_id,:analysis_text,:record_count,
                   :dataset_sha256,:rules_version,'PREPARED',:delivery_context,:created_at,:updated_at)
            """), {"report_id": report_id.bytes, "source_id": source_id,
                  "publication_id": bytes(payload["publication"]["publication_id"]),
                  "publication_version": int(payload["publication"]["version_no"]),
                  "contract_version": int(payload["publication"]["contract_version"]),
                  "map_version_id": bytes(payload["publication"]["map_version_id"]),
                  "region_key": region_code or "ALL", "recipient_key": recipient_key,
                  "recipient_email": recipient.identifier if recipient.kind == "email" else None,
                  "recipient_type": recipient.kind, "recipient_identifier": recipient.identifier,
                  "recipient_name": recipient.audit_name or recipient.label, "actor_user_id": actor_user_id,
                  "analysis_text": payload["analysis"], "record_count": payload["recordCount"],
                  "dataset_sha256": payload["digest"], "rules_version": RULES_VERSION,
                  "delivery_context": json.dumps({"artifacts": hashes, "fileName": payload["fileName"], "stages": {}, **({"dashboardScope": dashboard_scope} if dashboard_scope is not None else {})}, separators=(",", ":")),
                  "created_at": now, "updated_at": now})
    except Exception:
        for ext in ("xlsx", "png"):
            try:
                _artifact_path(staging_directory, report_id, ext).unlink(missing_ok=True)
            except OSError:
                pass
        raise
    pub = payload["publication"]
    return {"reportId": str(report_id), "sourceId": source_id,
            "publicationId": str(uuid.UUID(bytes=bytes(pub["publication_id"]))),
            "publicationVersion": int(pub["version_no"]), "contractVersion": int(pub["contract_version"]),
            "mapVersionId": str(uuid.UUID(bytes=bytes(pub["map_version_id"]))),
            "publishedAt": pub.get("published_at").isoformat(sep=" ") if isinstance(pub.get("published_at"), datetime) else str(pub.get("published_at") or ""),
            "operationalDate": payload["operationalDate"].isoformat(), "regionCode": region_code,
            "regionLabel": region_code or "Todas as regionais", "recipientKey": recipient.key,
            "recipientLabel": recipient.label, "recipientType": recipient.kind, "analysis": payload["analysis"],
            "recordCount": payload["recordCount"], "excelFileName": payload["fileName"],
            "periodStart": dashboard_scope.get("dateFrom") if dashboard_scope is not None else payload["operationalDate"].isoformat(),
            "periodEnd": dashboard_scope.get("dateTo") if dashboard_scope is not None else payload["operationalDate"].isoformat(),
            "scopeLabel": dashboard_scope.get("scopeLabel") if dashboard_scope is not None else region_code or "Publicação e regional selecionadas",
            "excelBytes": payload["xlsxBytes"], "pngBytes": payload["pngBytes"], "state": "PREPARED",
            "sendEnabled": True, "previouslySent": previously_sent}


def confirm_report(engine: Engine, *, report_id: str, actor_user_id: int, analysis_text: str, source_id: str,
                   settings: Settings | None = None) -> dict[str, object]:
    config = settings or get_settings()
    _source(source_id)
    _require_feature_enabled(source_id, config)
    if not analysis_text.strip():
        raise OperationalFeishuReportError("operational_feishu_text_required", 422)
    if len(analysis_text) > 20000:
        raise OperationalFeishuReportError("operational_feishu_text_too_long", 422)
    key, now = _uuid(report_id), _now()
    with engine.begin() as connection:
        lock = " FOR UPDATE" if engine.dialect.name == "mysql" else ""
        row = connection.execute(text(f"SELECT report_id,source_id,publication_id,publication_version,contract_version,map_version_id,region_key,recipient_key,recipient_email,recipient_type,recipient_identifier,recipient_name,report_state,delivery_context FROM {TABLE} WHERE report_id=:id AND actor_user_id=:actor" + lock), {"id": key, "actor": actor_user_id}).mappings().first()
        if row is None:
            raise OperationalFeishuReportError("operational_feishu_report_not_found", 404)
        if str(row["source_id"]) != source_id:
            raise OperationalFeishuReportError("operational_feishu_report_not_found", 404)
        if row["report_state"] not in {"PREPARED", "CONFIRMED"}:
            raise OperationalFeishuReportError("operational_feishu_confirmation_invalid", 409)
        region_code = None if row["region_key"] == "ALL" else str(row["region_key"])
        try:
            recipient_from_audit(
                str(row["recipient_key"]), region_code,
                row.get("recipient_type"), row.get("recipient_identifier"), row.get("recipient_name"),
                row.get("recipient_email"), config,
            )
        except ValueError as error:
            raise OperationalFeishuReportError(str(error), 409) from None
        if _report_scope(engine, row) is None:
            publication = _publication(engine, str(row["source_id"]), str(uuid.UUID(bytes=bytes(row["publication_id"]))), connection)
            _assert_report_publication_matches(row, publication)
        connection.execute(text(f"UPDATE {TABLE} SET report_state='CONFIRMED',analysis_text=:analysis,confirmed_at=:now,updated_at=:now WHERE report_id=:id AND actor_user_id=:actor"), {"analysis": analysis_text, "now": now, "id": key, "actor": actor_user_id})
    return {"reportId": report_id, "state": "CONFIRMED", "messageSent": False}


def _save_context(engine: Engine, report_id: bytes, state: str, context: dict[str, Any]) -> None:
    now = _now()
    with engine.begin() as connection:
        connection.execute(text(f"UPDATE {TABLE} SET report_state=:state,delivery_context=:context,updated_at=:now,sent_at=CASE WHEN :state='SENT' THEN :now ELSE sent_at END WHERE report_id=:id"), {"state": state, "context": json.dumps(context, separators=(",", ":")), "now": now, "id": report_id})


def send_report(engine: Engine, *, source_id: str, report_id: str, actor_user_id: int, staging_directory: Path,
                settings: Settings | None = None) -> dict[str, object]:
    config = settings or get_settings()
    requested_source_id = source_id
    _source(requested_source_id)
    _require_feature_enabled(requested_source_id, config)
    key, now = _uuid(report_id), _now()
    with engine.begin() as connection:
        lock = " FOR UPDATE" if engine.dialect.name == "mysql" else ""
        row = connection.execute(text(f"SELECT * FROM {TABLE} WHERE report_id=:id AND actor_user_id=:actor" + lock), {"id": key, "actor": actor_user_id}).mappings().first()
        if row is None:
            raise OperationalFeishuReportError("operational_feishu_report_not_found", 404)
        source_id = str(row["source_id"])
        if source_id != requested_source_id:
            raise OperationalFeishuReportError("operational_feishu_report_not_found", 404)
        report_region = None if row["region_key"] == "ALL" else str(row["region_key"])
        try:
            recipient = recipient_from_audit(
                str(row["recipient_key"]), report_region,
                row.get("recipient_type"), row.get("recipient_identifier"), row.get("recipient_name"),
                row.get("recipient_email"), config,
            )
        except ValueError as error:
            raise OperationalFeishuReportError(str(error), 409) from None
        if not _feature_enabled(source_id, config):
            raise OperationalFeishuReportError("operational_feishu_send_disabled", 503)
        if not config.feishu_oauth_app_id or not config.feishu_oauth_app_secret:
            raise OperationalFeishuReportError("feishu_application_not_configured", 503)
        if row["report_state"] != "CONFIRMED":
            raise OperationalFeishuReportError("operational_feishu_report_not_confirmed", 409)
        publication_id = str(uuid.UUID(bytes=bytes(row["publication_id"])))
        if engine.dialect.name == "mysql":
            # One publication row serializes different confirmed report IDs for
            # the same dataset, closing the race between duplicate checks.
            connection.execute(text("""
                SELECT publication_id FROM source_publications
                WHERE source_id=:source_id AND publication_id=:publication_id
                FOR UPDATE
            """), {"source_id": source_id, "publication_id": bytes(row["publication_id"])}).first()
        scope = _report_scope(engine, row)
        if scope is None:
            publication = _publication(engine, source_id, publication_id, connection)
            _assert_report_publication_matches(row, publication)
        duplicate = connection.execute(text(f"SELECT report_id FROM {TABLE} WHERE source_id=:source AND publication_id=:pub AND region_key=:region AND recipient_key=:recipient AND rules_version=:rules AND report_id<>:id AND report_state IN ('SENDING','PARTIAL','PENDING') LIMIT 1" + lock), {"source": source_id, "pub": bytes(row["publication_id"]), "region": row["region_key"], "recipient": row["recipient_key"], "rules": row["rules_version"], "id": key}).first()
        if duplicate:
            raise OperationalFeishuReportError("operational_feishu_delivery_unresolved", 409)
        region = None if row["region_key"] == "ALL" else str(row["region_key"])
        raw_context = row["delivery_context"]
        context = raw_context if isinstance(raw_context, dict) else json.loads(raw_context or "{}")
        if scope is None:
            current_digest, current_count = _digest_rows(connection, engine, source_id, publication, region)
            if current_count == 0:
                raise OperationalFeishuReportError("operational_feishu_no_records", 404)
            if not hmac.compare_digest(current_digest, bytes(row["dataset_sha256"])):
                raise OperationalFeishuReportError("operational_feishu_data_changed", 409)
        file_name = str(context.get("fileName") or "")
        if not file_name.endswith(".xlsx") or Path(file_name).name != file_name:
            raise OperationalFeishuReportError("operational_feishu_artifacts_changed", 409)
        stages = context.setdefault("stages", {})
        hashes = context.get("artifacts", {})
        report_uuid = uuid.UUID(bytes=key)
        for ext in ("xlsx", "png"):
            with _open_verified_artifact(staging_directory, report_uuid, ext, hashes.get(f"{ext}Sha256")):
                pass
        message_content = _feishu_interactive_content(str(row["analysis_text"]))
        connection.execute(text(f"UPDATE {TABLE} SET report_state='SENDING',updated_at=:now WHERE report_id=:id"), {"now": now, "id": key})
    try:
        with httpx.Client(timeout=HTTP_TIMEOUT) as client:
            token_result = _feishu_payload(client.post(TOKEN_URL, json={"app_id": config.feishu_oauth_app_id, "app_secret": config.feishu_oauth_app_secret}))
            token = token_result.get("tenant_access_token")
            if not isinstance(token, str) or not token:
                raise RuntimeError("feishu_token_unavailable")
            headers = {"Authorization": f"Bearer {token}"}
            if stages.get("text", {}).get("status") != "SENT":
                result = _feishu_payload(client.post(MESSAGE_URL, params={"receive_id_type": recipient.kind}, headers=headers, json={"receive_id": recipient.identifier, "msg_type": "interactive", "content": message_content}))
                stages["text"] = {"status": "SENT", "messageId": result["data"]["message_id"]}
                _save_context(engine, key, "SENDING", context)
            xlsx_stage = stages.setdefault("xlsx", {})
            if xlsx_stage.get("status") != "SENT":
                if not xlsx_stage.get("fileKey"):
                    with _open_verified_artifact(staging_directory, report_uuid, "xlsx", hashes.get("xlsxSha256")) as xlsx_file:
                        result = _feishu_payload(client.post(FILES_URL, headers=headers, data={"file_type": "stream", "file_name": file_name}, files={"file": (file_name, xlsx_file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}))
                    xlsx_stage["fileKey"] = result["data"]["file_key"]
                    _save_context(engine, key, "SENDING", context)
                result = _feishu_payload(client.post(MESSAGE_URL, params={"receive_id_type": recipient.kind}, headers=headers, json={"receive_id": recipient.identifier, "msg_type": "file", "content": json.dumps({"file_key": xlsx_stage["fileKey"]})}))
                xlsx_stage.update({"status": "SENT", "messageId": result["data"]["message_id"]})
                _save_context(engine, key, "SENDING", context)
            png_stage = stages.setdefault("png", {})
            if png_stage.get("status") != "SENT":
                if not png_stage.get("imageKey"):
                    with _open_verified_artifact(staging_directory, report_uuid, "png", hashes.get("pngSha256")) as png_file:
                        result = _feishu_payload(client.post(IMAGES_URL, headers=headers, data={"image_type": "message"}, files={"image": ("operational_preview.png", png_file, "image/png")}))
                    png_stage["imageKey"] = result["data"]["image_key"]
                    _save_context(engine, key, "SENDING", context)
                result = _feishu_payload(client.post(MESSAGE_URL, params={"receive_id_type": recipient.kind}, headers=headers, json={"receive_id": recipient.identifier, "msg_type": "image", "content": json.dumps({"image_key": png_stage["imageKey"]})}))
                png_stage.update({"status": "SENT", "messageId": result["data"]["message_id"]})
                _save_context(engine, key, "SENDING", context)
    except _FeishuRejected as error:
        _save_context(engine, key, "PARTIAL", context)
        raise OperationalFeishuReportError("operational_feishu_partial_delivery", 502) from error
    except Exception as error:
        _save_context(engine, key, "PENDING", context)
        logger.warning("Operational Feishu delivery needs reconciliation (%s)", type(error).__name__)
        raise OperationalFeishuReportError("operational_feishu_delivery_pending", 502) from None
    _save_context(engine, key, "SENT", context)
    for ext in ("xlsx", "png"):
        try:
            _artifact_path(staging_directory, uuid.UUID(bytes=key), ext).unlink(missing_ok=True)
        except OSError:
            logger.warning("Unable to remove successfully delivered operational report artifact")
    return {"reportId": report_id, "state": "SENT", "recipient": recipient.label, "messageIds": {name: stage.get("messageId") for name, stage in stages.items()}}
