import hashlib
from datetime import date, datetime, timezone
from io import BytesIO

import httpx
import pytest
from sqlalchemy import create_engine, text

from app.core.config import Settings
from app.services import no_movement_feishu_reports as service
from app.services.data_foundation.contracts import get_contract
from app.services.no_movement_feishu_reports import (
    NoMovementFeishuReportError,
    RECIPIENT_KEY,
    _FeishuRejected,
    _analysis,
    _feishu_interactive_content,
    _feishu_markdown,
    _feishu_payload,
    _make_png,
    _make_xlsx,
    _read_prepared_artifact,
    _remove_prepared_artifacts,
    _store_prepared_artifacts,
    _recipient_key,
    _utc_now,
    confirm_report,
    prepare_report,
    send_report,
)
from app.services.no_movement_operational_analysis import build_bilingual_analysis


def _response(status: int, payload: dict) -> httpx.Response:
    return httpx.Response(status, json=payload, request=httpx.Request("POST", "https://open.feishu.cn/test"))


def test_pilot_recipient_is_fixed_to_the_authorized_private_test_user():
    _recipient_key(RECIPIENT_KEY)
    with pytest.raises(NoMovementFeishuReportError) as error:
        _recipient_key("ANY_OTHER_RECIPIENT")
    assert error.value.status_code == 403
    assert error.value.code == "no_movement_report_recipient_invalid"


def _national_prepare_fixture(monkeypatch):
    import uuid

    engine = create_engine("sqlite://")
    with engine.begin() as connection:
        connection.exec_driver_sql("""
            CREATE TABLE no_movement_feishu_reports (
                report_id BLOB PRIMARY KEY, source_id TEXT, publication_id BLOB,
                publication_version INTEGER, contract_version INTEGER, map_version_id BLOB,
                region_key TEXT, recipient_key TEXT, recipient_email TEXT, recipient_type TEXT,
                recipient_identifier TEXT, recipient_name TEXT, actor_user_id INTEGER,
                analysis_text TEXT, interpretation TEXT, comments TEXT, recommendations TEXT,
                record_count INTEGER, total_stopped_orders INTEGER, dataset_sha256 BLOB,
                rules_version TEXT, report_state TEXT, delivery_context TEXT,
                created_at DATETIME, updated_at DATETIME
            )
        """)

    publication_id, map_version_id = uuid.uuid4(), uuid.uuid4()
    publication = {
        "publication_id": publication_id.bytes, "version_no": 12, "contract_version": 2,
        "map_version_id": map_version_id.bytes, "period_start": date(2026, 10, 26),
        "snapshot_as_of": datetime(2026, 10, 26, 9, 15), "publication_state": "PUBLISHED",
    }
    payload = {
        "publication": publication, "regionCode": None,
        "rows": [
            {"resolved_base_code": "001", "mapped_region": "SPS"},
            {"resolved_base_code": None, "mapped_region": None},
        ],
        "message": "Análise nacional", "recordCount": 2, "total": 23,
        "datasetDigest": b"d" * 32, "fileName": "sem_movimentacao_nacional_2026-10-26.xlsx",
        "xlsx": b"national xlsx", "png": b"national png", "regions": ["PR", "SPS"],
    }

    def report_payload(*args, publication_id, region_code, **kwargs):
        assert publication_id == str(uuid.UUID(bytes=publication["publication_id"]))
        assert region_code is None
        return payload

    monkeypatch.setattr("app.services.no_movement_feishu_reports._report_payload", report_payload)
    return engine, publication_id, payload


@pytest.mark.parametrize(
    "prior_state",
    [
        "PARTIAL",
        "PENDING",
        "SENDING",
    ],
)
def test_national_prepare_blocks_unresolved_prior_delivery(tmp_path, monkeypatch, prior_state):
    import uuid

    engine, publication_id, _ = _national_prepare_fixture(monkeypatch)
    prior_report_id = uuid.uuid4()
    with engine.begin() as connection:
        connection.execute(text("""
            INSERT INTO no_movement_feishu_reports
                (report_id, publication_id, region_key, recipient_key, actor_user_id,
                 rules_version, report_state)
            VALUES (:report_id, :publication_id, 'ALL', :recipient_key, 7, :rules_version, :state)
        """), {
            "report_id": prior_report_id.bytes, "publication_id": publication_id.bytes,
            "recipient_key": RECIPIENT_KEY, "rules_version": service.RULES_VERSION,
            "state": prior_state,
        })

    with pytest.raises(NoMovementFeishuReportError) as error:
        prepare_report(
            engine, publication_id=str(publication_id), region_code=None,
            recipient_key=RECIPIENT_KEY, actor_user_id=7, staging_directory=tmp_path,
        )

    assert error.value.code == "no_movement_report_delivery_unresolved"
    with engine.connect() as connection:
        rows = connection.execute(text("SELECT report_id, report_state FROM no_movement_feishu_reports")).all()
    assert rows == [(prior_report_id.bytes, prior_state)]
    assert not list((tmp_path / ".no-movement-feishu-reports").glob("*"))


def test_national_prepare_allows_new_report_after_sent_and_marks_prior_history(tmp_path, monkeypatch):
    import uuid

    engine, publication_id, _ = _national_prepare_fixture(monkeypatch)
    prior_report_id = uuid.uuid4()
    with engine.begin() as connection:
        connection.execute(text("""
            INSERT INTO no_movement_feishu_reports
                (report_id, publication_id, region_key, recipient_key, actor_user_id,
                 rules_version, report_state)
            VALUES (:report_id, :publication_id, 'ALL', :recipient_key, 7, :rules_version, 'SENT')
        """), {"report_id": prior_report_id.bytes, "publication_id": publication_id.bytes,
               "recipient_key": RECIPIENT_KEY, "rules_version": service.RULES_VERSION})

    prepared = prepare_report(
        engine, publication_id=str(publication_id), region_code=None,
        recipient_key=RECIPIENT_KEY, actor_user_id=7, staging_directory=tmp_path,
    )
    assert uuid.UUID(prepared["reportId"]) != prior_report_id
    assert prepared["previouslySent"] is True
    assert prepared["recordCount"] == 2
    with engine.connect() as connection:
        states = connection.execute(text(
            "SELECT report_state FROM no_movement_feishu_reports"
        )).scalars().all()
    assert sorted(states) == ["PREPARED", "SENT"]


@pytest.mark.parametrize("region_code", ["PR", "SPS"])
def test_regional_prepare_remains_independent_from_national_scope(tmp_path, monkeypatch, region_code):
    engine, publication_id, payload = _national_prepare_fixture(monkeypatch)

    def report_payload(*args, publication_id, region_code: str | None, **kwargs):
        assert region_code in {"PR", "SPS"}
        return {**payload, "regionCode": region_code, "rows": [payload["rows"][0]], "recordCount": 1}

    monkeypatch.setattr("app.services.no_movement_feishu_reports._report_payload", report_payload)
    prepared = prepare_report(
        engine, publication_id=str(publication_id), region_code=region_code,
        recipient_key=RECIPIENT_KEY, actor_user_id=7, staging_directory=tmp_path,
    )

    assert prepared["regionCode"] == region_code
    assert prepared["regionLabel"] == region_code
    assert prepared["recordCount"] == 1


def test_national_prepare_includes_unmapped_rows_and_ignores_completed_regional_report(tmp_path, monkeypatch):
    import uuid

    engine, publication_id, _ = _national_prepare_fixture(monkeypatch)
    with engine.begin() as connection:
        connection.execute(text("""
            INSERT INTO no_movement_feishu_reports
                (report_id, publication_id, region_key, recipient_key, actor_user_id,
                 rules_version, report_state)
            VALUES (:report_id, :publication_id, 'PR', :recipient_key, 7, :rules_version, 'SENT')
        """), {
            "report_id": uuid.uuid4().bytes, "publication_id": publication_id.bytes,
            "recipient_key": RECIPIENT_KEY, "rules_version": service.RULES_VERSION,
        })

    prepared = prepare_report(
        engine, publication_id=str(publication_id), region_code=None,
        recipient_key=RECIPIENT_KEY, actor_user_id=7, staging_directory=tmp_path,
    )

    assert prepared["regionCode"] is None
    assert prepared["regionLabel"] == "Todas as regionais"
    assert prepared["recordCount"] == 2
    with engine.connect() as connection:
        scopes = connection.execute(text(
            "SELECT region_key, report_state FROM no_movement_feishu_reports ORDER BY region_key"
        )).all()
    assert scopes == [("ALL", "PREPARED"), ("PR", "SENT")]


def test_sr_group_options_require_sr_in_the_publications_pinned_mapping(monkeypatch):
    import uuid

    from app.services import feishu_report_recipients

    engine = create_engine("sqlite://")
    sr_map_id, other_map_id = uuid.uuid4(), uuid.uuid4()
    sr_publication_id, other_publication_id = uuid.uuid4(), uuid.uuid4()
    with engine.begin() as connection:
        connection.exec_driver_sql("CREATE TABLE base_mapping_entries (map_version_id BLOB, regional_name TEXT)")
        connection.execute(text("INSERT INTO base_mapping_entries VALUES (:map,:region)"), {"map": sr_map_id.bytes, "region": "SR"})
        connection.execute(text("INSERT INTO base_mapping_entries VALUES (:map,:region)"), {"map": other_map_id.bytes, "region": "SPS"})

    publications = {
        str(sr_publication_id): {
            "publication_id": sr_publication_id.bytes, "version_no": 12, "contract_version": 2,
            "map_version_id": sr_map_id.bytes, "period_start": date(2026, 10, 26),
            "published_at_sao_paulo": None,
        },
        str(other_publication_id): {
            "publication_id": other_publication_id.bytes, "version_no": 13, "contract_version": 2,
            "map_version_id": other_map_id.bytes, "period_start": date(2026, 10, 26),
            "published_at_sao_paulo": None,
        },
    }
    monkeypatch.setattr(service, "_current_publication", lambda _engine, publication_id, _connection=None: publications[publication_id])
    enabled = Settings(_env_file=None, feishu_report_recipient_sr_enabled=True)
    monkeypatch.setattr(service, "get_settings", lambda: enabled)
    monkeypatch.setattr(feishu_report_recipients, "get_settings", lambda: enabled)

    options = service.get_report_options(engine, publication_id=str(sr_publication_id), region_code="SR")
    assert options["mapVersionId"] == str(sr_map_id)
    sr_option = next(item for item in options["recipients"] if item["key"] == "SR_COLLECTION_GROUP")
    assert sr_option["available"] is True

    with pytest.raises(NoMovementFeishuReportError) as missing_sr:
        service.get_report_options(engine, publication_id=str(other_publication_id), region_code="SR")
    assert missing_sr.value.code == "no_movement_feishu_region_invalid"

    national = service.get_report_options(engine, publication_id=str(sr_publication_id))
    national_sr = next(item for item in national["recipients"] if item["key"] == "SR_COLLECTION_GROUP")
    assert national_sr["available"] is False

    disabled = Settings(_env_file=None, feishu_report_recipient_sr_enabled=False)
    monkeypatch.setattr(service, "get_settings", lambda: disabled)
    monkeypatch.setattr(feishu_report_recipients, "get_settings", lambda: disabled)
    disabled_sr = service.get_report_options(engine, publication_id=str(sr_publication_id), region_code="SR")
    assert next(item for item in disabled_sr["recipients"] if item["key"] == "SR_COLLECTION_GROUP")["available"] is False


def test_explicit_feishu_http_rejection_is_distinguished_from_ambiguous_server_failure():
    with pytest.raises(_FeishuRejected):
        _feishu_payload(_response(400, {"code": 999, "msg": "rejected"}))
    with pytest.raises(RuntimeError, match="feishu_http_outcome_uncertain"):
        _feishu_payload(_response(503, {"code": 1}))


def test_feishu_api_rejection_is_detected_even_when_http_status_is_success():
    with pytest.raises(_FeishuRejected):
        _feishu_payload(_response(200, {"code": 999, "msg": "rejected"}))


def test_feishu_markdown_formats_bold_and_preserves_operational_text_and_literal_asterisks():
    analysis = (
        "📦 **Resumo SPS — 22/10/2026**\n"
        "Pedidos parados: **1.234** (12,5%)\n"
        "- **Base São Paulo**: 12\n"
        "- **网点 深圳**: 8\n"
        "Responsável: João / 张伟\n"
        "Nota literal: 2 * 3; escape: \\*não é negrito\\*; "
        "marcadores literais: \\*\\*texto literal\\*\\*; sem par: **texto"
    )

    rendered = _feishu_markdown(analysis)

    assert "**Resumo SPS — 22/10/2026**" in rendered
    assert "**1.234**" in rendered
    assert "- **Base São Paulo**: 12\n- **网点 深圳**: 8" in rendered
    assert "📦" in rendered and "João / 张伟" in rendered
    assert "2 \\* 3" in rendered
    assert r"\*não é negrito\*" in rendered
    assert r"\*\*texto literal\*\*" in rendered
    assert "**texto" not in rendered
    assert r"\*\*texto" in rendered


def test_feishu_interactive_content_keeps_large_multilingual_message_and_line_breaks():
    import json

    analysis = "**Análise atual**\n" + ("SPS — Operação / 运营 📦\n" * 800)
    content = _feishu_interactive_content(analysis)
    card = json.loads(content)

    assert card["schema"] == "2.0"
    assert card["body"]["elements"][0]["tag"] == "markdown"
    markdown = card["body"]["elements"][0]["content"]
    assert markdown.startswith("**Análise atual**\n")
    assert markdown.count("\n") == analysis.count("\n")
    assert "运营 📦" in markdown


def test_feishu_interactive_content_rejects_messages_over_platform_limit_without_truncation():
    analysis = "SPS — operação / 运营 📦\n" * 2000
    with pytest.raises(NoMovementFeishuReportError) as error:
        _feishu_interactive_content(analysis)
    assert error.value.code == "no_movement_report_message_too_large"
    assert error.value.status_code == 413


def test_disabled_pilot_blocks_send_before_database_or_http_access(monkeypatch):
    def forbidden_client(*args, **kwargs):
        pytest.fail("Feishu client must not be initialized while pilot flag is disabled")

    monkeypatch.setattr("app.services.no_movement_feishu_reports.httpx.Client", forbidden_client)
    with pytest.raises(NoMovementFeishuReportError) as error:
        send_report(None, report_id="not-read-when-disabled", actor_user_id=1, settings=Settings(feishu_no_movement_pilot_enabled=False))
    assert error.value.status_code == 503
    assert error.value.code == "no_movement_feishu_pilot_disabled"


def test_confirmation_persists_exact_edited_analysis_and_clears_legacy_text_fields(monkeypatch):
    engine = create_engine("sqlite://")
    with engine.begin() as connection:
        connection.exec_driver_sql("""
            CREATE TABLE no_movement_feishu_reports (
                report_id BLOB PRIMARY KEY, publication_id BLOB NOT NULL,
                region_key TEXT, recipient_key TEXT, recipient_email TEXT,
                recipient_type TEXT, recipient_identifier TEXT, recipient_name TEXT,
                actor_user_id INTEGER NOT NULL, report_state TEXT NOT NULL,
                analysis_text TEXT NOT NULL, interpretation TEXT NOT NULL,
                comments TEXT NOT NULL, recommendations TEXT NOT NULL,
                confirmed_at DATETIME, updated_at DATETIME
            )
        """)
        connection.execute(text("""
            INSERT INTO no_movement_feishu_reports
              (report_id,publication_id,region_key,recipient_key,recipient_email,actor_user_id,report_state,
               analysis_text,interpretation,comments,recommendations,confirmed_at,updated_at)
            VALUES (:report_id,:publication_id,'ALL','PHELIPPE_PRIVATE',:recipient_email,7,'PREPARED',
                    'generated','old','old','old',NULL,NULL)
        """), {"report_id": bytes.fromhex("72727272727272727272727272727272"), "publication_id": b"p" * 16,
              "recipient_email": service.RECIPIENT_EMAIL})
    monkeypatch.setattr(
        "app.services.no_movement_feishu_reports._current_publication",
        lambda *args, **kwargs: {},
    )
    edited = "分析 SPS\nLinha final editada — ç"
    result = confirm_report(engine, report_id="72727272-7272-7272-7272-727272727272", actor_user_id=7, analysis_text=edited)
    assert result == {"reportId": "72727272-7272-7272-7272-727272727272", "state": "CONFIRMED", "messageSent": False}
    with engine.connect() as connection:
        row = connection.execute(text("SELECT analysis_text, interpretation, comments, recommendations FROM no_movement_feishu_reports")).one()
    assert tuple(row) == (edited, "", "", "")


def test_confirmation_rejects_blank_analysis_before_database_access():
    with pytest.raises(NoMovementFeishuReportError) as error:
        confirm_report(None, report_id="not-used", actor_user_id=7, analysis_text="   ")
    assert error.value.code == "no_movement_report_text_required"


def test_regional_analysis_is_scoped_while_default_national_text_is_unchanged():
    common = {
        "operational_date": date(2026, 10, 26),
        "export_time": datetime(2026, 10, 26, 9, 15),
        "publication_version": 12,
        "total": 80,
        "groups": [
            {"mapped": True, "region": "SPS", "base_name": "Base A", "rm_responsible": "RGM 1", "total": 60},
            {"mapped": False, "region": "Sem mapeamento", "base_name": "Unidade X", "total": 20},
        ],
    }
    national = build_bilingual_analysis(**common)
    regional = build_bilingual_analysis(**{**common, "total": 60, "groups": common["groups"][:1]}, region_scope="SPS")
    assert "SEM MOVIMENTAÇÃO | 无流转**" in national
    assert "Volume nacional: 80 pacotes" in national
    assert "REGIONAL SPS | 区域 SPS" in regional
    assert "Volume regional: 60 pacotes" in regional


def test_pilot_xlsx_contract_keeps_all_eighteen_source_columns_and_analysis_total_includes_unmapped_national_rows():
    contract = get_contract("no_movement", 2)
    assert len(contract.columns) == 18
    publication = {
        "period_start": date(2026, 10, 26),
        "snapshot_as_of": datetime(2026, 10, 26, 9, 15),
        "version_no": 12,
    }
    rows = [
        {"stopped_orders": 60, "resolved_base_code": "001", "mapped_region": "SPS", "mapped_base_name": "Base A", "mapped_rm": "RGM 1", "source_base_code": "001", "source_base_name": "Base A", "reported_region": "SPS"},
        {"stopped_orders": 20, "resolved_base_code": None, "mapped_region": None, "mapped_base_name": None, "mapped_rm": None, "source_base_code": "999", "source_base_name": "Unidade X", "reported_region": "SPS"},
    ]
    national, national_total = _analysis(publication, rows, None)
    regional, regional_total = _analysis(publication, rows[:1], "SPS")
    assert national_total == 80 and "Volume nacional: 80 pacotes" in national
    assert "Unidade X (999) — **20 pacotes**" in national
    assert regional_total == 60 and "Volume regional: 60 pacotes" in regional


def test_real_xlsx_and_png_generation_preserves_contract_rows_and_regional_scope():
    from openpyxl import load_workbook
    from PIL import Image

    contract = get_contract("no_movement", 2)

    def make_row(row_no: int, stopped: int, region: str | None, code: str | None) -> dict:
        row = {column.key: None for column in contract.columns}
        row.update({
            "source_row_no": row_no,
            "reported_region": region,
            "source_base_code": code,
            "source_base_name": f"Loja {row_no}",
            "stopped_orders": stopped,
            "stopped_over_7_days": stopped // 2,
            "stopped_over_14_days": stopped // 3,
            "stopped_over_30_days": stopped // 4,
            "resolved_base_code": code,
            "mapped_region": region,
            "mapped_base_name": f"Base {row_no}" if code else None,
            "mapped_rm": "Responsável fictício" if code else None,
        })
        return row

    unmapped = make_row(4, 5, None, None)
    unmapped["reported_region"] = "SPS"
    rows = [
        make_row(2, 11, "SPS", "001"),
        make_row(3, 7, "SR", "002"),
        unmapped,  # Reported SPS is not enough to assign a DoomsDay regional scope.
    ]
    publication = {
        "period_start": date(2026, 10, 26),
        "snapshot_as_of": datetime(2026, 10, 26, 12, 30),
        "version_no": 17,
    }

    national_message, national_total = _analysis(publication, rows, None)
    sps_rows = [row for row in rows if row["mapped_region"] == "SPS" and row["resolved_base_code"]]
    sps_message, sps_total = _analysis(publication, sps_rows, "SPS")
    assert national_total == 23 and "Volume nacional: 23 pacotes" in national_message
    assert sps_total == 11 and "Volume regional: 11 pacotes" in sps_message
    assert "REGIONAL SPS" in sps_message and "REGIONAL SR" not in sps_message

    national_name, national_xlsx = _make_xlsx(rows, region_code=None, operational_date=publication["period_start"])
    sps_name, sps_xlsx = _make_xlsx(sps_rows, region_code="SPS", operational_date=publication["period_start"])
    assert national_name.endswith("_nacional_2026-10-26.xlsx")
    assert sps_name.endswith("_SPS_2026-10-26.xlsx")
    for content, expected_rows in ((national_xlsx, rows), (sps_xlsx, sps_rows)):
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        sheet = workbook["Dados"]
        exported = list(sheet.iter_rows(values_only=True))
        assert len(exported[0]) == 18
        assert exported[0] == tuple(column.header for column in contract.columns)
        assert len(exported) == len(expected_rows) + 1
        stopped_index = next(index for index, column in enumerate(contract.columns) if column.key == "stopped_orders")
        assert [row[stopped_index] for row in exported[1:]] == [row["stopped_orders"] for row in expected_rows]
        workbook.close()

    for content in (_make_png(rows, region_code=None, publication=publication, total=national_total),
                    _make_png(sps_rows, region_code="SPS", publication=publication, total=sps_total)):
        image = Image.open(BytesIO(content))
        image.verify()
        image = Image.open(BytesIO(content))
        assert image.format == "PNG" and image.size == (1800, 650)
        rgb = image.convert("RGB")
        assert rgb.getpixel((900, 260)) == (255, 255, 255)
        assert any(color == (230, 0, 18) for _, color in rgb.getcolors(maxcolors=1_200_000) or [])


def test_repeated_sps_xlsx_and_png_generation_is_byte_stable_for_same_input():
    contract = get_contract("no_movement", 2)
    row = {column.key: None for column in contract.columns}
    row.update({
        "source_row_no": 1, "source_base_code": "001", "source_base_name": "Base SPS",
        "stopped_orders": 12, "stopped_over_7_days": 6, "stopped_over_14_days": 3,
        "stopped_over_30_days": 1, "resolved_base_code": "001", "mapped_region": "SPS",
        "mapped_base_name": "Base SPS", "mapped_rm": "Responsável fictício",
    })
    publication = {
        "period_start": date(2026, 10, 26),
        "snapshot_as_of": datetime(2026, 10, 26, 9, 15),
        "version_no": 12,
    }
    first_xlsx = _make_xlsx([row], region_code="SPS", operational_date=publication["period_start"])[1]
    second_xlsx = _make_xlsx([row], region_code="SPS", operational_date=publication["period_start"])[1]
    first_png = _make_png([row], region_code="SPS", publication=publication, total=12)
    second_png = _make_png([row], region_code="SPS", publication=publication, total=12)
    assert first_xlsx == second_xlsx
    assert first_png == second_png
    changed_row = {**row, "stopped_orders": 13}
    changed_xlsx = _make_xlsx([changed_row], region_code="SPS", operational_date=publication["period_start"])[1]
    changed_png = _make_png([row], region_code="SPS", publication=publication, total=13)
    assert hashlib.sha256(first_xlsx).digest() != hashlib.sha256(changed_xlsx).digest()
    assert hashlib.sha256(first_png).digest() != hashlib.sha256(changed_png).digest()


def test_prepared_sps_artifacts_are_reused_and_real_changes_or_tampering_are_rejected(tmp_path):
    import hashlib
    import uuid

    report_id = uuid.uuid4()
    xlsx = b"xlsx bytes prepared for SPS"
    png = b"png bytes prepared for SPS"
    _store_prepared_artifacts(tmp_path, report_id, xlsx, png)
    assert _read_prepared_artifact(tmp_path, report_id, "xlsx", hashlib.sha256(xlsx).hexdigest()) == xlsx
    assert _read_prepared_artifact(tmp_path, report_id, "png", hashlib.sha256(png).hexdigest()) == png

    changed_xlsx_path = tmp_path / ".no-movement-feishu-reports" / f"{report_id.hex}.xlsx"
    changed_xlsx_path.write_bytes(b"tampered xlsx")
    with pytest.raises(NoMovementFeishuReportError) as xlsx_error:
        _read_prepared_artifact(tmp_path, report_id, "xlsx", hashlib.sha256(xlsx).hexdigest())
    assert xlsx_error.value.code == "no_movement_report_artifacts_changed"

    changed_png_path = tmp_path / ".no-movement-feishu-reports" / f"{report_id.hex}.png"
    changed_png_path.write_bytes(b"tampered png")
    with pytest.raises(NoMovementFeishuReportError) as png_error:
        _read_prepared_artifact(tmp_path, report_id, "png", hashlib.sha256(png).hexdigest())
    assert png_error.value.code == "no_movement_report_artifacts_changed"
    _remove_prepared_artifacts(tmp_path, report_id)
    assert not changed_xlsx_path.exists() and not changed_png_path.exists()


def test_prepare_and_confirm_sps_persists_exact_text_and_immutable_attachment_hashes(tmp_path, monkeypatch):
    import hashlib
    import json
    import uuid

    engine = create_engine("sqlite://")
    with engine.begin() as connection:
        connection.exec_driver_sql("""
            CREATE TABLE no_movement_feishu_reports (
                report_id BLOB PRIMARY KEY, source_id TEXT, publication_id BLOB,
                publication_version INTEGER, contract_version INTEGER, map_version_id BLOB,
                region_key TEXT, recipient_key TEXT, recipient_email TEXT, recipient_type TEXT,
                recipient_identifier TEXT, recipient_name TEXT, actor_user_id INTEGER,
                analysis_text TEXT, interpretation TEXT, comments TEXT, recommendations TEXT,
                record_count INTEGER, total_stopped_orders INTEGER, dataset_sha256 BLOB,
                rules_version TEXT, report_state TEXT, delivery_context TEXT,
                created_at DATETIME, updated_at DATETIME, confirmed_at DATETIME, sent_at DATETIME
            )
        """)

    publication_id = uuid.uuid4()
    map_version_id = uuid.uuid4()
    publication = {
        "publication_id": publication_id.bytes, "version_no": 12, "contract_version": 2,
        "map_version_id": map_version_id.bytes, "period_start": date(2026, 10, 26),
        "snapshot_as_of": datetime(2026, 10, 26, 9, 15), "publication_state": "PUBLISHED",
    }
    xlsx, png = b"prepared SPS xlsx", b"prepared SPS png"
    payload = {
        "publication": publication, "regionCode": "SPS", "rows": [{"resolved_base_code": "001", "mapped_region": "SPS"}],
        "message": "Análise SPS original", "recordCount": 1, "total": 12,
        "datasetDigest": b"d" * 32, "fileName": "sem_movimentacao_SPS_2026-10-26.xlsx",
        "xlsx": xlsx, "png": png, "regions": ["SPS"],
    }
    monkeypatch.setattr("app.services.no_movement_feishu_reports._report_payload", lambda *args, **kwargs: payload)
    monkeypatch.setattr("app.services.no_movement_feishu_reports._current_publication", lambda *args, **kwargs: publication)

    prepared = prepare_report(
        engine, publication_id=str(publication_id), region_code="SPS", recipient_key=RECIPIENT_KEY,
        actor_user_id=7, staging_directory=tmp_path,
    )
    report_uuid = uuid.UUID(prepared["reportId"])
    assert prepared["regionCode"] == "SPS"
    assert (tmp_path / ".no-movement-feishu-reports" / f"{report_uuid.hex}.xlsx").read_bytes() == xlsx
    assert (tmp_path / ".no-movement-feishu-reports" / f"{report_uuid.hex}.png").read_bytes() == png

    edited = "Análise SPS editada e confirmada — texto exato"
    confirm_report(engine, report_id=prepared["reportId"], actor_user_id=7, analysis_text=edited)
    with engine.connect() as connection:
        stored = connection.execute(text("SELECT analysis_text, report_state, delivery_context FROM no_movement_feishu_reports")).one()
    context = json.loads(stored.delivery_context)
    assert stored.analysis_text == edited and stored.report_state == "CONFIRMED"
    assert context["artifacts"] == {
        "xlsxSha256": hashlib.sha256(xlsx).hexdigest(),
        "pngSha256": hashlib.sha256(png).hexdigest(),
    }


@pytest.mark.parametrize(("recipient_key", "region_code", "receive_type", "receive_id", "recipient_flag", "recipient_name"), [
    (RECIPIENT_KEY, "SPS", "email", service.RECIPIENT_EMAIL, None, "Phelippe · Pessoa / 个人"),
    ("NATIONAL_RMS", None, "chat_id", "oc_8474b27a4b40adb3ecb8c6577773d222", "feishu_report_recipient_national_rms_enabled", "Grupo Nacional de RMs 全国RM"),
    ("MG_COLLECTION_GROUP", "MG", "chat_id", "oc_d3ef49bb0698066065c5556ce484c6ed", "feishu_report_recipient_mg_enabled", "MG 揽收沟通群 · Coordenação de Coleta First Mile"),
])
def test_send_uses_confirmed_artifacts_and_allowlisted_destination_without_regeneration(
    tmp_path, monkeypatch, recipient_key, region_code, receive_type, receive_id, recipient_flag, recipient_name,
):
    import uuid

    from app.services import no_movement_feishu_reports as service

    engine = create_engine("sqlite://")
    publication_id, map_version_id = uuid.uuid4(), uuid.uuid4()
    publication = {
        "publication_id": publication_id.bytes, "version_no": 12, "contract_version": 2,
        "map_version_id": map_version_id.bytes, "period_start": date(2026, 10, 26),
        "snapshot_as_of": datetime(2026, 10, 26, 9, 15), "publication_state": "PUBLISHED",
    }
    with engine.begin() as connection:
        connection.exec_driver_sql("""
            CREATE TABLE no_movement_feishu_reports (
                report_id BLOB PRIMARY KEY, source_id TEXT, publication_id BLOB,
                publication_version INTEGER, contract_version INTEGER, map_version_id BLOB,
                region_key TEXT, recipient_key TEXT, recipient_email TEXT, recipient_type TEXT,
                recipient_identifier TEXT, recipient_name TEXT, actor_user_id INTEGER,
                analysis_text TEXT, interpretation TEXT, comments TEXT, recommendations TEXT,
                record_count INTEGER, total_stopped_orders INTEGER, dataset_sha256 BLOB,
                rules_version TEXT, report_state TEXT, delivery_context TEXT,
                created_at DATETIME, updated_at DATETIME, confirmed_at DATETIME, sent_at DATETIME
            )
        """)
        connection.exec_driver_sql("CREATE TABLE source_snapshot_head (source_id TEXT, publication_id BLOB)")
        connection.exec_driver_sql("""
            CREATE TABLE source_publications (
                source_id TEXT, publication_id BLOB, version_no INTEGER, contract_version INTEGER,
                map_version_id BLOB, period_start DATE, snapshot_as_of DATETIME, publication_state TEXT
            )
        """)
        connection.execute(text("INSERT INTO source_snapshot_head VALUES ('no_movement', :publication_id)"), {"publication_id": publication_id.bytes})
        connection.execute(text("""
            INSERT INTO source_publications VALUES
            ('no_movement', :publication_id, 12, 2, :map_id, :period_start, :snapshot_as_of, 'PUBLISHED')
        """), {"publication_id": publication_id.bytes, "map_id": map_version_id.bytes,
              "period_start": publication["period_start"], "snapshot_as_of": publication["snapshot_as_of"]})

    xlsx, png = b"confirmed SPS workbook bytes", b"confirmed SPS image bytes"
    digest = b"d" * 32
    payload = {
        "publication": publication, "regionCode": region_code, "rows": [{"resolved_base_code": "001", "mapped_region": region_code}],
        "message": "Original analysis", "recordCount": 1, "total": 12, "datasetDigest": digest,
        "fileName": f"sem_movimentacao_{region_code}_2026-10-26.xlsx", "xlsx": xlsx, "png": png, "regions": [region_code],
    }
    include_artifact_calls = []

    def report_payload(*args, include_artifacts=True, **kwargs):
        include_artifact_calls.append(include_artifacts)
        return {**payload, "xlsx": xlsx if include_artifacts else None, "png": png if include_artifacts else None}

    monkeypatch.setattr(service, "_report_payload", report_payload)
    monkeypatch.setattr(service, "_current_publication", lambda *args, **kwargs: publication)
    settings_values = {recipient_flag: True} if recipient_flag else {}
    settings = Settings(feishu_no_movement_pilot_enabled=True, feishu_oauth_app_id="fake", feishu_oauth_app_secret="fake", **settings_values)
    monkeypatch.setattr(service, "get_settings", lambda: settings)
    from app.services import feishu_report_recipients
    monkeypatch.setattr(feishu_report_recipients, "get_settings", lambda: settings)
    prepared = prepare_report(
        engine, publication_id=str(publication_id), region_code=region_code, recipient_key=recipient_key,
        actor_user_id=7, staging_directory=tmp_path,
    )
    with engine.connect() as connection:
        audit_destination = connection.execute(text("""
            SELECT recipient_key, recipient_email, recipient_type, recipient_identifier, recipient_name
            FROM no_movement_feishu_reports WHERE report_id=:report_id
        """), {"report_id": uuid.UUID(prepared["reportId"]).bytes}).one()
    assert audit_destination == (
        recipient_key,
        receive_id if receive_type == "email" else None,
        receive_type,
        receive_id,
        recipient_name,
    )
    edited = f"**Resumo {region_code}**\nTexto confirmado — 内容原文 📦\nBases: **001, Base {region_code}**"
    confirm_report(engine, report_id=prepared["reportId"], actor_user_id=7, analysis_text=edited)
    replaced_publication_id = uuid.uuid4()
    with engine.begin() as connection:
        connection.execute(text("UPDATE source_snapshot_head SET publication_id=:new_id"), {"new_id": replaced_publication_id.bytes})
    with pytest.raises(NoMovementFeishuReportError) as replaced_error:
        send_report(engine, report_id=prepared["reportId"], actor_user_id=7, staging_directory=tmp_path, settings=settings)
    assert replaced_error.value.code == "no_movement_publication_not_current"
    with engine.begin() as connection:
        connection.execute(text("UPDATE source_snapshot_head SET publication_id=:old_id"), {"old_id": publication_id.bytes})
    captured = {"uploads": {}, "messages": [], "targets": []}

    class FakeClient:
        def __init__(self, timeout=None):
            self.timeout = timeout

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

        def post(self, url, **kwargs):
            if url == service.TOKEN_URL:
                return _response(200, {"code": 0, "tenant_access_token": "test-token"})
            if url == service.FILES_URL:
                captured["uploads"]["xlsx"] = kwargs["files"]["file"][1].read()
                return _response(200, {"code": 0, "data": {"file_key": "fake-file-key"}})
            if url == service.IMAGES_URL:
                captured["uploads"]["png"] = kwargs["files"]["image"][1].read()
                return _response(200, {"code": 0, "data": {"image_key": "fake-image-key"}})
            if url == service.MESSAGE_URL:
                body = kwargs["json"]
                captured["messages"].append(body)
                captured["targets"].append((kwargs["params"]["receive_id_type"], body["receive_id"]))
                return _response(200, {"code": 0, "data": {"message_id": f"message-{len(captured['messages'])}"}})
            raise AssertionError(f"Unexpected HTTP endpoint: {url}")

    monkeypatch.setattr(service.httpx, "Client", FakeClient)
    result = send_report(
        engine, report_id=prepared["reportId"], actor_user_id=7, staging_directory=tmp_path,
        settings=settings,
    )
    assert result["state"] == "SENT"
    assert captured["targets"] == [(receive_type, receive_id)] * 3
    assert include_artifact_calls == [True, False]
    assert captured["uploads"] == {"xlsx": xlsx, "png": png}
    import json
    assert captured["messages"][0]["msg_type"] == "interactive"
    card = json.loads(captured["messages"][0]["content"])
    assert card["schema"] == "2.0"
    assert card["body"]["elements"][0]["tag"] == "markdown"
    assert card["body"]["elements"][0]["content"] == _feishu_markdown(edited)
    assert f"**Resumo {region_code}**" in card["body"]["elements"][0]["content"]
    with engine.connect() as connection:
        persisted_text = connection.execute(text("SELECT analysis_text FROM no_movement_feishu_reports WHERE report_id=:id"),
                                            {"id": uuid.UUID(prepared["reportId"]).bytes}).scalar_one()
    assert persisted_text == edited
    assert not (tmp_path / ".no-movement-feishu-reports" / f"{uuid.UUID(prepared['reportId']).hex}.xlsx").exists()
    with pytest.raises(NoMovementFeishuReportError) as duplicate_error:
        send_report(engine, report_id=prepared["reportId"], actor_user_id=7,
                    staging_directory=tmp_path, settings=settings)
    assert duplicate_error.value.code == "no_movement_report_not_confirmed"
    second = prepare_report(
        engine, publication_id=str(publication_id), region_code=region_code, recipient_key=recipient_key,
        actor_user_id=7, staging_directory=tmp_path,
    )
    assert second["reportId"] != prepared["reportId"]
    assert second["previouslySent"] is True
    confirm_report(engine, report_id=second["reportId"], actor_user_id=7, analysis_text=edited)
    second_result = send_report(
        engine, report_id=second["reportId"], actor_user_id=7,
        staging_directory=tmp_path, settings=settings,
    )
    assert second_result["state"] == "SENT"
    assert captured["targets"] == [(receive_type, receive_id)] * 6
    assert include_artifact_calls == [True, False, True, False]
    if recipient_key == "NATIONAL_RMS":
        assert all(target == ("chat_id", "oc_8474b27a4b40adb3ecb8c6577773d222") for target in captured["targets"])
    with engine.connect() as connection:
        history = connection.execute(text(
            "SELECT report_id, report_state FROM no_movement_feishu_reports ORDER BY created_at, report_id"
        )).all()
    assert len(history) == 2
    assert all(row.report_state == "SENT" for row in history)


def test_mysql_datetime_values_are_created_as_naive_utc():
    before = datetime.now(timezone.utc).replace(tzinfo=None)
    value = _utc_now()
    after = datetime.now(timezone.utc).replace(tzinfo=None)
    assert value.tzinfo is None  # MySQL DATETIME(6) stores the explicit UTC wall time.
    assert before <= value <= after
