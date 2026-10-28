from __future__ import annotations

import io
import json
from datetime import date, datetime
import os
from pathlib import Path
import uuid

import pytest
from openpyxl import load_workbook
from PIL import Image

from app.core.config import Settings
from app.services import operational_feishu_reports as service
from app.services.no_movement_feishu_reports import _feishu_markdown


def test_source_allowlist_and_flags_are_separate_and_off_by_default():
    settings = Settings(_env_file=None)
    assert settings.feishu_collection_rate_report_enabled is False
    assert settings.feishu_pdd_collection_failure_report_enabled is False
    assert service._feature_enabled("collection_rate", settings) is False
    assert service._feature_enabled("pdd_collection_failure", settings) is False
    with pytest.raises(service.OperationalFeishuReportError):
        service._source("collection_monitoring")


def test_collection_rate_analysis_uses_official_weighted_formula_and_unmapped_national_rows():
    publication = {"version_no": 7}
    rows = [
        {"deadline_at": datetime(2026, 10, 22), "resolved_base_code": "001", "mapping_resolution": "CODE_MATCH", "mapped_region": "SPS", "mapped_base_name": "Base Um", "mapped_rm": "RM A", "collected_quantity": 90, "quantity_to_collect": 100, "order_quantity": 105},
        {"deadline_at": datetime(2026, 10, 22), "resolved_base_code": "002", "mapping_resolution": "NAME_APPROVED", "mapped_region": "SPE", "mapped_base_name": "Base Dois", "mapped_rm": "RM B", "collected_quantity": 40, "quantity_to_collect": 50, "order_quantity": 51},
        {"deadline_at": datetime(2026, 10, 22), "resolved_base_code": None, "mapping_resolution": "UNMAPPED", "mapped_region": None, "mapped_base_name": None, "mapped_rm": None, "collected_quantity": 10, "quantity_to_collect": 20, "order_quantity": 22},
    ]
    message, result, totals = service._analysis("collection_rate", publication, rows, None, date(2026, 10, 22))
    assert totals == {"collected": 140, "planned": 170, "totalOrders": 178, "rateNumerator": 140, "rateDenominator": 170}
    assert result["operationalDate"] == "2026-10-22"
    assert result["result"]["metric"]["numeratorField"] == "collected_quantity"
    assert result["result"]["metric"]["denominatorField"] == "quantity_to_collect"
    assert result["result"]["unmappedPlanned"] == 20
    assert "82,4%" in message
    assert "Volume sem regional/base DoomsDay" in message


def test_pdd_analysis_uses_collection_column_and_keeps_unmapped_in_national_totals():
    publication = {"version_no": 4}
    rows = [
        {"data_date": date(2026, 10, 22), "resolved_base_code": "001", "mapping_resolution": "CODE_MATCH", "mapped_region": "SPS", "mapped_base_name": "PDD Um", "mapped_rm": "RM SPS", "total_orders": 100, "orders_unscanned_at_collection": 7},
        {"data_date": date(2026, 10, 22), "resolved_base_code": None, "mapping_resolution": "UNMAPPED", "mapped_region": None, "mapped_base_name": None, "mapped_rm": None, "total_orders": 25, "orders_unscanned_at_collection": 3},
    ]
    message, result, totals = service._analysis("pdd_collection_failure", publication, rows, None, date(2026, 10, 22))
    assert totals == {"collected": 115, "planned": 125, "totalOrders": 125, "rateNumerator": 10, "rateDenominator": 125}
    assert result["result"]["metric"]["unscannedField"] == "orders_unscanned_at_collection"
    assert result["result"]["unmapped"]["totalOrders"] == 25
    assert "8,0%" in message


def test_collection_rate_excel_contains_all_contract_columns_and_all_rows():
    from app.services.data_foundation.contracts import get_contract

    contract = get_contract("collection_rate", 3)
    rows = [{column.key: None for column in contract.columns} for _ in range(2)]
    rows[0].update({"deadline_at": datetime(2026, 10, 22), "reported_region": "SPS", "order_quantity": 105, "quantity_to_collect": 100, "collected_quantity": 90})
    rows[1].update({"deadline_at": datetime(2026, 10, 21), "reported_region": None, "order_quantity": 22, "quantity_to_collect": 20, "collected_quantity": 10})
    name, content = service._xlsx("collection_rate", rows, date(2026, 10, 22), None)
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    assert name.endswith("2026-10-22.xlsx")
    assert workbook.sheetnames == ["Dados"]
    sheet = workbook["Dados"]
    values = list(sheet.values)
    assert len(values[0]) == 22
    assert len(values) == 3
    assert values[1][5] == 105


def test_pdd_excel_preserves_fifteen_columns_and_one_date_named_sheet_per_day():
    from app.services.data_foundation.contracts import get_contract

    contract = get_contract("pdd_collection_failure", 2)
    rows = [{column.key: None for column in contract.columns} for _ in range(2)]
    rows[0].update({"data_date": date(2026, 10, 21), "source_base_code": "001", "total_orders": 10, "orders_unscanned_at_collection": 1})
    rows[1].update({"data_date": date(2026, 10, 22), "source_base_code": "002", "total_orders": 20, "orders_unscanned_at_collection": 2})
    _, content = service._xlsx("pdd_collection_failure", rows, date(2026, 10, 22), "SPS")
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    assert workbook.sheetnames == ["2026-10-21", "2026-10-22"]
    assert all(len(list(workbook[sheet].values)[0]) == 15 and len(list(workbook[sheet].values)) == 2 for sheet in workbook.sheetnames)


def test_png_is_valid_image_and_reports_source_specific_metrics():
    publication = {"version_no": 3, "contract_version": 3}
    totals = {"totalOrders": 100, "rateNumerator": 98, "rateDenominator": 100, "collected": 98, "planned": 100}
    result = {"operationalDate": "2026-10-22", "result": {"regions": []}}
    content = service._png("collection_rate", publication, "SPS", totals, result)
    image = Image.open(io.BytesIO(content))
    assert image.format == "PNG"
    assert image.size == (1400, 760)


def test_recipient_is_fixed_before_database_or_artifact_access():
    with pytest.raises(service.OperationalFeishuReportError) as recipient_error:
        service.prepare_report(None, source_id="collection_rate", publication_id="bad", region_code=None,
                               recipient_key="arbitrary", actor_user_id=1, staging_directory=__import__("pathlib").Path("/tmp"),
                               settings=_enabled_settings("collection_rate"))
    assert recipient_error.value.code == "operational_feishu_recipient_invalid"


def _report_engine(source_id: str, *, database_url: str = "sqlite://", extra_rows: int = 0):
    from sqlalchemy import create_engine, text
    from app.services.data_foundation.contracts import get_contract

    source = service._source(source_id)
    contract = get_contract(source_id, source["contract"])
    engine = create_engine(database_url)
    pub_id, map_id = uuid.uuid4(), uuid.uuid4()
    with engine.begin() as connection:
        # The production sources are daily publications. Their dashboard
        # current rules are per-date metadata / source_daily_head, with no
        # source_snapshot_head row (the condition that caused the 409).
        connection.execute(text("CREATE TABLE source_publications(source_id TEXT, publication_id BLOB, version_no INTEGER, contract_version INTEGER, map_version_id BLOB, period_start TEXT, period_end TEXT, row_count INTEGER, snapshot_as_of TEXT, published_at TEXT, publication_state TEXT)"))
        connection.execute(text("CREATE TABLE source_publication_dates(source_id TEXT, publication_id BLOB, data_date TEXT, row_count INTEGER)"))
        connection.execute(text("CREATE TABLE source_daily_head(source_id TEXT, data_date TEXT, publication_id BLOB, PRIMARY KEY(source_id,data_date))"))
        connection.execute(text("CREATE TABLE base_mapping_entries(map_version_id BLOB, base_code TEXT, regional_name TEXT, base_name TEXT, rm_responsible TEXT)"))
        connection.execute(text(f"CREATE TABLE {source['table']}(source_id TEXT, publication_id BLOB, map_version_id BLOB, source_row_no INTEGER, source_sheet_name TEXT, resolved_base_code TEXT, mapping_resolution TEXT, " + ", ".join(f"`{column.key}` TEXT" for column in contract.columns) + ")"))
        connection.execute(text("""CREATE TABLE operational_feishu_reports(
          report_id BLOB, source_id TEXT, publication_id BLOB, publication_version INTEGER, contract_version INTEGER,
          map_version_id BLOB, region_key TEXT, recipient_key TEXT, recipient_email TEXT, recipient_type TEXT,
          recipient_identifier TEXT, recipient_name TEXT, actor_user_id INTEGER,
          analysis_text TEXT, record_count INTEGER, dataset_sha256 BLOB, rules_version TEXT, report_state TEXT,
          delivery_context TEXT, created_at TEXT, updated_at TEXT, confirmed_at TEXT, sent_at TEXT)"""))
        connection.execute(text("INSERT INTO source_publications VALUES (:s,:p,5,:cv,:m,'2026-10-22','2026-10-22',:count,'2026-10-22','2026-10-22 12:00:00','PUBLISHED')"), {"s": source_id, "p": pub_id.bytes, "cv": source["contract"], "m": map_id.bytes, "count": extra_rows + 2})
        connection.execute(text("INSERT INTO source_publication_dates VALUES (:s,:p,'2026-10-22',:count)"), {"s": source_id, "p": pub_id.bytes, "count": extra_rows + 2})
        if source_id == "pdd_collection_failure":
            connection.execute(text("INSERT INTO source_daily_head VALUES (:s,'2026-10-22',:p)"), {"s": source_id, "p": pub_id.bytes})
        connection.execute(text("INSERT INTO base_mapping_entries VALUES (:m,'001','SPS','Base SPS','RM SPS')"), {"m": map_id.bytes})
        connection.execute(text("INSERT INTO base_mapping_entries VALUES (:m,'002','SPE','Base SPE','RM SPE')"), {"m": map_id.bytes})
        values = {column.key: None for column in contract.columns}
        if source_id == "collection_rate":
            values.update({"deadline_at": "2026-10-22 12:00:00", "order_quantity": "12", "quantity_to_collect": "10", "collected_quantity": "9"})
        else:
            values.update({"data_date": "2026-10-22", "source_base_code": "001", "source_base_name": "Base SPS", "total_orders": "12", "orders_unscanned_at_collection": "2"})
        keys = ["source_id", "publication_id", "map_version_id", "source_row_no", "source_sheet_name", "resolved_base_code", "mapping_resolution"] + [column.key for column in contract.columns]
        inserted = {"source_id": source_id, "publication_id": pub_id.bytes, "map_version_id": map_id.bytes,
                    "source_row_no": 2, "source_sheet_name": "2026-10-22" if source["sheet_mode"] == "daily" else None,
                    "resolved_base_code": "001", "mapping_resolution": "CODE_MATCH", **values}
        connection.execute(text(f"INSERT INTO {source['table']} ({','.join(f'`{key}`' for key in keys)}) VALUES ({','.join(':'+key for key in keys)})"), inserted)
        second_values = {column.key: None for column in contract.columns}
        if source_id == "collection_rate":
            second_values.update({"deadline_at": "2026-10-22 12:00:00", "reported_region": "SPE", "order_quantity": "8", "quantity_to_collect": "8", "collected_quantity": "7"})
        else:
            second_values.update({"data_date": "2026-10-22", "source_base_code": "002", "source_base_name": "Base SPE", "total_orders": "8", "orders_unscanned_at_collection": "1"})
        second_row = {"source_id": source_id, "publication_id": pub_id.bytes, "map_version_id": map_id.bytes,
                      "source_row_no": 3, "source_sheet_name": "2026-10-22" if source["sheet_mode"] == "daily" else None,
                      "resolved_base_code": "002", "mapping_resolution": "CODE_MATCH", **second_values}
        connection.execute(text(f"INSERT INTO {source['table']} ({','.join(f'`{key}`' for key in keys)}) VALUES ({','.join(':'+key for key in keys)})"), second_row)
        insert_sql = text(f"INSERT INTO {source['table']} ({','.join(f'`{key}`' for key in keys)}) VALUES ({','.join(':'+key for key in keys)})")
        batch = []
        for index in range(extra_rows):
            base_code = "001" if index % 2 == 0 else "002"
            row_values = {column.key: None for column in contract.columns}
            if source_id == "collection_rate":
                row_values.update({"deadline_at": "2026-10-22 12:00:00", "order_quantity": "22", "quantity_to_collect": "20", "collected_quantity": "19"})
            else:
                row_values.update({"data_date": "2026-10-22", "source_base_code": base_code, "source_base_name": f"Base {base_code}", "total_orders": "20", "orders_unscanned_at_collection": "2"})
            batch.append({"source_id": source_id, "publication_id": pub_id.bytes, "map_version_id": map_id.bytes,
                          "source_row_no": index + 4, "source_sheet_name": "2026-10-22" if source["sheet_mode"] == "daily" else None,
                          "resolved_base_code": base_code, "mapping_resolution": "CODE_MATCH", **row_values})
            if len(batch) == 2_000:
                connection.execute(insert_sql, batch)
                batch.clear()
        if batch:
            connection.execute(insert_sql, batch)
    return engine, str(pub_id), map_id


@pytest.mark.parametrize("source_id", ["collection_rate", "pdd_collection_failure"])
def test_regional_payload_is_exactly_scoped_by_the_publication_mapping(source_id: str):
    engine, publication_id, _ = _report_engine(source_id)
    national = service._payload(engine, source_id, publication_id, None)
    regional = service._payload(engine, source_id, publication_id, "SPS")

    assert len(national["rows"]) == 2
    assert {row["mapped_region"] for row in national["rows"]} == {"SPS", "SPE"}
    assert len(regional["rows"]) == 1
    assert {row["mapped_region"] for row in regional["rows"]} == {"SPS"}
    assert {row["resolved_base_code"] for row in regional["rows"]} == {"001"}


@pytest.mark.parametrize("source_id", ["collection_rate", "pdd_collection_failure"])
def test_current_daily_publication_does_not_require_snapshot_head_and_rejects_superseded(source_id: str):
    from sqlalchemy import text

    engine, publication_id, map_id = _report_engine(source_id)
    current = service._publication(engine, source_id, publication_id)
    assert current["version_no"] == 5
    assert bytes(current["map_version_id"]) == map_id.bytes

    newer_id = uuid.uuid4()
    contract_version = service._source(source_id)["contract"]
    with engine.begin() as connection:
        connection.execute(text("INSERT INTO source_publications VALUES (:s,:p,6,:cv,:m,'2026-10-22','2026-10-22',2,'2026-10-22','2026-10-23 12:00:00','PUBLISHED')"),
                           {"s": source_id, "p": newer_id.bytes, "cv": contract_version, "m": map_id.bytes})
        connection.execute(text("INSERT INTO source_publication_dates VALUES (:s,:p,'2026-10-22',2)"),
                           {"s": source_id, "p": newer_id.bytes})
        if source_id == "pdd_collection_failure":
            connection.execute(text("UPDATE source_daily_head SET publication_id=:new WHERE source_id=:s AND data_date='2026-10-22'"),
                               {"s": source_id, "new": newer_id.bytes})

    assert service._publication(engine, source_id, str(newer_id))["version_no"] == 6
    with pytest.raises(service.OperationalFeishuReportError) as old_publication:
        service._publication(engine, source_id, publication_id)
    assert old_publication.value.code == "operational_feishu_publication_not_current"


@pytest.mark.parametrize("source_id", ["collection_rate", "pdd_collection_failure"])
@pytest.mark.parametrize("mutation", ["wrong_contract", "unpublished", "inconsistent_metadata"])
def test_current_daily_publication_requires_exact_contract_and_published_state(source_id: str, mutation: str):
    from sqlalchemy import text

    engine, publication_id, _ = _report_engine(source_id)
    column, value = {
        "wrong_contract": ("contract_version", 99),
        "unpublished": ("publication_state", "PREPARED"),
        "inconsistent_metadata": ("row_count", 99),
    }[mutation]
    with engine.begin() as connection:
        connection.execute(text(f"UPDATE source_publications SET {column}=:value WHERE publication_id=:p"),
                           {"value": value, "p": uuid.UUID(publication_id).bytes})
    with pytest.raises(service.OperationalFeishuReportError) as invalid:
        service._publication(engine, source_id, publication_id)
    assert invalid.value.code == "operational_feishu_publication_not_current"


@pytest.mark.parametrize("source_id", ["collection_rate", "pdd_collection_failure"])
@pytest.mark.parametrize("field,changed_value,error_code", [
    ("version_no", 6, "operational_feishu_publication_not_current"),
    ("map_version_id", uuid.UUID(int=999).bytes, "operational_feishu_mapping_changed"),
])
def test_confirmation_rejects_publication_version_or_mapping_changed_after_preparation(
    tmp_path: Path, source_id: str, field: str, changed_value: object, error_code: str,
):
    from sqlalchemy import text

    engine, publication_id, _ = _report_engine(source_id)
    prepared = service.prepare_report(
        engine, source_id=source_id, publication_id=publication_id, region_code="SPS",
        recipient_key=service.RECIPIENT_KEY, actor_user_id=71, staging_directory=tmp_path,
        settings=_enabled_settings(source_id),
    )
    with engine.begin() as connection:
        connection.execute(text(f"UPDATE source_publications SET {field}=:value WHERE publication_id=:p"),
                           {"value": changed_value, "p": uuid.UUID(publication_id).bytes})
    with pytest.raises(service.OperationalFeishuReportError) as changed:
        service.confirm_report(
            engine, report_id=prepared["reportId"], actor_user_id=71, analysis_text=prepared["analysis"],
            source_id=source_id, settings=_enabled_settings(source_id),
        )
    assert changed.value.code == error_code


@pytest.mark.parametrize("source_id", ["collection_rate", "pdd_collection_failure"])
def test_report_source_id_isolation_rejects_publication_from_another_source(source_id: str):
    engine, publication_id, _ = _report_engine(source_id)
    other_source = "pdd_collection_failure" if source_id == "collection_rate" else "collection_rate"
    with pytest.raises(service.OperationalFeishuReportError) as wrong_source:
        service._publication(engine, other_source, publication_id)
    assert wrong_source.value.code == "operational_feishu_publication_not_current"


@pytest.mark.parametrize("source_id", ["collection_rate", "pdd_collection_failure"])
def test_send_revalidates_daily_current_publication_before_any_feishu_request(
    tmp_path: Path, monkeypatch, source_id: str,
):
    from sqlalchemy import text

    engine, publication_id, map_id = _report_engine(source_id)
    fake = _FakeFeishuClient()
    monkeypatch.setattr(service.httpx, "Client", lambda **_kwargs: fake)
    prepared = service.prepare_report(
        engine, source_id=source_id, publication_id=publication_id, region_code="SPS",
        recipient_key=service.RECIPIENT_KEY, actor_user_id=71, staging_directory=tmp_path,
        settings=_enabled_settings(source_id),
    )
    service.confirm_report(
        engine, report_id=prepared["reportId"], actor_user_id=71, analysis_text=prepared["analysis"],
        source_id=source_id, settings=_enabled_settings(source_id),
    )
    newer_id = uuid.uuid4()
    with engine.begin() as connection:
        connection.execute(text("INSERT INTO source_publications VALUES (:s,:p,6,:cv,:m,'2026-10-22','2026-10-22',2,'2026-10-22','2026-10-23 12:00:00','PUBLISHED')"),
                           {"s": source_id, "p": newer_id.bytes, "cv": service._source(source_id)["contract"], "m": map_id.bytes})
        connection.execute(text("INSERT INTO source_publication_dates VALUES (:s,:p,'2026-10-22',2)"),
                           {"s": source_id, "p": newer_id.bytes})
        if source_id == "pdd_collection_failure":
            connection.execute(text("UPDATE source_daily_head SET publication_id=:new WHERE source_id=:s AND data_date='2026-10-22'"),
                               {"s": source_id, "new": newer_id.bytes})

    with pytest.raises(service.OperationalFeishuReportError) as replaced:
        service.send_report(
            engine, source_id=source_id, report_id=prepared["reportId"], actor_user_id=71,
            staging_directory=tmp_path, settings=_enabled_settings(source_id),
        )
    assert replaced.value.code == "operational_feishu_publication_not_current"
    assert fake.calls == []


@pytest.mark.parametrize("source_id", ["collection_rate", "pdd_collection_failure"])
def test_disabled_source_flag_blocks_options_prepare_confirm_and_send_before_publication_lookup(source_id: str, tmp_path: Path):
    disabled = Settings(_env_file=None)
    with pytest.raises(service.OperationalFeishuReportError) as options_error:
        service.get_options(None, source_id=source_id, publication_id=str(uuid.uuid4()), settings=disabled)
    assert options_error.value.code == "operational_feishu_send_disabled"

    with pytest.raises(service.OperationalFeishuReportError) as prepare_error:
        service.prepare_report(None, source_id=source_id, publication_id=str(uuid.uuid4()), region_code=None,
                               recipient_key=service.RECIPIENT_KEY, actor_user_id=1,
                               staging_directory=tmp_path, settings=disabled)
    assert prepare_error.value.code == "operational_feishu_send_disabled"

    with pytest.raises(service.OperationalFeishuReportError) as confirm_error:
        service.confirm_report(None, report_id=str(uuid.uuid4()), actor_user_id=1,
                               analysis_text="confirmed", source_id=source_id, settings=disabled)
    assert confirm_error.value.code == "operational_feishu_send_disabled"

    with pytest.raises(service.OperationalFeishuReportError) as send_error:
        service.send_report(None, source_id=source_id, report_id=str(uuid.uuid4()), actor_user_id=1,
                            staging_directory=tmp_path, settings=disabled)
    assert send_error.value.code == "operational_feishu_send_disabled"


@pytest.mark.parametrize("source_id", ["collection_rate", "pdd_collection_failure"])
def test_streamed_preparation_matches_reference_analysis_digest_and_excel_contract(tmp_path: Path, source_id: str):
    engine, publication_id, _ = _report_engine(source_id)
    source = service._source(source_id)
    date_field = source["date_field"]
    # Keep a second PDD date in its own sheet while the analysis remains on the
    # latest published date. Taxa V3 stays in its single contract sheet.
    if source["sheet_mode"] == "daily":
        with engine.begin() as connection:
            connection.execute(__import__("sqlalchemy").text(
                f"UPDATE {source['table']} SET {date_field}='2026-10-21' WHERE source_row_no=3"
            ), {})
    reference = service._payload(engine, source_id, publication_id, None)
    first_id = uuid.uuid4()
    streamed = service._stream_payload(engine, source_id=source_id, publication_id=publication_id,
                                       region_code=None, staging_directory=tmp_path, report_id=first_id)
    second_id = uuid.uuid4()
    repeated = service._stream_payload(engine, source_id=source_id, publication_id=publication_id,
                                       region_code=None, staging_directory=tmp_path, report_id=second_id)

    assert streamed["recordCount"] == len(reference["rows"])
    assert streamed["digest"] == reference["digest"]
    assert streamed["analysis"] == reference["analysis"]
    assert streamed["analysisResult"] == reference["analysisResult"]
    assert streamed["totals"] == reference["totals"]
    assert streamed["fileName"] == reference["fileName"]
    private = tmp_path / ".operational-feishu-reports"
    xlsx_path = private / f"{first_id.hex}.xlsx"
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    expected_sheets = ["Dados"] if source["sheet_mode"] == "single" else ["2026-10-21", "2026-10-22"]
    assert workbook.sheetnames == expected_sheets
    exported_rows = sum(sum(1 for _ in sheet.iter_rows(values_only=True)) - 1 for sheet in workbook.worksheets)
    assert exported_rows == len(reference["rows"])
    assert xlsx_path.read_bytes() == (private / f"{second_id.hex}.xlsx").read_bytes()
    assert (private / f"{first_id.hex}.png").read_bytes() == (private / f"{second_id.hex}.png").read_bytes()


@pytest.mark.parametrize("source_id", ["collection_rate", "pdd_collection_failure"])
def test_streaming_xlsx_size_limit_fails_without_ready_or_partial_artifacts(tmp_path: Path, monkeypatch, source_id: str):
    engine, publication_id, _ = _report_engine(source_id)
    monkeypatch.setattr(service, "MAX_FILE_BYTES", 512)
    report_id = uuid.uuid4()
    with pytest.raises(service.OperationalFeishuReportError) as error:
        service._stream_payload(engine, source_id=source_id, publication_id=publication_id,
                                region_code=None, staging_directory=tmp_path, report_id=report_id)
    assert error.value.code == "operational_feishu_xlsx_invalid"
    private = tmp_path / ".operational-feishu-reports"
    assert list(private.iterdir()) == []


@pytest.mark.parametrize("source_id", ["collection_rate", "pdd_collection_failure"])
def test_generation_timeout_leaves_no_report_or_partial_artifact(tmp_path: Path, monkeypatch, source_id: str):
    engine, publication_id, _ = _report_engine(source_id)
    monkeypatch.setattr(service, "MAX_GENERATION_SECONDS", 0.0)
    with pytest.raises(service.OperationalFeishuReportError) as error:
        service.prepare_report(engine, source_id=source_id, publication_id=publication_id, region_code=None,
                              recipient_key=service.RECIPIENT_KEY, actor_user_id=71, staging_directory=tmp_path,
                              settings=_enabled_settings(source_id))
    assert error.value.code == "operational_feishu_generation_timeout"
    assert error.value.status_code == 504
    with engine.connect() as connection:
        assert connection.execute(__import__("sqlalchemy").text("SELECT COUNT(*) FROM operational_feishu_reports")).scalar_one() == 0
    private = tmp_path / ".operational-feishu-reports"
    assert list(private.iterdir()) == []


@pytest.mark.parametrize("source_id", ["collection_rate", "pdd_collection_failure"])
def test_prepare_confirm_and_disabled_send_use_same_publication_map_and_immutable_artifacts(tmp_path: Path, monkeypatch, source_id: str):
    engine, publication_id, map_id = _report_engine(source_id)
    monkeypatch.setattr(service, "get_settings", lambda: Settings(_env_file=None))
    prepared = service.prepare_report(engine, source_id=source_id, publication_id=publication_id, region_code="SPS",
                                     recipient_key=service.RECIPIENT_KEY, actor_user_id=71, staging_directory=tmp_path,
                                     settings=_enabled_settings(source_id))
    assert prepared["state"] == "PREPARED"
    assert prepared["regionCode"] == "SPS"
    assert prepared["mapVersionId"] == str(map_id)
    assert prepared["recordCount"] == 1
    assert prepared["excelBytes"] > 0 and prepared["pngBytes"] > 0
    report_uuid = uuid.UUID(prepared["reportId"])
    private = tmp_path / ".operational-feishu-reports"
    assert (private / f"{report_uuid.hex}.xlsx").is_file()
    assert (private / f"{report_uuid.hex}.png").is_file()
    if os.name == "posix":
        assert (private.stat().st_mode & 0o777) == 0o700
        assert ((private / f"{report_uuid.hex}.xlsx").stat().st_mode & 0o777) == 0o600
    edited = prepared["analysis"] + "\n**Revisado pelo administrador**"
    confirmed = service.confirm_report(engine, report_id=prepared["reportId"], actor_user_id=71, analysis_text=edited,
                                       source_id=source_id, settings=_enabled_settings(source_id))
    assert confirmed["state"] == "CONFIRMED"
    with engine.connect() as connection:
        row = connection.execute(__import__("sqlalchemy").text("SELECT analysis_text,report_state FROM operational_feishu_reports WHERE report_id=:id"), {"id": report_uuid.bytes}).first()
    assert row[0] == edited and row[1] == "CONFIRMED"
    with pytest.raises(service.OperationalFeishuReportError) as disabled:
        service.send_report(engine, source_id=source_id, report_id=prepared["reportId"], actor_user_id=71,
                            staging_directory=tmp_path, settings=Settings(_env_file=None))
    assert disabled.value.code == "operational_feishu_send_disabled"
    assert (private / f"{report_uuid.hex}.xlsx").is_file()


class _FakeFeishuClient:
    def __init__(self, *, fail_url: str | None = None, ambiguous_message: bool = False):
        self.fail_url = fail_url
        self.ambiguous_message = ambiguous_message
        self.calls: list[tuple[str, dict[str, object]]] = []

    def __enter__(self):
        return self

    def __exit__(self, *_args):
        return False

    def post(self, url: str, **kwargs):
        import httpx

        self.calls.append((url, kwargs))
        if url == self.fail_url:
            if self.ambiguous_message:
                raise httpx.ReadTimeout("simulated timeout", request=httpx.Request("POST", url))
            return httpx.Response(403, json={"code": 999, "msg": "permission denied"})
        if url == service.TOKEN_URL:
            body = {"code": 0, "tenant_access_token": "mock-token"}
        elif url == service.FILES_URL:
            body = {"code": 0, "data": {"file_key": "mock-file-key"}}
        elif url == service.IMAGES_URL:
            body = {"code": 0, "data": {"image_key": "mock-image-key"}}
        else:
            body = {"code": 0, "data": {"message_id": f"message-{len(self.calls)}"}}
        return httpx.Response(200, json=body)


def _enabled_settings(source_id: str) -> Settings:
    return Settings(
        _env_file=None,
        **{
            service._source(source_id)["flag"]: True,
            "feishu_oauth_app_id": "test-app-id",
            "feishu_oauth_app_secret": "test-app-secret",
        },
    )


@pytest.mark.parametrize("collection_enabled,pdd_enabled", [
    (False, False), (True, False), (False, True), (True, True),
])
@pytest.mark.parametrize("source_id", ["collection_rate", "pdd_collection_failure"])
def test_availability_is_source_specific_and_validates_the_current_publication(
    source_id: str, collection_enabled: bool, pdd_enabled: bool,
):
    engine, publication_id, _ = _report_engine(source_id)
    settings = Settings(
        _env_file=None,
        feishu_collection_rate_report_enabled=collection_enabled,
        feishu_pdd_collection_failure_report_enabled=pdd_enabled,
    )
    expected_enabled = collection_enabled if source_id == "collection_rate" else pdd_enabled
    result = service.get_availability(
        engine if expected_enabled else None,
        source_id=source_id,
        publication_id=publication_id,
        settings=settings,
    )
    assert result == {"enabled": expected_enabled, "publicationValid": expected_enabled}


@pytest.mark.parametrize("source_id", ["collection_rate", "pdd_collection_failure"])
def test_availability_fails_closed_for_unavailable_publication_and_tracks_runtime_flag_changes(source_id: str):
    engine, publication_id, _ = _report_engine(source_id)
    flag_name = service._source(source_id)["flag"]
    disabled = Settings(_env_file=None, **{flag_name: False})
    enabled = Settings(_env_file=None, **{flag_name: True})

    assert service.get_availability(None, source_id=source_id, publication_id=publication_id, settings=disabled) == {
        "enabled": False, "publicationValid": False,
    }
    assert service.get_availability(engine, source_id=source_id, publication_id=publication_id, settings=enabled) == {
        "enabled": True, "publicationValid": True,
    }
    assert service.get_availability(None, source_id=source_id, publication_id=publication_id, settings=disabled) == {
        "enabled": False, "publicationValid": False,
    }
    invalid = str(uuid.uuid4())
    assert service.get_availability(engine, source_id=source_id, publication_id=invalid, settings=enabled) == {
        "enabled": True, "publicationValid": False,
    }


def test_availability_endpoint_requires_effective_matrix_admin_and_returns_only_booleans(monkeypatch):
    from types import SimpleNamespace
    from fastapi import HTTPException
    from app.api import data_foundation as api

    route = next(route for route in api.router.routes if getattr(route, "path", "") == "/v2/operational-feishu/{source_id}/availability")
    assert any(dependency.call is api.require_matrix_admin for dependency in route.dependant.dependencies)
    monkeypatch.setattr(api, "_require_feature", lambda: None)
    monkeypatch.setattr(api, "_engine", lambda: object())
    monkeypatch.setattr(api, "get_operational_feishu_report_availability", lambda *args, **kwargs: {
        "enabled": True, "publicationValid": True,
    })
    matrix_admin = SimpleNamespace(
        user=SimpleNamespace(platform_role="ADMIN"),
        identity={"administrative_scope": "matrix", "organizational_scope": "matrix", "role": "matrix"},
    )
    assert api.operational_feishu_report_availability(
        "collection_rate", str(uuid.uuid4()), matrix_admin,
    ) == {"enabled": True, "publicationValid": True}
    regional_admin = SimpleNamespace(
        user=SimpleNamespace(platform_role="ADMIN"),
        identity={"administrative_scope": "matrix", "organizational_scope": "regional", "role": "regional"},
    )
    with pytest.raises(HTTPException) as forbidden:
        api.operational_feishu_report_availability("collection_rate", str(uuid.uuid4()), regional_admin)
    assert forbidden.value.status_code == 403


@pytest.mark.parametrize("source_id", ["collection_rate", "pdd_collection_failure"])
def test_send_uses_confirmed_analysis_and_same_immutable_artifacts_without_real_network(tmp_path: Path, monkeypatch, source_id: str):
    engine, publication_id, _ = _report_engine(source_id)
    fake = _FakeFeishuClient()
    monkeypatch.setattr(service.httpx, "Client", lambda **_kwargs: fake)
    prepared = service.prepare_report(
        engine, source_id=source_id, publication_id=publication_id, region_code="SPS",
        recipient_key=service.RECIPIENT_KEY, actor_user_id=71, staging_directory=tmp_path,
        settings=_enabled_settings(source_id),
    )
    edited = prepared["analysis"] + "\n**Conferido manualmente**"
    service.confirm_report(engine, report_id=prepared["reportId"], actor_user_id=71, analysis_text=edited,
                           source_id=source_id, settings=_enabled_settings(source_id))

    result = service.send_report(
        engine, source_id=source_id, report_id=prepared["reportId"], actor_user_id=71,
        staging_directory=tmp_path, settings=_enabled_settings(source_id),
    )

    assert result["state"] == "SENT"
    sent_messages = [kwargs["json"] for url, kwargs in fake.calls if url == service.MESSAGE_URL]
    assert len(sent_messages) == 3
    assert sent_messages[0]["msg_type"] == "interactive"
    card = json.loads(sent_messages[0]["content"])
    rendered = card["body"]["elements"][0]["content"]
    assert rendered == _feishu_markdown(edited)
    assert "**Conferido manualmente**" in rendered
    assert sent_messages[1]["msg_type"] == "file"
    assert sent_messages[2]["msg_type"] == "image"
    with engine.connect() as connection:
        stored_state = connection.execute(
            __import__("sqlalchemy").text("SELECT report_state FROM operational_feishu_reports WHERE report_id=:id"),
            {"id": uuid.UUID(prepared["reportId"]).bytes},
        ).scalar_one()
    assert stored_state == "SENT"
    call_count = len(fake.calls)
    with pytest.raises(service.OperationalFeishuReportError) as duplicate:
        service.send_report(
            engine, source_id=source_id, report_id=prepared["reportId"], actor_user_id=71,
            staging_directory=tmp_path, settings=_enabled_settings(source_id),
        )
    assert duplicate.value.code == "operational_feishu_report_not_confirmed"
    assert len(fake.calls) == call_count

    second = service.prepare_report(
        engine, source_id=source_id, publication_id=publication_id, region_code="SPS",
        recipient_key=service.RECIPIENT_KEY, actor_user_id=71, staging_directory=tmp_path,
        settings=_enabled_settings(source_id),
    )
    assert second["previouslySent"] is True
    assert second["reportId"] != prepared["reportId"]
    service.confirm_report(
        engine, report_id=second["reportId"], actor_user_id=71, analysis_text=second["analysis"],
        source_id=source_id, settings=_enabled_settings(source_id),
    )
    resent = service.send_report(
        engine, source_id=source_id, report_id=second["reportId"], actor_user_id=71,
        staging_directory=tmp_path, settings=_enabled_settings(source_id),
    )
    assert resent["state"] == "SENT"
    assert len([kwargs for url, kwargs in fake.calls if url == service.MESSAGE_URL]) == 6
    with engine.connect() as connection:
        states = connection.execute(
            __import__("sqlalchemy").text("SELECT report_state FROM operational_feishu_reports")
        ).scalars().all()
    assert states.count("SENT") == 2


@pytest.mark.parametrize("source_id", ["collection_rate", "pdd_collection_failure"])
def test_national_rms_can_receive_full_publication_including_unmapped_rows(source_id: str, tmp_path: Path, monkeypatch):
    from sqlalchemy import text
    from app.services.data_foundation.contracts import get_contract

    engine, publication_id, map_id = _report_engine(source_id)
    source = service._source(source_id)
    contract = get_contract(source_id, source["contract"])
    values = {column.key: None for column in contract.columns}
    if source_id == "collection_rate":
        values.update({"deadline_at": "2026-10-22 12:00:00", "order_quantity": "5",
                       "quantity_to_collect": "4", "collected_quantity": "3"})
    else:
        values.update({"data_date": "2026-10-22", "total_orders": "5",
                       "orders_unscanned_at_collection": "2"})
    keys = ["source_id", "publication_id", "map_version_id", "source_row_no", "source_sheet_name",
            "resolved_base_code", "mapping_resolution"] + [column.key for column in contract.columns]
    parameters = {"source_id": source_id, "publication_id": uuid.UUID(publication_id).bytes,
                  "map_version_id": map_id.bytes, "source_row_no": 4,
                  "source_sheet_name": "2026-10-22" if source["sheet_mode"] == "daily" else None,
                  "resolved_base_code": None, "mapping_resolution": "UNMAPPED", **values}
    with engine.begin() as connection:
        connection.execute(text(f"INSERT INTO {source['table']} ({','.join(f'`{key}`' for key in keys)}) VALUES ({','.join(':'+key for key in keys)})"), parameters)
        connection.execute(text("UPDATE source_publications SET row_count=3 WHERE publication_id=:p"),
                           {"p": uuid.UUID(publication_id).bytes})
        connection.execute(text("UPDATE source_publication_dates SET row_count=3 WHERE publication_id=:p"),
                           {"p": uuid.UUID(publication_id).bytes})

    settings = _enabled_settings(source_id).model_copy(update={"feishu_report_recipient_national_rms_enabled": True})
    options = service.get_options(engine, source_id=source_id, publication_id=publication_id, settings=settings)
    national_option = next(item for item in options["recipients"] if item["key"] == "NATIONAL_RMS")
    assert national_option["available"] is True
    assert national_option["regionCode"] is None
    regional_options = service.get_options(
        engine, source_id=source_id, publication_id=publication_id, region_code="SPS", settings=settings,
    )
    assert next(item for item in regional_options["recipients"] if item["key"] == "NATIONAL_RMS")["available"] is False

    fake = _FakeFeishuClient()
    monkeypatch.setattr(service.httpx, "Client", lambda **_kwargs: fake)
    prepared = service.prepare_report(
        engine, source_id=source_id, publication_id=publication_id, region_code=None,
        recipient_key="NATIONAL_RMS", actor_user_id=71, staging_directory=tmp_path, settings=settings,
    )
    assert prepared["recordCount"] == 3
    assert prepared["regionCode"] is None
    assert prepared["previouslySent"] is False
    with engine.connect() as connection:
        audit = connection.execute(text("""
            SELECT region_key,recipient_email,recipient_type,recipient_identifier,recipient_name,report_state
            FROM operational_feishu_reports WHERE report_id=:id
        """), {"id": uuid.UUID(prepared["reportId"]).bytes}).one()
    assert audit == ("ALL", None, "chat_id", "oc_8474b27a4b40adb3ecb8c6577773d222",
                     "Grupo Nacional de RMs 全国RM", "PREPARED")
    service.confirm_report(
        engine, report_id=prepared["reportId"], actor_user_id=71, analysis_text=prepared["analysis"],
        source_id=source_id, settings=settings,
    )
    result = service.send_report(
        engine, source_id=source_id, report_id=prepared["reportId"], actor_user_id=71,
        staging_directory=tmp_path, settings=settings,
    )
    assert result["state"] == "SENT"
    messages = [kwargs for url, kwargs in fake.calls if url == service.MESSAGE_URL]
    assert len(messages) == 3
    assert all(call["params"] == {"receive_id_type": "chat_id"} for call in messages)
    assert all(call["json"]["receive_id"] == "oc_8474b27a4b40adb3ecb8c6577773d222" for call in messages)


@pytest.mark.parametrize("source_id", ["collection_rate", "pdd_collection_failure"])
@pytest.mark.parametrize("prior_state", ["SENDING", "PARTIAL", "PENDING"])
def test_prepare_blocks_unresolved_delivery_for_the_same_source_scope(source_id: str, prior_state: str, tmp_path: Path):
    from sqlalchemy import text

    engine, publication_id, _ = _report_engine(source_id)
    with engine.begin() as connection:
        connection.execute(text("""
            INSERT INTO operational_feishu_reports
              (report_id,source_id,publication_id,region_key,recipient_key,rules_version,report_state)
            VALUES (:id,:source,:publication,'SPS',:recipient,:rules,:state)
        """), {"id": uuid.uuid4().bytes, "source": source_id,
              "publication": uuid.UUID(publication_id).bytes, "recipient": service.RECIPIENT_KEY,
              "rules": service.RULES_VERSION, "state": prior_state})
    with pytest.raises(service.OperationalFeishuReportError) as failure:
        service.prepare_report(
            engine, source_id=source_id, publication_id=publication_id, region_code="SPS",
            recipient_key=service.RECIPIENT_KEY, actor_user_id=71, staging_directory=tmp_path,
            settings=_enabled_settings(source_id),
        )
    assert failure.value.code == "operational_feishu_delivery_unresolved"


@pytest.mark.parametrize("source_id", ["collection_rate", "pdd_collection_failure"])
def test_prepare_after_sent_uses_new_report_id_and_preserves_sent_history(source_id: str, tmp_path: Path):
    from sqlalchemy import text

    engine, publication_id, _ = _report_engine(source_id)
    prior_id = uuid.uuid4()
    with engine.begin() as connection:
        connection.execute(text("""
            INSERT INTO operational_feishu_reports
              (report_id,source_id,publication_id,region_key,recipient_key,rules_version,report_state)
            VALUES (:id,:source,:publication,'SPS',:recipient,:rules,'SENT')
        """), {"id": prior_id.bytes, "source": source_id,
              "publication": uuid.UUID(publication_id).bytes, "recipient": service.RECIPIENT_KEY,
              "rules": service.RULES_VERSION})
    prepared = service.prepare_report(
        engine, source_id=source_id, publication_id=publication_id, region_code="SPS",
        recipient_key=service.RECIPIENT_KEY, actor_user_id=71, staging_directory=tmp_path,
        settings=_enabled_settings(source_id),
    )
    assert uuid.UUID(prepared["reportId"]) != prior_id
    assert prepared["previouslySent"] is True
    with engine.connect() as connection:
        rows = connection.execute(text(
            "SELECT report_id,report_state FROM operational_feishu_reports ORDER BY report_id"
        )).all()
    assert len(rows) == 2
    assert sum(row.report_state == "SENT" for row in rows) == 1
    assert sum(row.report_state == "PREPARED" for row in rows) == 1


@pytest.mark.parametrize("source_id", ["collection_rate", "pdd_collection_failure"])
def test_send_routes_all_three_stages_to_the_server_allowlisted_regional_chat_id(tmp_path: Path, monkeypatch, source_id: str):
    from sqlalchemy import text

    engine, publication_id, _ = _report_engine(source_id)
    with engine.begin() as connection:
        connection.execute(text("UPDATE base_mapping_entries SET regional_name='MG'"))
    fake = _FakeFeishuClient()
    monkeypatch.setattr(service.httpx, "Client", lambda **_kwargs: fake)
    settings = _enabled_settings(source_id).model_copy(update={"feishu_report_recipient_mg_enabled": True})
    prepared = service.prepare_report(
        engine, source_id=source_id, publication_id=publication_id, region_code="MG",
        recipient_key="MG_COLLECTION_GROUP", actor_user_id=71, staging_directory=tmp_path,
        settings=settings,
    )
    with engine.connect() as connection:
        audit_destination = connection.execute(text("""
            SELECT recipient_key,recipient_email,recipient_type,recipient_identifier,recipient_name
            FROM operational_feishu_reports WHERE report_id=:report_id
        """), {"report_id": uuid.UUID(prepared["reportId"]).bytes}).one()
    assert audit_destination == (
        "MG_COLLECTION_GROUP", None, "chat_id", "oc_d3ef49bb0698066065c5556ce484c6ed",
        "MG 揽收沟通群 · Coordenação de Coleta First Mile",
    )
    service.confirm_report(
        engine, report_id=prepared["reportId"], actor_user_id=71, analysis_text=prepared["analysis"],
        source_id=source_id, settings=settings,
    )
    result = service.send_report(
        engine, source_id=source_id, report_id=prepared["reportId"], actor_user_id=71,
        staging_directory=tmp_path, settings=settings,
    )
    destinations = [kwargs for url, kwargs in fake.calls if url == service.MESSAGE_URL]
    assert result["state"] == "SENT"
    assert len(destinations) == 3
    assert all(call["params"] == {"receive_id_type": "chat_id"} for call in destinations)
    assert all(call["json"]["receive_id"] == "oc_d3ef49bb0698066065c5556ce484c6ed" for call in destinations)


@pytest.mark.parametrize("source_id", ["collection_rate", "pdd_collection_failure"])
def test_sr_group_options_require_sr_in_the_publications_pinned_mapping(source_id: str, monkeypatch):
    from sqlalchemy import text

    engine, publication_id, sr_map_id = _report_engine(source_id)
    other_map_id = uuid.uuid4()
    with engine.begin() as connection:
        connection.execute(text("INSERT INTO base_mapping_entries VALUES (:map,'003','SR','Base SR','RM SR')"), {"map": other_map_id.bytes})

    settings = _enabled_settings(source_id).model_copy(update={"feishu_report_recipient_sr_enabled": True})
    with pytest.raises(service.OperationalFeishuReportError) as missing_sr:
        service.get_options(engine, source_id=source_id, publication_id=publication_id, region_code="SR", settings=settings)
    assert missing_sr.value.code == "operational_feishu_region_invalid"

    with engine.begin() as connection:
        connection.execute(text("INSERT INTO base_mapping_entries VALUES (:map,'003','SR','Base SR','RM SR')"), {"map": sr_map_id.bytes})
    mapped = service.get_options(engine, source_id=source_id, publication_id=publication_id, region_code="SR", settings=settings)
    assert mapped["mapVersionId"] == str(sr_map_id)
    assert next(item for item in mapped["recipients"] if item["key"] == "SR_COLLECTION_GROUP")["available"] is True

    national = service.get_options(engine, source_id=source_id, publication_id=publication_id, settings=settings)
    assert next(item for item in national["recipients"] if item["key"] == "SR_COLLECTION_GROUP")["available"] is False

    disabled = settings.model_copy(update={"feishu_report_recipient_sr_enabled": False})
    scoped_but_disabled = service.get_options(
        engine, source_id=source_id, publication_id=publication_id, region_code="SR", settings=disabled,
    )
    assert next(item for item in scoped_but_disabled["recipients"] if item["key"] == "SR_COLLECTION_GROUP")["available"] is False


def test_send_rejects_modified_artifact_after_confirmation(tmp_path: Path, monkeypatch):
    source_id = "collection_rate"
    engine, publication_id, _ = _report_engine(source_id)
    fake = _FakeFeishuClient()
    monkeypatch.setattr(service.httpx, "Client", lambda **_kwargs: fake)
    prepared = service.prepare_report(
        engine, source_id=source_id, publication_id=publication_id, region_code="SPS",
        recipient_key=service.RECIPIENT_KEY, actor_user_id=71, staging_directory=tmp_path,
        settings=_enabled_settings(source_id),
    )
    service.confirm_report(engine, report_id=prepared["reportId"], actor_user_id=71,
                           analysis_text=prepared["analysis"], source_id=source_id,
                           settings=_enabled_settings(source_id))
    artifact = tmp_path / ".operational-feishu-reports" / f"{uuid.UUID(prepared['reportId']).hex}.xlsx"
    artifact.write_bytes(artifact.read_bytes() + b"tampered")
    with pytest.raises(service.OperationalFeishuReportError) as changed:
        service.send_report(
            engine, source_id=source_id, report_id=prepared["reportId"], actor_user_id=71,
            staging_directory=tmp_path, settings=_enabled_settings(source_id),
        )
    assert changed.value.code == "operational_feishu_artifacts_changed"
    assert fake.calls == []


@pytest.mark.parametrize("ambiguous", [False, True], ids=["explicit-rejection", "timeout-ambiguous"])
def test_partial_or_ambiguous_delivery_is_not_retried_automatically(tmp_path: Path, monkeypatch, ambiguous: bool):
    source_id = "collection_rate"
    engine, publication_id, _ = _report_engine(source_id)
    fake = _FakeFeishuClient(fail_url=service.FILES_URL if not ambiguous else service.MESSAGE_URL,
                            ambiguous_message=ambiguous)
    monkeypatch.setattr(service.httpx, "Client", lambda **_kwargs: fake)
    prepared = service.prepare_report(
        engine, source_id=source_id, publication_id=publication_id, region_code="SPS",
        recipient_key=service.RECIPIENT_KEY, actor_user_id=71, staging_directory=tmp_path,
        settings=_enabled_settings(source_id),
    )
    service.confirm_report(engine, report_id=prepared["reportId"], actor_user_id=71,
                           analysis_text=prepared["analysis"], source_id=source_id,
                           settings=_enabled_settings(source_id))
    expected_state = "PENDING" if ambiguous else "PARTIAL"
    expected_error = "operational_feishu_delivery_pending" if ambiguous else "operational_feishu_partial_delivery"
    with pytest.raises(service.OperationalFeishuReportError) as failure:
        service.send_report(
            engine, source_id=source_id, report_id=prepared["reportId"], actor_user_id=71,
            staging_directory=tmp_path, settings=_enabled_settings(source_id),
        )
    assert failure.value.code == expected_error
    with engine.connect() as connection:
        state = connection.execute(
            __import__("sqlalchemy").text("SELECT report_state FROM operational_feishu_reports WHERE report_id=:id"),
            {"id": uuid.UUID(prepared["reportId"]).bytes},
        ).scalar_one()
    assert state == expected_state
    call_count = len(fake.calls)
    with pytest.raises(service.OperationalFeishuReportError):
        service.send_report(
            engine, source_id=source_id, report_id=prepared["reportId"], actor_user_id=71,
            staging_directory=tmp_path, settings=_enabled_settings(source_id),
        )
    assert len(fake.calls) == call_count
