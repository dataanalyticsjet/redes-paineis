"""Disposable MySQL 8/InnoDB integration checks for the T4/T5 logical batch.

This module intentionally refuses non-loopback hosts and any database other than
the workflow's disposable ``redes_paineis_dados`` service database.
"""
from __future__ import annotations

import hashlib
import io
import os
import threading
import time
import uuid
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine, text
from sqlalchemy.engine import make_url

from app.services.data_foundation.contracts import contract_sha256, get_contract
from app.services.data_foundation.repository import _switch_t4t5_heads_if_newer
from app.services.data_foundation.service_principals import principal_can_publish
from app.services.data_foundation import t4t5_batches
from app.services.data_foundation.t4t5_batches import (
    SOURCE,
    VERSION,
    create_batch,
    normalized_manifest,
    publish_batch,
    publish_batch_safely,
    queue_batch_publication,
    queue_batch_validation,
    receive_part,
    validate_batch,
)
from app.services.data_foundation.xlsx_parser import PreviewSummary


pytestmark = pytest.mark.skipif(
    os.getenv("T4T5_MYSQL_TESTS") != "1",
    reason="requires the disposable MySQL 8 service from the dedicated Actions workflow",
)

PRINCIPAL = "windows-newbi-publisher"
MAP_ID = b"\x31" * 16
ALT_MAP_ID = b"\x32" * 16
REGIONS_DATES = ("2026-10-01", "2026-10-02")


@pytest.fixture(scope="module")
def mysql_engine():
    url_text = os.getenv("T4T5_MYSQL_TEST_URL", "")
    url = make_url(url_text)
    if url.drivername != "mysql+pymysql" or url.host not in {"127.0.0.1", "localhost", "::1"}:
        pytest.fail("T4T5 MySQL integration requires a loopback MySQL+pymysql URL")
    if url.database != "redes_paineis_dados":
        pytest.fail("T4T5 MySQL integration refuses a database other than redes_paineis_dados")
    if not url.password or url.password != "disposable-test-root":
        pytest.fail("T4T5 MySQL integration requires the workflow-only disposable credential")
    marker = os.getenv("T4T5_MYSQL_TEST_MARKER", "")
    if not marker.startswith("github-actions-"):
        pytest.fail("T4T5 MySQL integration requires the per-run GitHub Actions test marker")

    engine = create_engine(url, pool_size=8, max_overflow=0, pool_pre_ping=True)
    with engine.connect() as connection:
        version = str(connection.execute(text("SELECT VERSION()")).scalar_one())
        isolation = str(connection.execute(text("SELECT @@transaction_isolation")).scalar_one()).upper()
        if not version.startswith("8."):
            pytest.fail(f"expected MySQL 8, found {version}")
        if isolation != "REPEATABLE-READ":
            pytest.fail(f"expected REPEATABLE-READ, found {isolation}")
        actual_marker = connection.execute(text(
            "SELECT marker_value FROM ci_t4t5_homologation_marker"
        )).scalar_one_or_none()
        if actual_marker != marker:
            pytest.fail("database does not carry this GitHub Actions run's disposable marker")
        engine._t4t5_mysql_version = version  # type: ignore[attr-defined]
    yield engine
    engine.dispose()


def _seed_mapping(engine, map_id: bytes, *, version_no: int, code: str = "000001") -> None:
    source_job = uuid.uuid4().bytes
    with engine.begin() as connection:
        connection.execute(text("""
            INSERT INTO data_source_registry
              (source_id,display_name,description,source_state,history_policy)
            VALUES ('base_mapping_official','Synthetic mapping','CI fixture','ACTIVE','PENDING')
            ON DUPLICATE KEY UPDATE source_id=VALUES(source_id)
        """))
        connection.execute(text("""
            INSERT INTO source_contract_versions
              (source_id,contract_version,contract_state,expected_extension,history_policy,contract_sha256)
            VALUES ('base_mapping_official',1,'APPROVED','xlsx','PENDING',:hash)
            ON DUPLICATE KEY UPDATE source_id=VALUES(source_id)
        """), {"hash": b"c" * 32})
        connection.execute(text("""
            INSERT INTO import_jobs
              (job_id,source_id,contract_version,validation_map_version_id,reused_publication_id,
               submitted_by_user_id,submitted_by_service_principal_id,job_state,original_file_name,
               stored_file_key,file_size_bytes,file_sha256,request_id)
            VALUES (:job,'base_mapping_official',1,NULL,NULL,1,NULL,'PUBLISHED','ci-map.xlsx',
                    'ci-map.xlsx',1,:digest,NULL)
        """), {"job": source_job, "digest": hashlib.sha256(source_job).digest()})
        connection.execute(text("""
            INSERT INTO base_mapping_versions
              (map_version_id,version_no,source_id,source_job_id,contract_version,version_state,
               original_file_name,file_sha256,mapping_row_count,published_by_user_id,published_at)
            VALUES (:map,:version,'base_mapping_official',:job,1,'PUBLISHED','ci-map.xlsx',:digest,1,1,CURRENT_TIMESTAMP(6))
        """), {"map": map_id, "version": version_no, "job": source_job,
                "digest": hashlib.sha256(map_id).digest()})
        connection.execute(text("""
            INSERT INTO base_mapping_entries
              (map_version_id,base_code,base_name,source_row_no)
            VALUES (:map,:code,'Base Test',1)
        """), {"map": map_id, "code": code})
        if version_no == 1:
            connection.execute(text("""
                INSERT INTO base_mapping_current (current_key,map_version_id,changed_by_user_id)
                VALUES ('OFFICIAL',:map,1)
            """), {"map": map_id})


def _set_publication_access(engine, *, enabled: bool) -> None:
    with engine.begin() as connection:
        connection.execute(text("""
            UPDATE data_source_registry SET source_state='ACTIVE'
            WHERE source_id=:source
        """), {"source": SOURCE})
        connection.execute(text("""
            UPDATE source_contract_versions
            SET contract_state='APPROVED',contract_sha256=:hash
            WHERE source_id=:source AND contract_version=:version
        """), {"hash": contract_sha256(get_contract(SOURCE, VERSION)),
                "source": SOURCE, "version": VERSION})
        connection.execute(text("""
            UPDATE data_service_principal_contract_scopes SET is_active=:active
            WHERE service_principal_id=:principal AND source_id=:source AND contract_version=:version
        """), {"active": int(enabled), "principal": PRINCIPAL, "source": SOURCE, "version": VERSION})


def _part_workbook(day: str, row_count: int) -> bytes:
    contract = get_contract(SOURCE, VERSION)
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(contract.sheet_name)
    sheet.append(list(contract.headers))
    for row_no in range(row_count):
        sheet.append([
            date.fromisoformat(day), "SPS", "Base Test", "Synthetic customer", "Synthetic shop",
            f"{row_no:08d}", "Synthetic driver", "E-commerce", 1, 2, 3, 4, 5, 6, 7,
        ])
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _manifest_part(day: str, content: bytes, rows: int, suffix: str = "") -> dict[str, object]:
    return {
        "dataDate": day,
        "fileName": f"t4t5-{day}{suffix}.xlsx",
        "sizeBytes": len(content),
        "sha256": hashlib.sha256(content).hexdigest(),
        "rowCount": rows,
    }


def _create_import_batch(engine, parts: list[dict[str, object]], key: str) -> dict[str, object]:
    return create_batch(engine, principal=PRINCIPAL, idempotency_key=key, raw_parts=parts,
                       map_version=MAP_ID, request_id=f"ci-{key}")


def _seed_competing_publication(engine, version: int, dates: tuple[str, ...], map_id: bytes) -> bytes:
    job_id, publication_id = uuid.uuid4().bytes, uuid.uuid4().bytes
    digest = hashlib.sha256(f"synthetic-publication-{version}".encode()).digest()
    with engine.begin() as connection:
        connection.execute(text("""
            INSERT INTO import_jobs
              (job_id,source_id,contract_version,validation_map_version_id,reused_publication_id,
               submitted_by_user_id,submitted_by_service_principal_id,job_state,original_file_name,
               stored_file_key,file_size_bytes,file_sha256,request_id)
            VALUES (:job,:source,2,:map,NULL,NULL,:principal,'PUBLISHED','ci-candidate.manifest',
                    'ci-candidate',1,:digest,NULL)
        """), {"job": job_id, "source": SOURCE, "map": map_id,
                "principal": PRINCIPAL, "digest": digest})
        connection.execute(text("""
            INSERT INTO source_publications
              (publication_id,source_id,version_no,job_id,contract_version,map_version_id,file_sha256,
               publication_state,history_policy,period_start,period_end,row_count,mapped_row_count,
               unmapped_row_count,missing_identifier_row_count,ambiguous_row_count,published_at)
            VALUES (:pub,:source,:version,:job,2,:map,:digest,'PUBLISHED','DAILY_REPLACE_DATES',
                    :start,:end,2,2,0,0,0,CURRENT_TIMESTAMP(6))
        """), {"pub": publication_id, "source": SOURCE, "version": version, "job": job_id,
                "map": map_id, "digest": digest,
                "start": date.fromisoformat(dates[0]), "end": date.fromisoformat(dates[-1])})
        connection.execute(text("""
            INSERT INTO source_publication_dates (source_id,publication_id,data_date,row_count)
            VALUES (:source,:pub,:day,1)
        """), [{"source": SOURCE, "pub": publication_id, "day": date.fromisoformat(day)} for day in dates])
    return publication_id


def test_reviewed_sql_is_installed_on_mysql8_with_v1_preserved_and_v2_inactive(mysql_engine):
    engine = mysql_engine
    assert engine._t4t5_mysql_version.startswith("8.")  # type: ignore[attr-defined]
    with engine.connect() as connection:
        tables = connection.execute(text("""
            SELECT table_name,engine FROM information_schema.tables
            WHERE table_schema=DATABASE() AND table_name IN
              ('data_import_batches','data_import_batch_parts','import_jobs','source_daily_head',
               'source_publications','source_publication_dates','fact_collection_monitoring_t4_t5')
        """)).all()
        assert {name for name, _engine in tables} == {
            "data_import_batches", "data_import_batch_parts", "import_jobs", "source_daily_head",
            "source_publications", "source_publication_dates", "fact_collection_monitoring_t4_t5",
        }
        assert all(engine_name == "InnoDB" for _name, engine_name in tables)
        versions = connection.execute(text("""
            SELECT contract_version,contract_state,contract_sha256
            FROM source_contract_versions WHERE source_id=:source ORDER BY contract_version
        """), {"source": SOURCE}).all()
        assert [row[0] for row in versions] == [1, 2]
        assert versions[0][1] == "DISCOVERED"
        assert bytes(versions[0][2]) == contract_sha256(get_contract(SOURCE, 1))
        assert versions[1][1] == "DISCOVERED"
        assert bytes(versions[1][2]) == contract_sha256(get_contract(SOURCE, 2))
        column_counts = connection.execute(text("""
            SELECT contract_version,COUNT(*) FROM source_contract_columns
            WHERE source_id=:source GROUP BY contract_version ORDER BY contract_version
        """), {"source": SOURCE}).all()
        assert column_counts == [(1, 15), (2, 15)]
        assert connection.execute(text("""
            SELECT is_active FROM data_service_principal_contract_scopes
            WHERE service_principal_id=:principal AND source_id=:source AND contract_version=2
        """), {"principal": PRINCIPAL, "source": SOURCE}).scalar_one() == 0
        daily_head_check = connection.execute(text("""
            SELECT cc.check_clause FROM information_schema.check_constraints cc
            JOIN information_schema.table_constraints tc
              ON tc.constraint_schema=cc.constraint_schema AND tc.constraint_name=cc.constraint_name
            WHERE tc.constraint_schema=DATABASE() AND tc.table_name='source_daily_head'
              AND tc.constraint_name='chk_daily_head_source'
        """)).scalar_one()
        assert SOURCE in str(daily_head_check)
        fk_names = set(connection.execute(text("""
            SELECT constraint_name FROM information_schema.referential_constraints
            WHERE constraint_schema=DATABASE() AND table_name IN ('data_import_batches','data_import_batch_parts')
        """)).scalars())
        assert {"fk_t4t5_batch_job", "fk_t4t5_batch_contract", "fk_t4t5_batch_principal",
                "fk_t4t5_batch_map", "fk_t4t5_batch_publication", "fk_t4t5_batch_part_job"} <= fk_names


def test_mysql8_logical_batch_publication_resume_concurrency_and_rollback(
    mysql_engine, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
):
    engine = mysql_engine
    _seed_mapping(engine, MAP_ID, version_no=9001)
    # The V2 SQL must not grant use implicitly. Contract/scope checks reject it first.
    with engine.connect() as connection:
        assert not principal_can_publish(connection, PRINCIPAL, SOURCE, VERSION)
    original = _part_workbook(REGIONS_DATES[0], 1)
    with pytest.raises(Exception, match="source_contract_not_approved"):
        _create_import_batch(engine, [_manifest_part(REGIONS_DATES[0], original, 1)], "contract-gate-001")

    _set_publication_access(engine, enabled=True)
    with engine.connect() as connection:
        assert principal_can_publish(connection, PRINCIPAL, SOURCE, VERSION)
    part_rows = (5_000, 5_001)
    payloads = [_part_workbook(day, count) for day, count in zip(REGIONS_DATES, part_rows)]
    parts = [_manifest_part(day, content, count) for day, content, count in zip(REGIONS_DATES, payloads, part_rows)]
    with ThreadPoolExecutor(max_workers=2) as pool:
        created_pair = list(pool.map(
            lambda _n: _create_import_batch(engine, parts, "mysql-integration-batch-001"), range(2)
        ))
    created = created_pair[0]
    job_id = str(created["jobId"])
    assert created["state"] == "RECEIVING"
    assert created_pair[1]["jobId"] == job_id
    assert created["mapVersionId"] == str(uuid.UUID(bytes=MAP_ID))
    assert _create_import_batch(engine, parts, "mysql-integration-batch-001")["jobId"] == job_id
    changed_manifest = [dict(parts[0], sha256="f" * 64), parts[1]]
    with pytest.raises(Exception, match="t4t5_batch_idempotency_conflict"):
        _create_import_batch(engine, changed_manifest, "mysql-integration-batch-001")

    for day, content in zip(REGIONS_DATES, payloads):
        receive_part(engine, job_id=job_id, principal=PRINCIPAL, day=day,
                     directory=tmp_path, stream=io.BytesIO(content))
    queued = queue_batch_validation(engine, job_id=job_id, principal=PRINCIPAL)
    assert queued["startWorker"] is True
    with engine.begin() as connection:
        connection.execute(text("""
            UPDATE data_import_batches SET updated_at=CURRENT_TIMESTAMP(6)-INTERVAL 3 HOUR,
              progress_rows=4321 WHERE job_id=UUID_TO_BIN(:job,0)
        """), {"job": job_id})
    resumed = queue_batch_validation(engine, job_id=job_id, principal=PRINCIPAL)
    assert resumed["startWorker"] is True
    with engine.connect() as connection:
        state, rows = connection.execute(text("""
            SELECT batch_state,progress_rows FROM data_import_batches WHERE job_id=UUID_TO_BIN(:job,0)
        """), {"job": job_id}).one()
        assert (state, rows) == ("VALIDATING", 0)
    validate_batch(engine, job_id=job_id, principal=PRINCIPAL, directory=tmp_path)
    publish_queue = queue_batch_publication(engine, job_id=job_id, principal=PRINCIPAL,
                                            idempotency_key="mysql-publish-batch-001")
    assert publish_queue["startWorker"] is True
    publish_batch(engine, job_id=job_id, principal=PRINCIPAL, directory=tmp_path,
                  idempotency_key="mysql-publish-batch-001", request_id="mysql-publish-001")
    with engine.connect() as connection:
        batch = connection.execute(text("""
            SELECT batch.batch_state,batch.publication_id,batch.map_version_id,
                   publication.version_no,publication.row_count
            FROM data_import_batches batch JOIN source_publications publication
              ON publication.publication_id=batch.publication_id
            WHERE batch.job_id=UUID_TO_BIN(:job,0)
        """), {"job": job_id}).one()
        assert batch[0] == "PUBLISHED"
        assert bytes(batch[2]) == MAP_ID
        assert batch[3] == 1 and batch[4] == sum(part_rows)
        fact_summary = connection.execute(text("""
            SELECT COUNT(*),COUNT(DISTINCT map_version_id),COUNT(DISTINCT data_date)
            FROM fact_collection_monitoring_t4_t5 WHERE publication_id=:publication
        """), {"publication": batch[1]}).one()
        assert fact_summary == (sum(part_rows), 1, 2)
        heads = connection.execute(text("""
            SELECT COUNT(DISTINCT publication_id),COUNT(*) FROM source_daily_head
            WHERE source_id=:source AND data_date BETWEEN '2026-10-01' AND '2026-10-02'
        """), {"source": SOURCE}).one()
        assert heads == (1, 2)

    # A newly official DoomsDay version cannot be substituted after the batch
    # pinned its map; validation fails closed and does not create a publication.
    _seed_mapping(engine, ALT_MAP_ID, version_no=9002)
    next_day = "2026-10-03"
    next_payload = _part_workbook(next_day, 1)
    map_pinned_batch = _create_import_batch(
        engine, [_manifest_part(next_day, next_payload, 1)], "mysql-map-pinned-002"
    )
    pinned_job_id = str(map_pinned_batch["jobId"])
    receive_part(engine, job_id=pinned_job_id, principal=PRINCIPAL, day=next_day,
                 directory=tmp_path, stream=io.BytesIO(next_payload))
    queue_batch_validation(engine, job_id=pinned_job_id, principal=PRINCIPAL)
    with engine.begin() as connection:
        connection.execute(text("""
            UPDATE base_mapping_current SET map_version_id=:map,changed_by_user_id=1
            WHERE current_key='OFFICIAL'
        """), {"map": ALT_MAP_ID})
    with pytest.raises(Exception, match="official_mapping_changed_repreview"):
        validate_batch(engine, job_id=pinned_job_id, principal=PRINCIPAL, directory=tmp_path)
    with engine.begin() as connection:
        connection.execute(text("""
            UPDATE base_mapping_current SET map_version_id=:map,changed_by_user_id=1
            WHERE current_key='OFFICIAL'
        """), {"map": MAP_ID})
    with engine.connect() as connection:
        state, pinned_map = connection.execute(text("""
            SELECT batch.batch_state,batch.map_version_id FROM data_import_batches batch
            WHERE batch.job_id=UUID_TO_BIN(:job,0)
        """), {"job": pinned_job_id}).one()
        assert state == "INVALID" and bytes(pinned_map) == MAP_ID
        assert connection.execute(text("""
            SELECT COUNT(*) FROM source_publications WHERE job_id=UUID_TO_BIN(:job,0)
        """), {"job": pinned_job_id}).scalar_one() == 0

    # Publish two competing versions over the same dates with independent InnoDB sessions.
    lower_pub = _seed_competing_publication(engine, 2, REGIONS_DATES, MAP_ID)
    higher_pub = _seed_competing_publication(engine, 3, REGIONS_DATES, MAP_ID)
    low_locked, let_low_commit, high_started = threading.Event(), threading.Event(), threading.Event()
    worker_errors: list[BaseException] = []

    def switch(publication_id: bytes, version: int, pause: bool = False) -> None:
        try:
            with engine.begin() as connection:
                summary = PreviewSummary(source_id=SOURCE, file_sha256="0" * 64, contract_version=VERSION,
                                        mapping_version_available=True, map_version_id=MAP_ID,
                                        per_date_counts={day: 1 for day in REGIONS_DATES})
                _switch_t4t5_heads_if_newer(connection, publication_id=publication_id,
                                           version_no=version, summary=summary)
                if pause:
                    low_locked.set()
                    if not let_low_commit.wait(10):
                        raise AssertionError("test coordinator timed out")
        except BaseException as error:  # bubble worker exceptions to pytest thread
            worker_errors.append(error)

    with ThreadPoolExecutor(max_workers=2) as pool:
        low_future = pool.submit(switch, lower_pub, 2, True)
        assert low_locked.wait(10)
        def high_switch() -> None:
            high_started.set()
            switch(higher_pub, 3)
        high_future = pool.submit(high_switch)
        assert high_started.wait(10)
        time.sleep(0.2)  # give the second connection time to block on the source row lock
        assert not high_future.done()
        with engine.connect() as observer:
            observed = observer.execute(text("""
                SELECT COUNT(DISTINCT head.publication_id),COUNT(*)
                FROM source_daily_head head WHERE head.source_id=:source
                  AND head.data_date BETWEEN '2026-10-01' AND '2026-10-02'
            """), {"source": SOURCE}).one()
            # Low publication is uncommitted; a separate REPEATABLE READ session sees the old set.
            assert observed == (1, 2)
            old_publication = observer.execute(text("""
                SELECT publication_id FROM source_daily_head
                WHERE source_id=:source AND data_date='2026-10-01'
            """), {"source": SOURCE}).scalar_one()
        assert bytes(old_publication) == bytes(batch[1])
        let_low_commit.set()
        low_future.result(timeout=15)
        high_future.result(timeout=15)
    assert not worker_errors
    with engine.connect() as connection:
        assert set(connection.execute(text("""
            SELECT DISTINCT publication_id FROM source_daily_head
            WHERE source_id=:source AND data_date BETWEEN '2026-10-01' AND '2026-10-02'
        """), {"source": SOURCE}).scalars()) == {higher_pub}

    # A failure while receiving a later part leaves earlier verified bytes
    # resumable, without creating a publication or exposing a partial date set.
    partial_days = ("2026-12-01", "2026-12-02")
    partial_payloads = [_part_workbook(day, 1) for day in partial_days]
    partial_parts = [_manifest_part(day, content, 1) for day, content in zip(partial_days, partial_payloads)]
    partial = _create_import_batch(engine, partial_parts, "mysql-partial-receive-003")
    partial_job_id = str(partial["jobId"])
    receive_part(engine, job_id=partial_job_id, principal=PRINCIPAL, day=partial_days[0],
                 directory=tmp_path, stream=io.BytesIO(partial_payloads[0]))
    altered = bytearray(partial_payloads[1])
    altered[-1] ^= 1
    with pytest.raises(Exception, match="t4t5_batch_part_integrity_mismatch"):
        receive_part(engine, job_id=partial_job_id, principal=PRINCIPAL, day=partial_days[1],
                     directory=tmp_path, stream=io.BytesIO(altered))
    with engine.connect() as connection:
        part_states = connection.execute(text("""
            SELECT upload_state FROM data_import_batch_parts
            WHERE job_id=UUID_TO_BIN(:job,0) ORDER BY data_date
        """), {"job": partial_job_id}).scalars().all()
        assert part_states == ["RECEIVED", "MISSING"]
        assert connection.execute(text("""
            SELECT COUNT(*) FROM source_publications WHERE job_id=UUID_TO_BIN(:job,0)
        """), {"job": partial_job_id}).scalar_one() == 0
        assert connection.execute(text("""
            SELECT COUNT(*) FROM source_daily_head
            WHERE source_id=:source AND data_date BETWEEN '2026-12-01' AND '2026-12-02'
        """), {"source": SOURCE}).scalar_one() == 0
    receive_part(engine, job_id=partial_job_id, principal=PRINCIPAL, day=partial_days[1],
                 directory=tmp_path, stream=io.BytesIO(partial_payloads[1]))
    queue_batch_validation(engine, job_id=partial_job_id, principal=PRINCIPAL)
    validate_batch(engine, job_id=partial_job_id, principal=PRINCIPAL, directory=tmp_path)
    publish_key = "mysql-resume-publish-003"
    assert queue_batch_publication(engine, job_id=partial_job_id, principal=PRINCIPAL,
                                   idempotency_key=publish_key)["startWorker"] is True
    with engine.begin() as connection:
        connection.execute(text("""
            UPDATE data_import_batches SET updated_at=CURRENT_TIMESTAMP(6)-INTERVAL 3 HOUR
            WHERE job_id=UUID_TO_BIN(:job,0)
        """), {"job": partial_job_id})
    assert queue_batch_publication(engine, job_id=partial_job_id, principal=PRINCIPAL,
                                   idempotency_key=publish_key)["startWorker"] is True
    publish_batch(engine, job_id=partial_job_id, principal=PRINCIPAL, directory=tmp_path,
                  idempotency_key=publish_key, request_id="mysql-resume-publish-003")
    with engine.connect() as connection:
        assert connection.execute(text("""
            SELECT batch_state FROM data_import_batches WHERE job_id=UUID_TO_BIN(:job,0)
        """), {"job": partial_job_id}).scalar_one() == "PUBLISHED"
        assert connection.execute(text("""
            SELECT COUNT(*) FROM source_daily_head
            WHERE source_id=:source AND data_date BETWEEN '2026-12-01' AND '2026-12-02'
        """), {"source": SOURCE}).scalar_one() == 2

    # Inject a failure after the real final transaction has switched both
    # pointers but before commit. Rollback must retain the former publication.
    failed_payloads = [_part_workbook(day, 1) for day in partial_days]
    failed_parts = [_manifest_part(day, content, 1) for day, content in zip(partial_days, failed_payloads)]
    failed = _create_import_batch(engine, failed_parts, "mysql-precommit-failure-004")
    failed_job_id = str(failed["jobId"])
    prior_heads = {}
    with engine.connect() as connection:
        prior_heads = dict(connection.execute(text("""
            SELECT data_date,publication_id FROM source_daily_head
            WHERE source_id=:source AND data_date BETWEEN '2026-12-01' AND '2026-12-02'
        """), {"source": SOURCE}).all())
        assert len(prior_heads) == 2
    for day, content in zip(partial_days, failed_payloads):
        receive_part(engine, job_id=failed_job_id, principal=PRINCIPAL, day=day,
                     directory=tmp_path, stream=io.BytesIO(content))
    queue_batch_validation(engine, job_id=failed_job_id, principal=PRINCIPAL)
    validate_batch(engine, job_id=failed_job_id, principal=PRINCIPAL, directory=tmp_path)
    failed_publish_key = "mysql-precommit-failure-publish-004"
    queue_batch_publication(engine, job_id=failed_job_id, principal=PRINCIPAL,
                            idempotency_key=failed_publish_key)
    original_switch = t4t5_batches._switch_t4t5_heads_if_newer

    def fail_after_head_updates(connection, **kwargs):
        original_switch(connection, **kwargs)
        raise RuntimeError("synthetic failure immediately before final commit")

    monkeypatch.setattr(t4t5_batches, "_switch_t4t5_heads_if_newer", fail_after_head_updates)
    publish_batch_safely(engine, job_id=failed_job_id, principal=PRINCIPAL, directory=tmp_path,
                         idempotency_key=failed_publish_key, request_id="mysql-precommit-failure-004")
    monkeypatch.setattr(t4t5_batches, "_switch_t4t5_heads_if_newer", original_switch)
    with engine.connect() as connection:
        after_heads = dict(connection.execute(text("""
            SELECT data_date,publication_id FROM source_daily_head
            WHERE source_id=:source AND data_date BETWEEN '2026-12-01' AND '2026-12-02'
        """), {"source": SOURCE}).all())
        failure_state = connection.execute(text("""
            SELECT batch.batch_state,publication.publication_state
            FROM data_import_batches batch JOIN source_publications publication
              ON publication.job_id=batch.job_id
            WHERE batch.job_id=UUID_TO_BIN(:job,0)
        """), {"job": failed_job_id}).one()
        assert after_heads == prior_heads
        assert failure_state == ("FAILED", "PREPARED")

    # Failure after the first date update must roll back the entire head switch.
    bad_summary = PreviewSummary(source_id=SOURCE, file_sha256="0" * 64, contract_version=VERSION,
                                mapping_version_available=True, map_version_id=MAP_ID,
                                per_date_counts={"2026-10-01": 1, "not-a-date": 1})
    with pytest.raises(ValueError):
        with engine.begin() as connection:
            _switch_t4t5_heads_if_newer(connection, publication_id=lower_pub,
                                       version_no=4, summary=bad_summary)
    with engine.connect() as connection:
        assert set(connection.execute(text("""
            SELECT DISTINCT publication_id FROM source_daily_head
            WHERE source_id=:source AND data_date BETWEEN '2026-10-01' AND '2026-10-02'
        """), {"source": SOURCE}).scalars()) == {higher_pub}

    # Over-one-million logical row-count manifests stay small: InnoDB stores
    # only 26 date-partition metadata rows during this proportional limit check.
    large_parts = [{"dataDate": f"2026-11-{day:02}", "fileName": f"large-{day:02}.xlsx",
                    "sizeBytes": 1, "sha256": hashlib.sha256(f"p{day}".encode()).hexdigest(),
                    "rowCount": 40_000} for day in range(1, 27)]
    large = _create_import_batch(engine, large_parts, "mysql-large-manifest-001")
    assert large["expectedTotalRows"] == 1_040_000
    assert large["expectedPartCount"] == 26
