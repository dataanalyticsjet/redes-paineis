from __future__ import annotations

import hashlib
import io
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
import uuid

import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine, text

from app.services.data_foundation.contracts import contract_sha256, get_contract
from app.services.data_foundation.xlsx_parser import MappingResolver, parse_workbook
from app.services.data_foundation.t4t5_batches import (
    MAX_BATCH_ROWS, MAX_PART_BYTES, SOURCE, VERSION, _touch_validation_progress,
    normalized_manifest, queue_batch_validation, validate_part_summary, save_part_stream,
)


def _part(day: str, rows: int = 0, payload: bytes = b"part") -> dict[str, object]:
    return {"dataDate": day, "fileName": f"t4t5-{day}.xlsx", "sizeBytes": len(payload),
            "sha256": hashlib.sha256(payload).hexdigest(), "rowCount": rows}


def test_v2_contract_keeps_the_same_15_columns_and_v1_default():
    v1 = get_contract(SOURCE, 1)
    v2 = get_contract(SOURCE, 2)
    assert get_contract(SOURCE).version == 1
    assert len(v1.columns) == len(v2.columns) == 15
    assert v1.headers == v2.headers
    assert v1.history_policy == v2.history_policy == "DAILY_REPLACE_DATES"
    assert contract_sha256(v1) != contract_sha256(v2)
    assert contract_sha256(v2).hex() == "49bccfa5e597fa7c26928d10e445afd198a1c099406ae62cdaf173215da4e547"


def test_manifest_accepts_monthly_logical_batch_over_one_million_rows():
    parts = [_part(f"2026-10-{day:02}", 40_000) for day in range(1, 27)]
    manifest, digest = normalized_manifest(parts)
    assert sum(part["rowCount"] for part in manifest["parts"]) == 1_040_000
    assert len(manifest["parts"]) == 26
    assert len(digest) == 32
    assert manifest["contractVersion"] == 2


def test_manifest_rejects_missing_or_overlapping_daily_partitions():
    with pytest.raises(Exception, match="t4t5_batch_dates_not_contiguous"):
        normalized_manifest([_part("2026-10-01"), _part("2026-10-03")])
    with pytest.raises(Exception, match="t4t5_batch_manifest_invalid"):
        normalized_manifest([_part("2026-10-01"), _part("2026-10-01")])


def test_manifest_limits_each_part_and_total_explicitly():
    huge = _part("2026-10-01")
    huge["sizeBytes"] = MAX_PART_BYTES + 1
    with pytest.raises(Exception, match="t4t5_batch_limit_exceeded"):
        normalized_manifest([huge])
    with pytest.raises(Exception, match="t4t5_batch_total_limit_exceeded"):
        normalized_manifest([_part(f"2026-10-{day:02}", 80_000) for day in range(1, 27)])


def test_part_upload_resumes_only_when_size_and_digest_match(tmp_path: Path):
    content = b"validated daily workbook bytes"
    digest = hashlib.sha256(content).hexdigest()
    first = save_part_stream(tmp_path, "t4t5-batch-a-2026-10-01.xlsx", io.BytesIO(content), len(content), digest)
    assert first == (len(content), digest)
    with pytest.raises(Exception, match="t4t5_batch_part_integrity_mismatch"):
        save_part_stream(tmp_path, "t4t5-batch-a-2026-10-01.xlsx", io.BytesIO(b"changed"), len(content), digest)
    assert (tmp_path / "t4t5-batch-a-2026-10-01.xlsx").read_bytes() == content


def test_v2_part_parser_preserves_fifteen_columns_and_same_day_duplicates(tmp_path: Path):
    contract = get_contract(SOURCE, VERSION)
    path = tmp_path / "daily.xlsx"
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Dados")
    sheet.append(list(contract.headers))
    row = [date(2026, 10, 1), "SPS", "Base A", "Cliente", "Loja", "0000123", "Motorista",
           "E-commerce", 1, 2, 3, 4, 5, 6, 7]
    sheet.append(row)
    sheet.append(row)
    workbook.save(path)
    resolver = MappingResolver(map_version_id=b"m" * 16, exact_base_name_candidates={"Base A": ("BASE-1",)})
    parsed = parse_workbook(path, contract, resolver)
    assert parsed.input_row_count == 2
    assert parsed.valid_row_count == 2
    assert parsed.exact_duplicate_count == 1
    assert parsed.per_date_counts == {"2026-10-01": 2}
    assert parsed.can_publish
    assert parsed.map_version_id == b"m" * 16


def test_batch_rejects_rows_from_another_operational_date(tmp_path: Path):
    contract = get_contract(SOURCE, VERSION)
    path = tmp_path / "daily.xlsx"
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Dados")
    sheet.append(list(contract.headers))
    sheet.append([date(2026, 10, 2), "SPS", "Base A", "Cliente", "Loja", "123", "Motorista",
                  "E-commerce", 1, 2, 3, 4, 5, 6, 7])
    workbook.save(path)
    parsed = parse_workbook(path, contract, MappingResolver(map_version_id=b"m" * 16))
    assert parsed.can_publish
    with pytest.raises(Exception, match="t4t5_batch_part_date_mismatch"):
        validate_part_summary(parsed, "2026-10-01", 1)


def _validation_engine(state: str, updated_at: datetime):
    engine = create_engine("sqlite://")
    job_id = uuid.uuid4()
    with engine.begin() as connection:
        connection.execute(text("""
            CREATE TABLE import_jobs (
              job_id BLOB PRIMARY KEY, source_id TEXT, contract_version INTEGER,
              validation_map_version_id BLOB, reused_publication_id BLOB, job_state TEXT,
              original_file_name TEXT, stored_file_key TEXT, file_size_bytes INTEGER, file_sha256 BLOB,
              input_row_count INTEGER, valid_row_count INTEGER, error_count INTEGER,
              submitted_by_user_id INTEGER, submitted_by_service_principal_id TEXT
            )
        """))
        connection.execute(text("""
            CREATE TABLE data_import_batches (
              job_id BLOB PRIMARY KEY, batch_state TEXT, updated_at TEXT,
              expected_total_rows INTEGER, expected_part_count INTEGER, validation_error_code TEXT,
              progress_rows INTEGER DEFAULT 0
            )
        """))
        connection.execute(text("""
            CREATE TABLE data_import_batch_parts (
              job_id BLOB, data_date TEXT, part_no INTEGER, original_file_name TEXT,
              expected_size_bytes INTEGER, expected_sha256 BLOB, expected_row_count INTEGER,
              stored_file_key TEXT, upload_state TEXT, validation_state TEXT,
              valid_row_count INTEGER, mapped_row_count INTEGER, unmapped_row_count INTEGER
            )
        """))
        connection.execute(text("""
            INSERT INTO import_jobs
              (job_id,source_id,contract_version,validation_map_version_id,job_state,
               submitted_by_service_principal_id)
            VALUES (:job,:source,:version,:map,'PARSING','test-principal')
        """), {"job": job_id.bytes, "source": SOURCE, "version": VERSION, "map": b"m" * 16})
        connection.execute(text("""
            INSERT INTO data_import_batches
              (job_id,batch_state,updated_at,expected_total_rows,expected_part_count)
            VALUES (:job,:state,:updated,1,1)
        """), {"job": job_id.bytes, "state": state, "updated": updated_at.strftime("%Y-%m-%d %H:%M:%S")})
        connection.execute(text("""
            INSERT INTO data_import_batch_parts
              (job_id,data_date,part_no,original_file_name,expected_size_bytes,expected_sha256,
               expected_row_count,stored_file_key,upload_state,validation_state)
            VALUES (:job,'2026-10-01',1,'part.xlsx',10,:digest,1,'part.xlsx','RECEIVED','PENDING')
        """), {"job": job_id.bytes, "digest": b"h" * 32})
    return engine, str(job_id)


def test_validation_still_running_is_not_queued_twice():
    engine, job_id = _validation_engine("VALIDATING", datetime.now(timezone.utc).replace(tzinfo=None))
    try:
        result = queue_batch_validation(engine, job_id=job_id, principal="test-principal")
        assert result["state"] == "VALIDATING"
        assert result["startWorker"] is False
    finally:
        engine.dispose()


def test_stale_validation_can_be_resumed_after_worker_interruption():
    engine, job_id = _validation_engine("VALIDATING", datetime.now(timezone.utc).replace(tzinfo=None) - timedelta(hours=3))
    try:
        result = queue_batch_validation(engine, job_id=job_id, principal="test-principal")
        assert result["state"] == "VALIDATING"
        assert result["startWorker"] is True
        with engine.connect() as connection:
            assert connection.execute(text("SELECT job_state FROM import_jobs WHERE job_id=:job"),
                                     {"job": uuid.UUID(job_id).bytes}).scalar_one() == "PARSING"
            assert connection.execute(text("SELECT julianday('now')-julianday(updated_at)<1.0/24 FROM data_import_batches WHERE job_id=:job"),
                                     {"job": uuid.UUID(job_id).bytes}).scalar_one()
    finally:
        engine.dispose()


def test_validation_progress_heartbeat_prevents_stale_reclaim():
    engine, job_id = _validation_engine("VALIDATING", datetime.now(timezone.utc).replace(tzinfo=None) - timedelta(hours=3))
    try:
        _touch_validation_progress(engine, uuid.UUID(job_id).bytes, 12_345)
        with engine.connect() as connection:
            updated_at = connection.execute(text("SELECT updated_at FROM data_import_batches WHERE job_id=:job"),
                                            {"job": uuid.UUID(job_id).bytes}).scalar_one()
            assert connection.execute(text("SELECT progress_rows FROM data_import_batches WHERE job_id=:job"),
                                     {"job": uuid.UUID(job_id).bytes}).scalar_one() == 12_345
            assert connection.execute(text("SELECT julianday('now')-julianday(:updated)<2.0/24"),
                                      {"updated": updated_at}).scalar_one()
    finally:
        engine.dispose()
