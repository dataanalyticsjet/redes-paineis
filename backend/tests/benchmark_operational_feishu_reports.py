"""Opt-in synthetic load benchmark for operational Feishu report generation.

Run from ``backend`` with the project virtualenv:
``python tests/benchmark_operational_feishu_reports.py``.
No production connection or Feishu request is made.
"""
from __future__ import annotations

import argparse
import ctypes
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import time
import uuid

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook
from PIL import Image
from sqlalchemy import create_engine, text

from app.services import operational_feishu_reports as service
from test_operational_feishu_reports import _report_engine


def _peak_working_set() -> int:
    class Counters(ctypes.Structure):
        _fields_ = [
            ("cb", ctypes.c_ulong), ("PageFaultCount", ctypes.c_ulong),
            ("PeakWorkingSetSize", ctypes.c_size_t), ("WorkingSetSize", ctypes.c_size_t),
            ("QuotaPeakPagedPoolUsage", ctypes.c_size_t), ("QuotaPagedPoolUsage", ctypes.c_size_t),
            ("QuotaPeakNonPagedPoolUsage", ctypes.c_size_t), ("QuotaNonPagedPoolUsage", ctypes.c_size_t),
            ("PagefileUsage", ctypes.c_size_t), ("PeakPagefileUsage", ctypes.c_size_t),
        ]

    counters = Counters()
    counters.cb = ctypes.sizeof(counters)
    ctypes.windll.kernel32.GetCurrentProcess.restype = ctypes.c_void_p
    process = ctypes.windll.kernel32.GetCurrentProcess()
    get_info = ctypes.windll.psapi.GetProcessMemoryInfo
    get_info.argtypes = [ctypes.c_void_p, ctypes.POINTER(Counters), ctypes.c_ulong]
    get_info.restype = ctypes.c_bool
    if not get_info(process, ctypes.byref(counters), counters.cb):
        raise OSError("GetProcessMemoryInfo failed")
    return int(counters.PeakWorkingSetSize)


def _worker(args: argparse.Namespace) -> dict[str, object]:
    baseline_rss = _peak_working_set()
    database_path = Path(args.database)
    engine = create_engine(f"sqlite:///{database_path.as_posix()}")
    with engine.connect() as connection:
        row = connection.execute(text("SELECT map_version_id FROM source_publications WHERE source_id=:s AND publication_id=:p"),
                                 {"s": args.source, "p": uuid.UUID(args.publication).bytes}).first()
        if row is None:
            raise RuntimeError("benchmark fixture publication missing")
        map_id = bytes(row[0])
    started = time.perf_counter()
    if args.mode == "legacy":
        payload = service._payload(engine, args.source, args.publication, None)
        report_id = uuid.uuid4()
        artifact_dir = Path(args.stage)
        artifact_dir.mkdir(parents=True, exist_ok=True)
        xlsx_path = artifact_dir / f"{report_id.hex}.xlsx"
        xlsx_path.write_bytes(payload["xlsx"])
        png_content = payload["png"]
        row_count = len(payload["rows"])
        totals = payload["totals"]
        sheet_count = None
    else:
        report_id = uuid.uuid4()
        payload = service._stream_payload(engine, source_id=args.source, publication_id=args.publication,
                                          region_code=None, staging_directory=Path(args.stage), report_id=report_id)
        artifact_dir = Path(args.stage) / ".operational-feishu-reports"
        xlsx_path = artifact_dir / f"{report_id.hex}.xlsx"
        png_path = artifact_dir / f"{report_id.hex}.png"
        png_content = None
        row_count = int(payload["recordCount"])
        totals = payload["totals"]
        sheet_count = None
    elapsed = time.perf_counter() - started
    peak_rss = _peak_working_set()
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    exported_rows = 0
    sheet_names = workbook.sheetnames
    widths = []
    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, ())
        widths.append(len(header))
        exported_rows += sum(1 for _ in rows)
    workbook.close()
    if png_content is not None:
        image = Image.open(__import__("io").BytesIO(png_content))
    else:
        image = Image.open(png_path)
    image.verify()
    result = {
        "source": args.source, "mode": args.mode, "inputRows": args.rows,
        "exportedRows": exported_rows, "rowCount": row_count,
        "sheets": sheet_names, "columnWidths": widths,
        "xlsxBytes": xlsx_path.stat().st_size, "pngBytes": len(png_content) if png_content is not None else png_path.stat().st_size,
        "elapsedSeconds": round(elapsed, 3), "peakRssDeltaMiB": round(max(0, peak_rss - baseline_rss) / 1024 / 1024, 2),
        "totals": totals, "pngValid": True,
    }
    if row_count != args.rows or exported_rows != args.rows:
        raise AssertionError(f"row mismatch: expected={args.rows} returned={row_count} exported={exported_rows}")
    expected_width = 22 if args.source == "collection_rate" else 15
    if any(width != expected_width for width in widths):
        raise AssertionError(f"contract width mismatch: expected={expected_width} actual={widths}")
    engine.dispose()
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--worker", action="store_true")
    parser.add_argument("--database")
    parser.add_argument("--publication")
    parser.add_argument("--source")
    parser.add_argument("--rows", type=int)
    parser.add_argument("--mode", choices=("legacy", "stream"), default="stream")
    parser.add_argument("--stage")
    args = parser.parse_args()
    if args.worker:
        print(json.dumps(_worker(args), ensure_ascii=False))
        return

    temp_root = Path(os.environ.get("TEMP") or os.environ.get("TMP") or tempfile.gettempdir())
    temp_root.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="newbi-feishu-benchmark-", dir=temp_root) as temp_dir:
        root = Path(temp_dir)
        for source in ("collection_rate", "pdd_collection_failure"):
            for row_count in (10_000, 50_000, 100_000):
                database = root / f"{source}-{row_count}.sqlite"
                engine, publication_id, _ = _report_engine(source, database_url=f"sqlite:///{database.as_posix()}", extra_rows=row_count - 2)
                engine.dispose()
                modes = ("legacy", "stream") if row_count == 10_000 else ("stream",)
                for mode in modes:
                    stage = root / f"stage-{source}-{row_count}-{mode}"
                    completed = subprocess.run(
                        [sys.executable, str(Path(__file__).resolve()), "--worker", "--database", str(database),
                         "--publication", publication_id, "--source", source, "--rows", str(row_count),
                         "--mode", mode, "--stage", str(stage)],
                        check=False, capture_output=True, text=True, cwd=Path(__file__).resolve().parents[1],
                        env=os.environ.copy(),
                    )
                    if completed.returncode:
                        raise RuntimeError(completed.stderr.strip() or f"benchmark worker exited {completed.returncode}")
                    print(completed.stdout.strip(), flush=True)


if __name__ == "__main__":
    main()
